"""Cobertura da exportação CSV do Raio-X CRM."""

import csv
import io
import sys
import unittest
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import polars as pl
from fastapi import HTTPException

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))
from backend.api.services.analytics import crm_export  # noqa: E402

# Índices das colunas do CSV.
NUM, CRM, NOME, SITUACAO, VALOR = 3, 4, 5, 6, 7


def _ok(df):
    return SimpleNamespace(error=None, df=df)


class CrmRaioxExportTests(unittest.TestCase):
    def setUp(self):
        self.transactions = pl.DataFrame({
            "dt_janela": ["2024-01-03", "2024-01-02", "2024-02-01", "2024-01-03"],
            "data_hora": ["2024-01-03 09:20:00", "2024-01-02 08:10:00", "2024-02-01 10:00:00", "2024-01-03 09:25:00"],
            "num_autorizacao": ["0002", "0001", "0003", "0004"],
            "id_medico": ["123/PA", "123/PA", "456/PA", "456/PA"],
            "valor_pago": [19.5, 10.25, 5.0, 7.0],
        })
        self.medicos = pl.DataFrame({
            "id_medico": ["123/PA", "456/PA"],
            "no_medico": ["Médico Teste", "=HYPERLINK(\"x\")"],
        })
    def _export(self, inicio=None, fim=None, transactions=None):
        with patch.object(
            crm_export,
            "sync_crm_raiox_tx",
            return_value=_ok(self.transactions if transactions is None else transactions),
        ), patch.object(crm_export, "scan_dados_medico", return_value=self.medicos.lazy()):
            filename, chunks = crm_export.export_crm_raiox_csv("11.483.531/0001-07", inicio, fim)
            body = b"".join(chunks)
        return filename, list(csv.reader(io.StringIO(body.decode("utf-8-sig")), delimiter=";")), body

    def test_export_is_sorted_and_filtered_with_doctor_names(self):
        filename, rows, body = self._export(date(2024, 1, 1), date(2024, 1, 31))
        self.assertTrue(body.startswith(b"\xef\xbb\xbf"))
        self.assertEqual(filename, "crm_raiox_dias_alertados_11483531000107_2024-01-01_2024-01-31.csv")
        self.assertEqual(len(rows), 4)
        self.assertEqual(rows[1], [
            "11.483.531/0001-07", "02/01/2024", "08:10:00", '="0001"', "123/PA", "Médico Teste",
            "Localizado", "10,25",
        ])
        self.assertEqual([row[NUM] for row in rows[1:]], ['="0001"', '="0002"', '="0004"'])
        self.assertEqual(rows[2][1:4], ["03/01/2024", "09:20:00", '="0002"'])

    def test_header_has_no_alert_columns(self):
        _, rows, _ = self._export()
        self.assertEqual(rows[0], [
            "CNPJ", "Data", "Hora", "Número da autorização", "CRM/UF",
            "Nome do médico", "Situação do CRM", "Valor pago (R$)",
        ])

    def test_authorization_number_is_exported_as_excel_text_literal(self):
        long_number = self.transactions.with_columns(
            pl.Series("num_autorizacao", ["000123456789012345678", 'x"=1', "0003", "0004"])
        )
        _, rows, body = self._export(transactions=long_number)
        numbers = {row[NUM] for row in rows[1:]}
        self.assertIn('="000123456789012345678"', numbers)
        self.assertIn('="x""=1"', numbers)
        self.assertIn(b'"=""000123456789012345678"""', body)

    def test_csv_escapes_spreadsheet_formula_in_doctor_name(self):
        _, rows, _ = self._export()
        by_num = {row[NUM]: row for row in rows[1:]}
        self.assertEqual(by_num['="0003"'][NOME], "'=hyperlink(\"x\")")

    def test_doctor_name_is_exported_in_title_case(self):
        self.medicos = pl.DataFrame({
            "id_medico": ["123/PA", "456/PA"],
            "no_medico": ["JOAO ANTONIO SPOTT DE OLIVEIRA-BOZA", "MARIA/LIDIA"],
        })
        _, rows, _ = self._export()
        nomes = {row[CRM]: row[NOME] for row in rows[1:]}
        self.assertEqual(nomes["123/PA"], "Joao Antonio Spott De Oliveira-Boza")
        self.assertEqual(nomes["456/PA"], "Maria/Lidia")

    def test_crm_not_found_in_registry_is_flagged_in_own_column(self):
        unknown = self.transactions.with_columns(
            pl.Series("id_medico", ["123/PA", "4200667/SC", "456/PA", "456/PA"])
        )
        _, rows, _ = self._export(transactions=unknown)
        by_crm = {row[CRM]: row for row in rows[1:]}
        self.assertEqual(by_crm["4200667/SC"][NOME:VALOR], ["", "Não localizado na base do CFM"])
        self.assertEqual(by_crm["123/PA"][NOME:VALOR], ["Médico Teste", "Localizado"])
        self.assertEqual(rows[0][SITUACAO], "Situação do CRM")

    def test_csv_escapes_tab_and_carriage_return_prefixes(self):
        self.assertEqual(crm_export._csv_text("\t=1+1"), "'\t=1+1")
        self.assertEqual(crm_export._csv_text("\r=1+1"), "'\r=1+1")
        self.assertEqual(crm_export._csv_text("  @SUM(A1)"), "'  @SUM(A1)")
        self.assertEqual(crm_export._csv_text("123/PA"), "123/PA")

    def test_invalid_range_and_empty_range_are_explicit(self):
        with self.assertRaises(HTTPException) as invalid:
            crm_export.export_crm_raiox_csv("11483531000107", date(2024, 2, 1), date(2024, 1, 1))
        self.assertEqual(invalid.exception.status_code, 422)
        with self.assertRaises(HTTPException) as empty:
            self._export(date(2025, 1, 1), date(2025, 1, 31))
        self.assertEqual(empty.exception.status_code, 404)

    def test_missing_required_transaction_field_fails_before_stream(self):
        missing = self.transactions.drop("valor_pago")
        with self.assertRaises(HTTPException) as ctx:
            self._export(transactions=missing)
        self.assertEqual(ctx.exception.status_code, 503)
        self.assertIn("valor_pago", ctx.exception.detail)

    def test_null_value_and_invalid_timestamp_fail_before_stream(self):
        null_value = self.transactions.with_columns(pl.lit(None).cast(pl.Float64).alias("valor_pago"))
        with self.assertRaises(HTTPException) as null_ctx:
            self._export(transactions=null_value)
        self.assertEqual(null_ctx.exception.status_code, 500)
        self.assertIn("campo obrigatório", null_ctx.exception.detail)
        invalid_time = self.transactions.with_columns(pl.lit("data inválida").alias("data_hora"))
        with self.assertRaises(HTTPException) as time_ctx:
            self._export(transactions=invalid_time)
        self.assertEqual(time_ctx.exception.status_code, 500)
        self.assertIn("data/hora inválida", time_ctx.exception.detail)


class CrmRaioxXlsxExportTests(unittest.TestCase):
    """Reaproveita os fixtures do CSV e valida a pasta de trabalho Excel."""

    def setUp(self):
        CrmRaioxExportTests.setUp(self)
        self.perfil = pl.DataFrame({
            "cnpj": ["11483531000107"],
            "razao_social": ["FARMA TESTE LTDA"],
            "no_municipio": ["SAO JOAO BATISTA"],
            "uf": ["sc"],
        })

    def _export_xlsx(self, inicio=None, fim=None, transactions=None, perfil=None):
        from openpyxl import load_workbook

        with patch.object(crm_export, "sync_crm_raiox_tx",
                          return_value=_ok(self.transactions if transactions is None else transactions)), \
            patch.object(crm_export, "scan_dados_medico", return_value=self.medicos.lazy()), \
            patch.object(crm_export, "get_df_perfil_estabelecimento",
                         return_value=self.perfil if perfil is None else perfil):
            filename, content = crm_export.export_crm_raiox_xlsx("11.483.531/0001-07", inicio, fim)
        return filename, load_workbook(io.BytesIO(content))

    def test_workbook_has_detail_and_summary_sheets(self):
        filename, wb = self._export_xlsx()
        self.assertTrue(filename.endswith(".xlsx"))
        self.assertEqual(wb.sheetnames, ["Autorizações", "Por dia", "Por médico"])

    def test_detail_sheet_uses_real_types_and_title_case(self):
        _, wb = self._export_xlsx()
        ws = wb["Autorizações"]
        self.assertEqual(ws["A2"].value, "Raio-X CRM · Autorizações dos dias alertados")
        self.assertEqual(ws["A3"].value, "Farma Teste Ltda")
        self.assertIn("CNPJ 11.483.531/0001-07", ws["A4"].value)
        self.assertIn("Sao Joao Batista/SC", ws["A4"].value)
        self.assertIn("Período: 02/01/2024 a 01/02/2024", ws["A5"].value)
        header = [c.value for c in ws[10]]
        self.assertEqual(header, ["Data", "Hora", "Nº da autorização", "CRM/UF", "Nome do médico",
                                  "Situação do CRM", "Valor pago"])
        primeira = [c.value for c in ws[11]]
        self.assertEqual(primeira[0].date().isoformat(), "2024-01-02")
        self.assertEqual(primeira[1].isoformat(), "08:10:00")
        self.assertEqual(primeira[2], "0001")  # texto, zeros preservados
        self.assertEqual(primeira[4], "Médico Teste")
        self.assertEqual(primeira[6], 10.25)
        self.assertEqual(ws["A11"].number_format, "dd/mm/yyyy")
        self.assertEqual(ws["G11"].number_format, '"R$" #,##0.00')
        self.assertEqual(ws.freeze_panes, "A11")

    def test_detail_sheet_kpis_and_total_row(self):
        _, wb = self._export_xlsx()
        ws = wb["Autorizações"]
        self.assertEqual(ws["A8"].value, 4)            # autorizações
        self.assertEqual(ws["C8"].value, 3)            # dias alertados
        self.assertEqual(ws["D8"].value, 2)            # médicos distintos
        self.assertEqual(ws["E8"].value, 0)            # CRMs não localizados
        self.assertAlmostEqual(ws["F8"].value, 41.75)  # valor total
        total = [c.value for c in ws[15]]
        self.assertEqual(total[0], "Total")
        self.assertEqual(total[2], "=SUBTOTAL(103,C11:C14)")
        self.assertEqual(total[6], "=SUBTOTAL(109,G11:G14)")

    def test_formula_like_text_is_stored_as_text(self):
        _, wb = self._export_xlsx()
        ws = wb["Autorizações"]
        nomes = [ws.cell(row=r, column=5).value for r in range(11, 15)]
        self.assertIn('=hyperlink("x")', nomes)
        self.assertEqual(ws.cell(row=11, column=5).data_type, "s")

    def test_summary_sheets_totals(self):
        _, wb = self._export_xlsx()
        por_dia = wb["Por dia"]
        dias = [(r[0].value.date().isoformat(), r[1].value, r[3].value)
                for r in por_dia.iter_rows(min_row=8, max_row=10)]
        self.assertEqual(dias, [("2024-01-02", 1, 10.25), ("2024-01-03", 2, 26.5), ("2024-02-01", 1, 5.0)])
        por_medico = wb["Por médico"]
        primeiro = [c.value for c in por_medico[8]]
        self.assertEqual(primeiro[:3], ["123/PA", "Médico Teste", "Localizado"])
        self.assertAlmostEqual(primeiro[5], 29.75)
        self.assertAlmostEqual(primeiro[6], 29.75 / 41.75)

    def test_unknown_crm_is_flagged_in_both_sheets(self):
        unknown = self.transactions.with_columns(
            pl.Series("id_medico", ["123/PA", "4200667/SC", "456/PA", "456/PA"])
        )
        _, wb = self._export_xlsx(transactions=unknown)
        self.assertEqual(wb["Autorizações"]["E8"].value, 1)
        situacoes = {r[0].value: r[2].value for r in wb["Por médico"].iter_rows(min_row=8, max_row=10)}
        self.assertEqual(situacoes["4200667/SC"], "Não localizado na base do CFM")

    def test_missing_profile_is_explicit(self):
        with self.assertRaises(HTTPException) as ctx:
            self._export_xlsx(perfil=self.perfil.clear())
        self.assertEqual(ctx.exception.status_code, 404)


if __name__ == "__main__":
    unittest.main()
