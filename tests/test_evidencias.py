"""Testes da cesta de evidências (serviço e endpoints)."""

import io
import json
import sys
import tempfile
import types
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))
try:
    import backend.api.services.analytics.crm_export  # noqa: F401
except ModuleNotFoundError:
    # Máquina sem as dependências do cache (SQL Server/Parquet): o pacote analytics
    # entra sem o __init__ (que importa tudo) e os produtores de cache viram stubs.
    # Só serve para importar; a farmácia é sempre substituída nos testes.
    import backend.api.services  # noqa: F401
    for _nome in [n for n in sys.modules if n.startswith("backend.api.services.analytics")]:
        del sys.modules[_nome]
    _pkg = types.ModuleType("backend.api.services.analytics")
    _pkg.__path__ = [str(Path(__file__).resolve().parents[1] / "backend" / "api" / "services" / "analytics")]
    sys.modules["backend.api.services.analytics"] = _pkg
    for _nome, _attrs in {
        "cache_producers": {},
        "cache_producers.crm": {"sync_crm_raiox_tx": None},
        "data_cache": {"get_df_perfil_estabelecimento": None, "scan_dados_medico": None},
    }.items():
        _mod = types.ModuleType(_nome)
        _mod.__dict__.update(_attrs)
        sys.modules[_nome] = _mod

from fastapi import FastAPI
from fastapi.testclient import TestClient

from backend.api.endpoints import evidencias as evidencias_endpoints
from backend.api.services.evidencias import (
    EvidenciaDuplicadaError,
    EvidenciaNaoEncontradaError,
    EvidenciasError,
    EvidenciasService,
)
from backend.api.services import evidencias_export
from backend.api.services.analytics.crm_export import _Farmacia
from backend.api.services.preferences import PreferencesService

CNPJ_A = "04570047000141"
CNPJ_B = "05363613000107"


def dia(cnpj=CNPJ_A, data="2021-01-29", **extra):
    return {"cnpj": cnpj, "tipo": "dia", "dt_janela": data, "hora": None,
            "num_autorizacao": None, "snapshot": {"qtd": 21}, "nota": "", **extra}


class EvidenciasBase(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.original_dir = PreferencesService.BASE_DIR
        PreferencesService._set_base_dir(Path(self.temp.name))

    def tearDown(self):
        PreferencesService._set_base_dir(self.original_dir)
        self.temp.cleanup()


class EvidenciasServiceTests(EvidenciasBase):
    def test_primeiro_uso_lista_vazia_sem_criar_arquivo(self):
        self.assertEqual(EvidenciasService.listar(), [])
        self.assertFalse(EvidenciasService._file_path().exists())

    def test_criar_listar_e_filtrar_por_cnpj(self):
        a = EvidenciasService.criar(dia(nota="  rajada  "))
        EvidenciasService.criar(dia(cnpj=CNPJ_B))
        self.assertEqual(a["nota"], "rajada")
        self.assertEqual(len(EvidenciasService.listar()), 2)
        self.assertEqual([i["id"] for i in EvidenciasService.listar(CNPJ_A)], [a["id"]])

    def test_duplicidade_por_tipo(self):
        EvidenciasService.criar(dia())
        with self.assertRaises(EvidenciaDuplicadaError):
            EvidenciasService.criar(dia())
        hora = {**dia(), "tipo": "hora", "hora": 11}
        EvidenciasService.criar(hora)
        EvidenciasService.criar({**hora, "hora": 12})
        with self.assertRaises(EvidenciaDuplicadaError):
            EvidenciasService.criar(hora)
        aut = {**dia(), "tipo": "autorizacao", "hora": 11, "num_autorizacao": "123"}
        EvidenciasService.criar(aut)
        with self.assertRaises(EvidenciaDuplicadaError):
            EvidenciasService.criar({**aut, "hora": 12})

    def test_atualizar_nota_e_remover(self):
        a = EvidenciasService.criar(dia())
        atualizado = EvidenciasService.atualizar_nota(a["id"], " nova ")
        self.assertEqual(atualizado["nota"], "nova")
        EvidenciasService.remover(a["id"])
        self.assertEqual(EvidenciasService.listar(), [])
        with self.assertRaises(EvidenciaNaoEncontradaError):
            EvidenciasService.remover(a["id"])
        with self.assertRaises(EvidenciaNaoEncontradaError):
            EvidenciasService.atualizar_nota(a["id"], "x")

    def test_remover_por_cnpj_e_resumo(self):
        EvidenciasService.criar(dia())
        EvidenciasService.criar(dia(data="2021-01-30"))
        EvidenciasService.criar(dia(cnpj=CNPJ_B))
        resumo = {r["cnpj"]: r["quantidade"] for r in EvidenciasService.resumo_por_cnpj()}
        self.assertEqual(resumo, {CNPJ_A: 2, CNPJ_B: 1})
        self.assertEqual(EvidenciasService.remover_por_cnpj(CNPJ_A), 2)
        self.assertEqual(EvidenciasService.remover_por_cnpj(CNPJ_A), 0)
        self.assertEqual([i["cnpj"] for i in EvidenciasService.listar()], [CNPJ_B])

    def test_gravacao_mantem_backup_da_versao_anterior(self):
        EvidenciasService.criar(dia())
        EvidenciasService.criar(dia(data="2021-01-30"))
        backup = json.loads(EvidenciasService._backup_path().read_text(encoding="utf-8"))
        self.assertEqual(len(backup["evidencias"]), 1)

    def test_arquivo_corrompido_falha_sem_sobrescrever(self):
        EvidenciasService._file_path().parent.mkdir(parents=True, exist_ok=True)
        EvidenciasService._file_path().write_text('{"evidencias":', encoding="utf-8")
        with self.assertRaises(EvidenciasError):
            EvidenciasService.listar()
        with self.assertRaises(EvidenciasError):
            EvidenciasService.criar(dia())
        self.assertEqual(EvidenciasService._file_path().read_text(encoding="utf-8"), '{"evidencias":')

    def test_principal_ausente_com_backup_nao_recria_vazio(self):
        EvidenciasService.criar(dia())
        EvidenciasService.criar(dia(data="2021-01-30"))
        EvidenciasService._file_path().unlink()
        with self.assertRaises(EvidenciasError):
            EvidenciasService.listar()
        self.assertFalse(EvidenciasService._file_path().exists())


class EvidenciasEndpointTests(EvidenciasBase):
    def setUp(self):
        super().setUp()
        app = FastAPI()
        app.include_router(evidencias_endpoints.router, prefix="/evidencias")
        self.client = TestClient(app)

    def test_fluxo_completo(self):
        r = self.client.post("/evidencias", json=dia())
        self.assertEqual(r.status_code, 201)
        ev = r.json()
        self.assertEqual(self.client.post("/evidencias", json=dia()).status_code, 409)
        self.assertEqual(len(self.client.get("/evidencias", params={"cnpj": CNPJ_A}).json()), 1)
        self.assertEqual(self.client.get("/evidencias/resumo").json()[0]["quantidade"], 1)
        r = self.client.patch(f"/evidencias/{ev['id']}", json={"nota": "ok"})
        self.assertEqual(r.json()["nota"], "ok")
        self.assertEqual(self.client.delete(f"/evidencias/{ev['id']}").status_code, 204)
        self.assertEqual(self.client.delete(f"/evidencias/{ev['id']}").status_code, 404)
        self.client.post("/evidencias", json=dia())
        self.assertEqual(self.client.delete(f"/evidencias/cnpj/{CNPJ_A}").json(), {"removidas": 1})

    def test_validacao_por_tipo(self):
        self.assertEqual(self.client.post("/evidencias", json={**dia(), "tipo": "hora"}).status_code, 422)
        self.assertEqual(self.client.post("/evidencias", json={**dia(), "hora": 3}).status_code, 422)
        self.assertEqual(self.client.post("/evidencias", json={**dia(), "tipo": "autorizacao", "hora": 3}).status_code, 422)
        self.assertEqual(self.client.post("/evidencias", json=dia(cnpj="123")).status_code, 422)
        self.assertEqual(self.client.post("/evidencias", json=dia(data="29/01/2021")).status_code, 422)

    def test_arquivo_corrompido_retorna_503(self):
        EvidenciasService._file_path().parent.mkdir(parents=True, exist_ok=True)
        EvidenciasService._file_path().write_text("[]", encoding="utf-8")
        r = self.client.get("/evidencias")
        self.assertEqual(r.status_code, 503)
        self.assertIn("preservado", r.json()["detail"])


class EvidenciasExportTests(EvidenciasBase):
    def setUp(self):
        super().setUp()
        app = FastAPI()
        app.include_router(evidencias_endpoints.router, prefix="/evidencias")
        self.client = TestClient(app)
        farmacia = _Farmacia(razao_social="Farmacia Teste Ltda", municipio="Descanso", uf="SC")
        self.patcher = patch.object(evidencias_export, "_load_farmacia", return_value=farmacia)
        self.patcher.start()

    def tearDown(self):
        self.patcher.stop()
        super().tearDown()

    def test_sem_evidencias_retorna_404(self):
        self.assertEqual(self.client.get(f"/evidencias/cnpj/{CNPJ_A}/exportar").status_code, 404)

    def test_planilha_com_uma_linha_por_item_em_ordem_cronologica(self):
        from openpyxl import load_workbook

        EvidenciasService.criar({**dia(data="2021-02-01"), "snapshot": {"qtd": 30, "alertas": ["Volume Atípico"]},
                                 "nota": '=HYPERLINK("x")'})
        EvidenciasService.criar({**dia(), "tipo": "hora", "hora": 11, "snapshot": {"qtd": 48, "alertas": []}})
        EvidenciasService.criar({**dia(), "tipo": "autorizacao", "hora": 11, "num_autorizacao": "0001",
                                 "snapshot": {"horario": "11:10:32", "crm": "3998/SC", "medico": "Ana Souza",
                                              "valor": 24.9, "alertas": ["U#1"]}})
        EvidenciasService.criar(dia(cnpj=CNPJ_B))

        r = self.client.get(f"/evidencias/cnpj/{CNPJ_A}/exportar")
        self.assertEqual(r.status_code, 200)
        self.assertIn(f"evidencias_{CNPJ_A}.xlsx", r.headers["content-disposition"])
        ws = load_workbook(io.BytesIO(r.content))["Evidências"]
        self.assertEqual(ws["A2"].value, "Cesta de evidências")
        self.assertEqual(ws["A3"].value, "Farmacia Teste Ltda")
        self.assertEqual(ws["A8"].value, 3)  # KPI total, só do CNPJ pedido
        cabecalho = [c.value for c in ws[10]]
        self.assertEqual(cabecalho[0], "Tipo")
        self.assertEqual(cabecalho[9], "Nota do auditor")
        linhas = [[c.value for c in row] for row in ws.iter_rows(min_row=11, max_row=13)]
        self.assertEqual([l[0] for l in linhas], ["Hora", "Autorização", "Dia"])
        self.assertEqual(linhas[0][2], "11h às 11h59")
        self.assertEqual(linhas[0][7], 48)
        self.assertEqual(linhas[1][2], "11:10:32")
        self.assertEqual(linhas[1][3], "0001")
        self.assertEqual(linhas[1][6], 24.9)
        self.assertEqual(linhas[2][2], "Dia todo")
        self.assertEqual(linhas[2][8], "Volume Atípico")
        self.assertEqual(linhas[2][9], '=HYPERLINK("x")')
        self.assertEqual(ws.cell(13, 10).data_type, "s")  # nota nunca vira fórmula


if __name__ == "__main__":
    unittest.main()
