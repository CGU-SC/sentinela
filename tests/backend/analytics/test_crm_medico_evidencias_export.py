from datetime import date, datetime
from io import BytesIO
from types import SimpleNamespace

import openpyxl
import polars as pl

from backend.api.services.analytics import crm_medico_evidencias_export as export


def _evidence_data():
    unico = pl.DataFrame(
        {
            "id_cnpj": [1], "dt": ["2024-01-03"], "dt_ini_hora": [datetime(2024, 1, 3, 9)],
            "dt_fim_hora": [datetime(2024, 1, 3, 10)], "razao_social": ["Farmacia Um"],
            "cnpj": ["11111111000111"], "municipio": ["Sao Paulo"], "uf": ["SP"],
            "nu_autorizacoes": [8], "nu_minutos": [20], "taxa_hora": [24.5], "id_severidade": [1],
        }
    )
    multiplos = pl.DataFrame(
        {
            "id_cnpj": [1], "dt": ["2024-01-03"], "dt_ini_hora": [datetime(2024, 1, 3, 9)],
            "dt_fim_hora": [datetime(2024, 1, 3, 10)], "razao_social": ["Farmacia Um"],
            "cnpj": ["11111111000111"], "municipio": ["Sao Paulo"], "uf": ["SP"],
            "nu_autorizacoes_crm": [3], "nu_autorizacoes_total": [8], "nu_crms": [2],
            "nu_minutos": [20], "taxa_hora": [8.0], "id_severidade": [2],
        }
    )
    distancia = pl.DataFrame(
        {
            "competencia": [202401], "cnpj_a": ["11111111000111"], "razao_social_a": ["Farmacia Um"],
            "no_municipio_a": ["Sao Paulo"], "sg_uf_a": ["SP"], "nu_prescricoes_a": [5],
            "cnpj_b": ["22222222000122"], "razao_social_b": ["Farmacia Dois"],
            "no_municipio_b": ["Rio de Janeiro"], "sg_uf_b": ["RJ"], "nu_prescricoes_b": [6],
            "distancia_km": [360.5], "vl_autorizacoes_total": [110.0],
        }
    )
    return SimpleNamespace(unico=unico, multiplos=multiplos, distancia=distancia)


def test_excel_export_has_three_evidence_sheets_and_typed_rows(monkeypatch):
    monkeypatch.setattr(export, "evidencias_do_medico", lambda *args: _evidence_data())
    monkeypatch.setattr(export, "nome_do_medico", lambda _: "ana maria")
    monkeypatch.setattr(export, "_period_bounds", lambda start, end: (start or date(2024, 1, 1), end or date(2024, 1, 31)))
    filename, content = export.export_crm_medico_evidencias_xlsx(
        "123/SP", date(2024, 1, 1), date(2024, 1, 31)
    )
    assert filename == "crm_evidencias_123_SP_202401_202401.xlsx"
    workbook = openpyxl.load_workbook(BytesIO(content), data_only=False)
    assert workbook.sheetnames == ["Sequências único CRM", "Sequências múltiplos CRMs", "Farmácias distantes"]
    unico, multiplo, distancia = [workbook[name] for name in workbook.sheetnames]
    assert unico["A9"].value == datetime(2024, 1, 3)
    assert unico["B9"].value.hour == 9 and unico["B9"].value.minute == 0
    assert unico["G9"].value == 8 and unico["J9"].value == "Alta"
    assert multiplo["G9"].value == 3 and multiplo["H9"].value == 8 and multiplo["L9"].value == "Grave"
    assert distancia["A9"].value == "01/2024" and distancia["J9"].value == 360
    assert "Ana Maria" in unico["A3"].value


def test_export_helpers_write_blank_date_numeric_and_string_cells_and_empty_notice():
    assert export._soma(pl.DataFrame(schema={"x": pl.Int64}), "x") == 0.0
    assert export._soma(pl.DataFrame({"x": [2.5, 3.5]}), "x") == 6.0
    assert export._data("2024-03-04 12:00") == date(2024, 3, 4)

    import xlsxwriter

    stream = BytesIO()
    workbook = xlsxwriter.Workbook(stream, {"in_memory": True})
    formats = export._formats(workbook)
    worksheet = workbook.add_worksheet()
    export._tabela(
        worksheet,
        formats,
        [("Col A", 10, "texto"), ("Col B", 10, "data"), ("Col C", 10, "inteiro")],
        [["texto", date(2024, 1, 2), 4], [None, None, 5]],
        [("label", "Total"), ("contagem", 2, "total_inteiro"), ("vazio",)],
    )
    empty = workbook.add_worksheet("Vazia")
    export._tabela(empty, formats, [("Coluna", 10, "texto")], [], [])
    workbook.close()
    reopened = openpyxl.load_workbook(BytesIO(stream.getvalue()), data_only=False)
    assert reopened.worksheets[0]["A9"].value == "texto"
    assert reopened.worksheets[0]["B9"].value == datetime(2024, 1, 2)
    assert reopened.worksheets[0]["C9"].value == 4
    assert reopened["Vazia"]["A9"].value == "Nenhuma evidência no período."


def test_export_applies_pharmacy_and_municipality_labels(monkeypatch):
    monkeypatch.setattr(export, "evidencias_do_medico", lambda *args: _evidence_data())
    monkeypatch.setattr(export, "nome_do_medico", lambda _: None)
    monkeypatch.setattr(export, "_period_bounds", lambda start, end: (date(2024, 1, 1), date(2024, 1, 31)))
    monkeypatch.setattr(export, "_cadastro", lambda _: pl.DataFrame({"cnpj": ["11111111000111"], "razao_social": ["Farmacia Um"]}))
    monkeypatch.setattr(export, "nome_do_municipio", lambda _: ("sao paulo", "SP"))

    pharmacy_file, pharmacy_bytes = export.export_crm_medico_evidencias_xlsx("123/SP", id_cnpj=1)
    city_file, city_bytes = export.export_crm_medico_evidencias_xlsx("123/SP", id_ibge7=3550308)
    assert pharmacy_file == city_file
    assert len(pharmacy_bytes) > 0 and len(city_bytes) > 0

