from datetime import date
from types import SimpleNamespace

import polars as pl
import pytest
from docx import Document
from fastapi import HTTPException

from backend.api.services.analytics import nota_tecnica_criticidades as criticidades


def test_criticality_formatting_and_validation_helpers():
    assert criticidades._normalize_ascii_upper("  Doença_de Parkinson ") == "DOENCA DE PARKINSON"
    assert criticidades._clinica_meta_key("Doença de Parkinson", "idade menor_50") == (
        "DOENCA DE PARKINSON", "IDADE_MENOR_50"
    )
    assert criticidades._cnpj_digits("12.345.678/0001-95") == "12345678000195"
    assert criticidades._format_cnpj_pt("12345678000195") == "12.345.678/0001-95"
    assert criticidades._format_int_pt(1234.6) == "1.235"
    assert criticidades._format_brl_pt(1234.5) == "R$ 1.234,50"
    assert criticidades._format_optional_decimal(None) == "não calculado"
    assert criticidades._format_optional_decimal(1.5) == "1,50"
    assert criticidades._format_optional_percent(None) == "não calculado"
    assert criticidades._format_optional_percent(25) == "25,00%"
    assert criticidades._format_optional_ratio_percent(0.25) == "25,00%"
    assert criticidades._format_optional_ratio_percent(None) == "não calculado"
    assert criticidades._ratio(1, 0) is None
    assert criticidades._ratio(2, 4) == 0.5
    assert criticidades._format_periodo_anos_item({"ano_inicio": 2020, "ano_fim": 2020}) == "em 2020"
    assert criticidades._format_periodo_anos_item({"ano_inicio": 2020, "ano_fim": 2024}) == "no período de 2020 a 2024"
    assert criticidades._sum_numeric([{"n": 2}, {"n": 3}], "n") == 5
    assert criticidades._format_indicador_quadro_value(12.345, "pct3") == "12,345%"
    assert criticidades._format_indicador_quadro_value(12.345, "pct") == "12,35%"
    assert criticidades._format_indicador_quadro_value(12.345, "dec") == "12,35"
    assert criticidades._format_indicador_quadro_value(12.345, "val") == "R$ 12,35"
    assert criticidades._format_indicador_quadro_optional(None, "val") == "—"
    assert criticidades._format_risco_regional(2) == "2,00x"
    assert criticidades._format_risco_regional(None) == "—"
    assert criticidades._indicador_valor_farmacia_header("val") == "Valor Farmácia"
    assert criticidades._indicador_valor_farmacia_header("pct3") == "Percentual Farmácia"
    assert criticidades._indicador_valor_farmacia_header("dec") == "Índice Farmácia"
    assert criticidades._vez_ou_vezes("1,00") == "vez"
    assert criticidades._vez_ou_vezes("1,01") == "vezes"
    assert criticidades._vez_ou_vezes("invalido") == "vezes"
    assert criticidades._is_parkinson_menor_50_item({
        "patologia": "Doença de Parkinson", "regra_clinica": "IDADE_MENOR_50"
    })
    assert criticidades._count_incompatibilidade_patologica_tables({"ranking_patologias": [
        {"patologia": "Doença de Parkinson", "regra_clinica": "IDADE_MENOR_50"},
        {"patologia": "Diabetes", "regra_clinica": "IDADE_MENOR_20"},
    ]}) == 7

    with pytest.raises(RuntimeError, match="CNPJ invalido"):
        criticidades._cnpj_digits("123")
    with pytest.raises(RuntimeError, match="obrigatorias"):
        criticidades._require_columns(pl.DataFrame({"a": [1]}), {"b"}, "fonte")
    with pytest.raises(RuntimeError, match="Valor inteiro obrigatorio"):
        criticidades._format_int_pt("bad")
    with pytest.raises(RuntimeError, match="Valor financeiro obrigatorio"):
        criticidades._format_brl_pt("bad")
    with pytest.raises(RuntimeError, match="decimal opcional invalido"):
        criticidades._format_optional_decimal("bad")
    with pytest.raises(RuntimeError, match="proporcional opcional invalido"):
        criticidades._format_optional_ratio_percent("bad")
    with pytest.raises(RuntimeError, match="Razao obrigatoria invalida"):
        criticidades._ratio("bad", 1)
    with pytest.raises(RuntimeError, match="Periodo anual obrigatorio"):
        criticidades._format_periodo_anos_item({})
    with pytest.raises(RuntimeError, match="Periodo anual invalido"):
        criticidades._format_periodo_anos_item({"ano_inicio": 2024, "ano_fim": 2020})
    with pytest.raises(RuntimeError, match="Valor numerico obrigatorio"):
        criticidades._sum_numeric([{"n": "bad"}], "n")
    with pytest.raises(RuntimeError, match="sem valor calculado"):
        criticidades._format_indicador_quadro_value(None, "pct")
    with pytest.raises(RuntimeError, match="Formato de indicador critico"):
        criticidades._format_indicador_quadro_value(1, "unknown")
    with pytest.raises(RuntimeError, match="Valor de indicador critico invalido"):
        criticidades._format_indicador_quadro_value("invalido", "pct")
    with pytest.raises(RuntimeError, match="cabecalho"):
        criticidades._indicador_valor_farmacia_header("unknown")


def test_criticality_zero_baseline_and_regional_sort_rules():
    assert criticidades._optional_float(None) is None
    assert criticidades._optional_float("bad") is None
    assert criticidades._optional_float(float("nan")) is None
    assert criticidades._optional_float("2.5") == 2.5
    assert criticidades._is_zero_baseline_critical("falecidos", 2, 0, None)
    assert not criticidades._is_zero_baseline_critical("teto", 2, 0, None)
    assert not criticidades._is_zero_baseline_critical("falecidos", 0, 0, None)
    assert criticidades._risco_regional_sort_value("teto", 2, 1, 3) == 3
    assert criticidades._risco_regional_sort_value("falecidos", 2, 0, None) == float("inf")
    with pytest.raises(RuntimeError, match="sem risco regional"):
        criticidades._risco_regional_sort_value("teto", 2, 0, None)


def _matrix_row_for_indicators(cnpj="12345678000195", critical_keys=None):
    critical_keys = set(critical_keys if critical_keys is not None else criticidades._INDICATOR_FLAGS)
    row = {"cnpj": cnpj, "id_regiao_saude": "R-1"}
    for key, mapping in criticidades.INDICATOR_MAPPING.items():
        value_col, mediana_col, _, _, risk_col, _, _ = mapping
        row[value_col] = 10.0
        row[mediana_col] = 5.0
        row[risk_col] = 2.0
        flags = criticidades._INDICATOR_FLAGS.get(key)
        if flags:
            row[flags[0].lower()] = 0
            row[flags[1].lower()] = int(key in critical_keys)
    row[criticidades._VOLUME_ATIPICO_VALOR_AUMENTO_COL] = 123.45
    return row


def test_dynamic_critical_indicators_are_selected_ordered_and_formatted(monkeypatch):
    row = _matrix_row_for_indicators()
    frame = pl.DataFrame([row])
    monkeypatch.setattr(criticidades, "_get_matriz_dinamica_nota", lambda *args: (frame, row))

    found = criticidades._get_criticos("12.345.678/0001-95")
    assert "teto" in found and "percentual_nao_comprovacao" in found
    ordered = criticidades._get_criticos_ordenados_por_risco("12345678000195", found)
    assert ordered[0] == "percentual_nao_comprovacao"
    assert ordered.index("falecidos") < ordered.index("teto")

    summary = criticidades._build_indicadores_criticos_quadro("12345678000195")
    by_key = {item["key"]: item for item in summary}
    assert by_key["volume_atipico"]["valor"] == "R$ 123,45"
    assert by_key["volume_atipico"]["bookmark"] == "tabela_evolucao_financeira"
    assert by_key["percentual_nao_comprovacao"]["bookmark"] == "secao6_percentual_nao_comprovacao"
    assert all(item["status"] == "CRÍTICO" for item in summary)


def test_critical_indicator_ordering_and_matrix_row_access_empty_and_success(monkeypatch):
    assert criticidades._get_criticos_ordenados_por_risco("12345678000195", set()) == []
    row = _matrix_row_for_indicators()
    frame = pl.DataFrame([row])
    monkeypatch.setattr(criticidades, "_get_matriz_dinamica_nota", lambda *args: (frame, row))
    assert criticidades._get_matriz_row_context("12345678000195", None, None, "indicador") == row


def test_matrix_metric_contexts_return_none_when_target_is_absent(monkeypatch):
    monkeypatch.setattr(criticidades, "_get_matriz_row_context", lambda *_args: None)
    builders = (
        criticidades._build_teto_context,
        criticidades._build_polimedicamento_context,
        criticidades._build_ticket_medio_context,
        criticidades._build_receita_paciente_context,
        criticidades._build_per_capita_context,
        criticidades._build_alto_custo_context,
        criticidades._build_vendas_rapidas_context,
        criticidades._build_recorrencia_sistemica_context,
        criticidades._build_dias_pico_context,
        criticidades._build_incompatibilidade_patologica_context,
    )

    for builder in builders:
        assert builder("12345678000195", None, None) is None


def test_metric_contexts_use_unbounded_period_and_coerce_invalid_numbers(monkeypatch):
    row = {key: "invalido" for key in _matrix_row_for_indicators()}
    monkeypatch.setattr(criticidades, "_get_matriz_row_context", lambda *_args: row)
    builders = (
        criticidades._build_polimedicamento_context,
        criticidades._build_ticket_medio_context,
        criticidades._build_receita_paciente_context,
        criticidades._build_per_capita_context,
        criticidades._build_alto_custo_context,
        criticidades._build_vendas_rapidas_context,
        criticidades._build_recorrencia_sistemica_context,
        criticidades._build_dias_pico_context,
    )

    for builder in builders:
        context = builder("12345678000195", None, None)
        assert context["periodo_desc"] == "no período analisado"
        assert 0.0 in context.values()


@pytest.mark.parametrize(
    ("data_inicio", "data_fim", "expected"),
    [
        (date(2024, 1, 1), None, "a partir de janeiro de 2024"),
        (None, date(2024, 3, 1), "até março de 2024"),
        (None, None, "no período analisado"),
    ],
)
def test_deceased_context_formats_open_ended_periods(monkeypatch, data_inicio, data_fim, expected):
    monkeypatch.setattr(
        criticidades,
        "get_falecidos_data",
        lambda *_args: SimpleNamespace(
            transacoes=[SimpleNamespace(cpf="123", data_autorizacao=None)],
            summary=SimpleNamespace(total_autorizacoes=1, cpfs_distintos=1, valor_total=10.0),
        ),
    )

    result = criticidades._build_falecidos_context(
        "12345678000195", "SP", data_inicio, data_fim
    )

    assert result["periodo_desc"] == expected

    monkeypatch.setattr(criticidades, "_FORCAR_TODOS_CRITICOS_NOTA_TECNICA", True)
    assert criticidades._get_criticos("12345678000195") == {
        key for key, _, _ in criticidades._SECAO5_MAP
    }
    assert criticidades._get_criticos_ordenados_por_risco("12345678000195", {"teto", "falecidos"}) == [
        "falecidos", "teto"
    ]


def test_dynamic_critical_indicator_errors_fail_visibly(monkeypatch):
    monkeypatch.setattr(
        criticidades, "_get_matriz_dinamica_nota",
        lambda *args: (_ for _ in ()).throw(RuntimeError("matriz offline")),
    )
    with pytest.raises(RuntimeError, match="Falha ao identificar indicadores criticos"):
        criticidades._get_criticos("12345678000195")
    with pytest.raises(RuntimeError, match="Matriz de risco dinamica indisponivel"):
        criticidades._get_matriz_row_context("12345678000195", None, None, "teto")

    bad = {"cnpj": "12345678000195"}
    monkeypatch.setattr(criticidades, "_get_matriz_dinamica_nota", lambda *args: (pl.DataFrame([bad]), bad))
    with pytest.raises(RuntimeError, match="sem mapeamento de colunas"):
        criticidades._get_criticos_ordenados_por_risco("12345678000195", {"not_mapped"})
    with pytest.raises(RuntimeError, match="sem colunas obrigatorias"):
        criticidades._get_criticos_ordenados_por_risco("12345678000195", {"teto"})
    with pytest.raises(RuntimeError, match="sem colunas obrigatorias"):
        criticidades._build_indicadores_criticos_quadro("12345678000195")


def test_clinical_context_and_parkinson_demographics_use_validated_sources(monkeypatch):
    cnpj = "12345678000195"
    monkeypatch.setattr(criticidades, "get_df_perfil_estabelecimento", lambda: pl.DataFrame({"cnpj": [cnpj]}))
    valid = pl.DataFrame({"cnpj": [cnpj, "99999999000191"], "score": [1, 2]})
    monkeypatch.setattr(criticidades, "_build_dynamic_matriz_risco", lambda **kwargs: valid)
    full, target = criticidades._get_matriz_dinamica_nota(cnpj)
    assert full.height == 2 and target["score"] == 1
    monkeypatch.setattr(criticidades, "_build_dynamic_matriz_risco", lambda **kwargs: pl.DataFrame())
    with pytest.raises(RuntimeError, match="sem linhas"):
        criticidades._get_matriz_dinamica_nota(cnpj)
    monkeypatch.setattr(criticidades, "_build_dynamic_matriz_risco", lambda **kwargs: pl.DataFrame({"cnpj": ["99999999000191"]}))
    with pytest.raises(RuntimeError, match="nao encontrado"):
        criticidades._get_matriz_dinamica_nota(cnpj)
    monkeypatch.setattr(criticidades, "_build_dynamic_matriz_risco", lambda **kwargs: pl.DataFrame({"cnpj": [cnpj, cnpj]}))
    with pytest.raises(RuntimeError, match="mais de uma linha"):
        criticidades._get_matriz_dinamica_nota(cnpj)

    demographic = pl.DataFrame(
        {
            "id_ibge7": ["1234567"] * 3, "ano_censo": [2022] * 3,
            "idade_min": [0, 50, 80], "nu_populacao": [500, 400, 100],
        }
    )
    monkeypatch.setattr(criticidades, "get_df_dados_ibge_demografia", lambda: demographic)
    context = criticidades._build_parkinson_demografia_context(
        {"id_ibge7": 1234567, "municipio": "SAO PAULO", "uf": "sp"},
        [{"qtd_cpfs_distintos": 4, "ano_base": 2023}, {"qtd_cpfs_distintos": 5, "ano_base": 2022}],
    )
    assert context["populacao_total"] == 1000
    assert context["populacao_50_mais"] == 500
    assert context["faixas_etarias"][-1]["faixa"] == "80+"
    assert context["ano_observado"] == 2022
    with pytest.raises(RuntimeError, match="Evolucao anual de Parkinson obrigatoria"):
        criticidades._build_parkinson_demografia_context({}, [])
    with pytest.raises(RuntimeError, match="municipio/UF/id_ibge7"):
        criticidades._build_parkinson_demografia_context({}, [{"qtd_cpfs_distintos": 1, "ano_base": 2020}])


def test_regional_indicator_context_ranks_within_id_region(monkeypatch):
    key = "ticket_medio"
    value_col, median_col, _, _, risk_col, _, _ = criticidades.INDICATOR_MAPPING[key]
    frame = pl.DataFrame(
        {
            "cnpj": ["12345678000195", "99999999000191", "22222222000191"],
            "id_regiao_saude": ["R1", "R1", "R2"],
            value_col: [10.0, 20.0, 99.0], median_col: [5.0, 10.0, 5.0], risk_col: [2.0, 4.0, 9.0],
        }
    )
    monkeypatch.setattr(criticidades, "_get_matriz_dinamica_nota", lambda *args: (frame, frame.row(0, named=True)))

    result = criticidades._build_indicador_regional_context("12345678000195", key)
    assert result["id_regiao_saude"] == "R1"
    assert result["posicao_regional"] == 2 and result["total_regional"] == 2
    assert result["valor"] == "R$ 10,00" and result["risco_regional"] == 2.0
    assert result["percentil_regional"] == 0.0
    with pytest.raises(RuntimeError, match="Indicador sem mapeamento"):
        criticidades._build_indicador_regional_context("12345678000195", "unknown")


def test_falecidos_context_counts_distinct_cpfs_and_uses_requested_period(monkeypatch):
    transacoes = [
        SimpleNamespace(cpf="111", data_autorizacao=date(2024, 1, 5)),
        SimpleNamespace(cpf="111", data_autorizacao=date(2024, 2, 5)),
        SimpleNamespace(cpf="222", data_autorizacao=None),
    ]
    falecidos = SimpleNamespace(
        transacoes=transacoes,
        summary=SimpleNamespace(total_autorizacoes=7, cpfs_distintos=99, valor_total=123.0),
    )
    monkeypatch.setattr(criticidades, "get_falecidos_data", lambda *args: falecidos)
    result = criticidades._build_falecidos_context("12345678000195", "DF", None, None)
    assert result["total_autorizacoes"] == 7 and result["cpfs_distintos"] == 2
    assert "janeiro" in result["periodo_desc"] and "fevereiro" in result["periodo_desc"]
    result = criticidades._build_falecidos_context(
        "12345678000195", "DF", date(2023, 1, 1), date(2023, 12, 31)
    )
    assert "2023" in result["periodo_desc"]
    monkeypatch.setattr(criticidades, "get_falecidos_data", lambda *args: SimpleNamespace(transacoes=[]))
    assert criticidades._build_falecidos_context("12345678000195", "DF", None, None) is None
    monkeypatch.setattr(
        criticidades, "get_falecidos_data", lambda *args: (_ for _ in ()).throw(RuntimeError("offline"))
    )
    assert criticidades._build_falecidos_context("12345678000195", "DF", None, None) is None


def test_incompatibilidade_context_normalizes_service_payload_and_handles_unavailable(monkeypatch):
    row = {
        "pct_clinico": 12, "clinico_valor_suspeito": 20, "med_clinico_reg": 2, "med_clinico_uf": 3,
        "med_clinico_br": 4, "risco_clinico_reg": 6, "risco_clinico_uf": 4, "risco_clinico_br": 3,
        "id_cnpj": 7,
    }
    monkeypatch.setattr(criticidades, "_get_matriz_row_context", lambda *args: row)
    service_result = SimpleNamespace(model_dump=lambda: {
        "patologias": [{
            "patologia": "Doença de Parkinson", "regra_clinica": "IDADE_MENOR_50",
            "municipal_resumo": [{"grupo": "total"}], "ranking_municipal": [{"cnpj": "123"}],
            "demografia_parkinson": {
                "populacao_total": 1000, "populacao_50_mais": 500, "casos_esperados": 4.3,
                "cpfs_observados": 8, "razao_observado_esperado": None, "prevalencia_50_mais": 0.0086,
            },
        }]
    })
    monkeypatch.setattr(criticidades, "get_incompatibilidade_patologica_data", lambda *args: service_result)

    result = criticidades._build_incompatibilidade_patologica_context(
        "12345678000195", date(2024, 1, 1), date(2024, 12, 31)
    )

    assert result["periodo_desc"] == "no ano de 2024"
    assert result["ranking_patologias"][0]["municipal"]["top20"][0]["cnpj"] == "123"
    assert result["ranking_patologias"][0]["demografia_parkinson"]["qtd_cpfs_distintos_observado"] == 8
    assert result["ranking_patologias"][0]["demografia_parkinson"]["razao_observado_esperado"] == 8 / 4.3

    monkeypatch.setattr(
        criticidades, "get_incompatibilidade_patologica_data",
        lambda *args: (_ for _ in ()).throw(HTTPException(status_code=404, detail="sem dados")),
    )
    unavailable = criticidades._build_incompatibilidade_patologica_context("12345678000195", None, None)
    assert unavailable["unavailable"] is True and unavailable["motivo"] == "sem dados"
    monkeypatch.setattr(
        criticidades, "get_incompatibilidade_patologica_data",
        lambda *args: (_ for _ in ()).throw(HTTPException(status_code=503, detail="offline")),
    )
    with pytest.raises(RuntimeError, match="Detalhamento clinico indisponivel"):
        criticidades._build_incompatibilidade_patologica_context("12345678000195", None, None)


def test_clinical_municipal_comparison_aggregates_pharmacies_and_validates_profiles():
    item = {"patologia": "Doença de Parkinson", "regra_clinica": "IDADE_MENOR_50"}
    clinic = pl.DataFrame(
        {
            "id_cnpj": [1, 2], "patologia": ["DOENCA DE PARKINSON"] * 2,
            "regra_clinica": ["IDADE_MENOR_50"] * 2,
            "qtd_cpfs_distintos": [10, 20], "qtd_cpfs_incompativeis": [2, 5],
            "qtd_autorizacoes_incompativeis": [3, 7], "valor_incompativel_pago": [30.0, 70.0],
        }
    )
    profiles = pl.DataFrame(
        {
            "id_cnpj": [1, 2], "cnpj": ["12345678000195", "99999999000191"],
            "razao_social": ["Farmácia Alvo", "Outra Farmácia"],
        }
    )

    result = criticidades._build_clinica_municipal_context(item, clinic, profiles, 1)

    assert result["qtd_farmacias_municipio"] == 2
    assert result["qtd_cpfs_incompativeis_municipio"] == 7
    assert result["top20"][0]["is_alvo"] is False
    assert result["top20"][0]["participacao_municipal"] == 0.7
    assert result["resumo"][1]["qtd_cpfs_incompativeis"] == 5.0
    with pytest.raises(RuntimeError, match="sem linhas para o recorte"):
        criticidades._build_clinica_municipal_context(item, clinic.head(0), profiles, 1)
    with pytest.raises(RuntimeError, match="sem CNPJ alvo"):
        criticidades._build_clinica_municipal_context(item, clinic, profiles, 99)
    with pytest.raises(RuntimeError, match="Cache de perfil dos estabelecimentos sem colunas"):
        criticidades._build_clinica_municipal_context(item, clinic, pl.DataFrame({"id_cnpj": [1]}), 1)


def test_teto_context_requires_finite_nonnegative_financial_fields(monkeypatch):
    row = {
        "pct_teto": 15, "teto_valor": 20, "teto_valor_total": 30, "valor_total_vendas": 50,
        "med_teto_reg": 5, "med_teto_uf": 6, "med_teto_br": 7,
        "risco_teto_reg": 3, "risco_teto_uf": 2, "risco_teto_br": 1,
    }
    monkeypatch.setattr(criticidades, "_get_matriz_row_context", lambda *args: row)
    annual = criticidades._build_teto_context(
        "12345678000195", date(2023, 1, 1), date(2023, 12, 31)
    )
    assert annual["periodo_desc"] == "no ano de 2023" and annual["valor_suspeito"] == 20
    monkeypatch.setattr(criticidades, "_get_matriz_row_context", lambda *args: {**row, "teto_valor": float("inf")})
    with pytest.raises(RuntimeError, match="teto_valor invalido"):
        criticidades._build_teto_context("12345678000195", None, None)


def test_matrix_context_builders_cover_all_standard_indicators(monkeypatch):
    matrix_values = {
        "pct_polimedicamento": 12.5, "polimedicamento_valor": 125.0,
        "med_polimedicamento_reg": 5.0, "med_polimedicamento_uf": 4.0, "med_polimedicamento_br": 3.0,
        "risco_polimedicamento_reg": 2.5, "risco_polimedicamento_uf": 2.0, "risco_polimedicamento_br": 1.5,
        "val_ticket_medio": 20.0, "med_ticket_reg": 10.0, "med_ticket_uf": 9.0, "med_ticket_br": 8.0,
        "risco_ticket_reg": 2.0, "risco_ticket_uf": 1.8, "risco_ticket_br": 1.6,
        "val_receita_paciente": 200.0, "med_receita_paciente_reg": 100.0, "med_receita_paciente_uf": 90.0, "med_receita_paciente_br": 80.0,
        "risco_receita_paciente_reg": 2.0, "risco_receita_paciente_uf": 1.8, "risco_receita_paciente_br": 1.6,
        "val_per_capita": 30.0, "med_per_capita_reg": 15.0, "med_per_capita_uf": 12.0, "med_per_capita_br": 10.0,
        "risco_per_capita_reg": 2.0, "risco_per_capita_uf": 1.8, "risco_per_capita_br": 1.6,
        "pct_alto_custo": 10.0, "alto_custo_valor": 100.0, "med_alto_custo_reg": 5.0, "med_alto_custo_uf": 4.0,
        "med_alto_custo_br": 3.0, "risco_alto_custo_reg": 2.0, "risco_alto_custo_uf": 1.8, "risco_alto_custo_br": 1.6,
        "pct_vendas_rapidas": 20.0, "vendas_rapidas_valor": 200.0, "med_vendas_rapidas_reg": 10.0,
        "med_vendas_rapidas_uf": 8.0, "med_vendas_rapidas_br": 6.0, "risco_vendas_rapidas_reg": 2.0,
        "risco_vendas_rapidas_uf": 1.8, "risco_vendas_rapidas_br": 1.6,
        "pct_recorrencia_sistemica": 15.0, "recorrencia_valor_sistemico": 150.0,
        "med_recorrencia_sistemica_reg": 7.5, "med_recorrencia_sistemica_uf": 6.0,
        "med_recorrencia_sistemica_br": 5.0, "risco_recorrencia_sistemica_reg": 2.0,
        "risco_recorrencia_sistemica_uf": 1.8, "risco_recorrencia_sistemica_br": 1.6,
        "pct_pico": 25.0, "pico_valor_top3_dias": 250.0, "med_pico_reg": 12.5, "med_pico_uf": 10.0,
        "med_pico_br": 8.0, "risco_pico_reg": 2.0, "risco_pico_uf": 1.8, "risco_pico_br": 1.6,
        "pct_teto": 25.0, "teto_valor": 250.0, "teto_valor_total": 1000.0, "valor_total_vendas": 1200.0,
        "med_teto_reg": 12.5, "med_teto_uf": 10.0, "med_teto_br": 8.0,
        "risco_teto_reg": 2.0, "risco_teto_uf": 1.8, "risco_teto_br": 1.6,
    }
    monkeypatch.setattr(criticidades, "_get_matriz_row_context", lambda *_args: matrix_values)
    inicio, fim = date(2020, 1, 1), date(2024, 12, 31)
    builders = [
        (criticidades._build_teto_context, "valor_suspeito", 250.0),
        (criticidades._build_polimedicamento_context, "percentual", 12.5),
        (criticidades._build_ticket_medio_context, "valor", 20.0),
        (criticidades._build_receita_paciente_context, "valor", 200.0),
        (criticidades._build_per_capita_context, "valor", 30.0),
        (criticidades._build_alto_custo_context, "valor_suspeito", 100.0),
        (criticidades._build_vendas_rapidas_context, "percentual", 20.0),
        (criticidades._build_recorrencia_sistemica_context, "percentual", 15.0),
        (criticidades._build_dias_pico_context, "percentual", 25.0),
    ]
    for builder, field, expected in builders:
        result = builder("12345678000195", inicio, fim)
        assert result is not None
        assert result["periodo_desc"] == "no período de 2020 a 2024"
        assert result[field] == expected


def test_criticality_text_and_summary_tables_render_all_standard_indicators():
    doc = Document()
    common = {
        "periodo_desc": "no período de 2020 a 2024", "percentual": 25.0,
        "valor_suspeito": 100.0, "valor": 200.0,
        "valor_monitorado": 500.0, "valor_total_vendas": 700.0,
        "multiplicador_regiao": 2.0, "multiplicador_uf": 1.5, "multiplicador_brasil": 1.2,
    }
    criticidades._add_falecidos_criticidade_text(
        doc, "7.1", "Farmácia Exemplo",
        {"total_autorizacoes": 1234, "cpfs_distintos": 12, "valor_total": 987.65, "periodo_desc": "no período analisado"},
        bookmark_name="critico_falecidos",
    )
    criticidades._add_teto_text(doc, "7.2", "Farmácia Exemplo", common, bookmark_name="critico_teto")
    criticidades._add_polimedicamento_text(doc, "7.3", "Farmácia Exemplo", common, bookmark_name="critico_poli")
    criticidades._add_ticket_medio_text(doc, "7.4", "Farmácia Exemplo", common, bookmark_name="critico_ticket")
    criticidades._add_receita_paciente_text(doc, "7.5", "Farmácia Exemplo", common, bookmark_name="critico_receita")
    criticidades._add_per_capita_text(doc, "7.6", "Farmácia Exemplo", common, bookmark_name="critico_capita")
    criticidades._add_alto_custo_text(doc, "7.7", "Farmácia Exemplo", common, bookmark_name="critico_custo")
    criticidades._add_vendas_rapidas_text(doc, "7.8", "Farmácia Exemplo", common, bookmark_name="critico_rapidas")
    criticidades._add_recorrencia_sistemica_text(doc, "7.9", "Farmácia Exemplo", common, bookmark_name="critico_recorrencia")
    criticidades._add_dias_pico_text(doc, "7.10", "Farmácia Exemplo", common, bookmark_name="critico_pico")
    criticidades._add_indicador_regional_table(
        doc,
        {"indicador_key": "teto", "indicador_label": "Teto", "formato": "pct", "id_regiao_saude": "R-1",
         "valor": "25,00%", "mediana_regional": "12,50%", "risco_regional": 2.0,
         "posicao_regional": 2, "total_regional": 20, "percentil_regional": 94.7},
        20,
    )
    criticidades._add_indicador_regional_table(
        doc,
        {"indicador_key": "dispersao_geografica", "indicador_label": "Dispersão", "formato": "pct",
         "id_regiao_saude": "R-1", "valor": "12,00%", "mediana_regional": "6,00%",
         "risco_regional": 2.0, "posicao_regional": 1, "total_regional": 10, "percentil_regional": 100.0},
        21,
    )
    criticidades._add_indicadores_criticos_quadro(
        doc,
        [
            {"indicador": "Teto", "valor": "25,00%", "mediana_regional": "12,50%", "risco_regional": 2.0, "status": "CRÍTICO", "bookmark": "critico_teto"},
            {"indicador": "Ticket médio", "valor": "R$ 200,00", "mediana_regional": "R$ 100,00", "risco_regional": None, "status": "CRÍTICO"},
        ],
        22,
    )
    criticidades._add_indicadores_criticos_quadro(doc, [], 23)
    text = " ".join(p.text for p in doc.paragraphs) + " " + " ".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    assert "R$ 987,65" in text
    assert "R$ 100,00" in text
    assert "Farmácias Região" in text
    assert "CRÍTICO" in text


def test_clinical_detail_tables_and_parkinson_narrative_cover_branches():
    doc = Document()
    item = {
        "objeto": "Doença de Parkinson", "titulo": "Parkinson menor de 50 anos", "criterio": "IDADE_MENOR_50",
        "ano_inicio": 2020, "ano_fim": 2024,
        "qtd_cpfs_distintos": 40, "qtd_cpfs_incompativeis": 10,
        "qtd_autorizacoes_incompativeis": 12, "valor_incompativel_pago": 1500.0,
        "maior_percentil_regional_qtd_cpfs_incompativeis": 0.95,
        "maior_participacao_cpfs_incompativeis_regiao": 0.5,
        "melhor_rank_regional_qtd_cpfs_incompativeis": 1,
        "evolucao_anual": [
            {"ano_base": 2024, "qtd_cpfs_distintos": 40, "qtd_cpfs_incompativeis": 10,
             "percentual_cpfs_incompativeis": None, "qtd_autorizacoes": 100,
             "qtd_autorizacoes_incompativeis": 12, "valor_incompativel_pago": 1500.0}
        ],
        "municipal": {
            "resumo": [
                {"grupo": "Farmácia analisada", "qtd_farmacias": 1, "valor_incompativel_pago": 1500.0},
                {"grupo": "Demais farmácias", "qtd_farmacias": 4, "valor_incompativel_pago": 500.0},
                {"grupo": "Total municipal", "qtd_farmacias": 5, "valor_incompativel_pago": 2000.0},
            ],
            "top20": [{"posicao": 1, "cnpj": "12345678000195", "razao_social": "Farmácia Exemplo",
                       "valor_incompativel_pago": 1500.0, "participacao_municipal": 0.75}],
        },
        "demografia_parkinson": {
            "municipio": "Campinas", "uf": "SP", "populacao_total": 100000,
            "populacao_50_mais": 30000, "percentual_50_mais": 0.3,
            "prevalencia_referencia": 0.03, "casos_esperados": 900.0,
            "ano_observado": 2024, "qtd_cpfs_distintos_observado": 1200,
            "razao_observado_esperado": 1.33, "percentual_superior": 33.0,
            "faixas_etarias": [
                {"faixa": "0-49", "populacao": 70000, "destacar_50_mais": False},
                {"faixa": "50+", "populacao": 30000, "destacar_50_mais": True},
            ],
        },
    }
    criticidades._add_clinica_evolucao_anual_table(doc, item, 1)
    criticidades._add_clinica_municipio_resumo_table(doc, item, 2)
    criticidades._add_clinica_municipio_top20_table(doc, item, 3)
    criticidades._add_parkinson_demografia_table(doc, item["demografia_parkinson"], 4)
    criticidades._add_parkinson_demografia_text(doc, item, 5)
    criticidades._add_parkinson_demografia_text(
        doc,
        {"demografia_parkinson": {**item["demografia_parkinson"], "percentual_superior": -10.0}},
        6,
    )
    criticidades._add_parkinson_gtin_sem_comprovacao_text(
        doc,
        {"ranking_patologias": [{"patologia": "Doença de Parkinson", "regra_clinica": "IDADE_MENOR_50"}]},
        {"rows": [
            {"patologia": "Doença de Parkinson", "descricao": "Medicamento A", "gtin": "123",
             "valor_vendas": 1000.0, "valor_sem_comprovacao": 250.0},
            {"patologia": "Doença de Parkinson", "descricao": "Medicamento B", "gtin": "456",
             "valor_vendas": 100.0, "valor_sem_comprovacao": 10.0},
        ]},
        7,
    )
    assert len(doc.tables) == 6
    all_text = " ".join(p.text for p in doc.paragraphs) + " " + " ".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    assert "abaixo dessa estimativa" in all_text
    assert "Medicamento A" in all_text


def test_geographic_detail_table_rolls_up_remaining_states_and_regional_text():
    doc = Document()
    origin_rows = [
        {"uf_paciente": f"UF{i}", "is_outra_uf": i % 2 == 0, "qtd_autorizacoes": i + 1,
         "valor_autorizado": 100.0 + i, "percentual_sobre_total": 2.0 + i,
         "percentual_sobre_outra_uf": (5.0 + i) if i % 2 == 0 else None}
        for i in range(14)
    ]
    geographic = {
        "periodo_desc": "no período analisado", "percentual": 25.0, "percentual_financeiro_outra_uf": 18.5,
        "total_autorizacoes_outra_uf": 35, "total_valor_outra_uf": 2500.0,
        "total_valor_origem": 10000.0, "total_autorizacoes_origem": 100,
        "multiplicador_regiao": 2.0, "multiplicador_uf": 1.5, "multiplicador_brasil": 1.2,
        "origem_uf_rows": origin_rows,
    }
    criticidades._add_dispersao_geografica_origem_uf_table(doc, "Farmácia Exemplo", geographic, 10)
    criticidades._add_dispersao_geografica_regional_text(doc, "Farmácia Exemplo", geographic)
    assert "Demais UFs" in " ".join(cell.text for row in doc.tables[0].rows for cell in row.cells)
    assert "2,00 vezes" in " ".join(p.text for p in doc.paragraphs)


def test_geographic_context_builds_period_scoped_rows_and_fails_on_bad_sources(monkeypatch):
    cnpj = "12345678000195"
    matrix_row = {
        "pct_geografico": 25.0, "med_geografico_reg": 10.0, "med_geografico_uf": 8.0,
        "med_geografico_br": 5.0, "risco_geografico_reg": 2.5, "risco_geografico_uf": 3.0,
        "risco_geografico_br": 5.0,
    }
    monkeypatch.setattr(criticidades, "_get_matriz_row_context", lambda *_args: matrix_row)
    farmacias = pl.DataFrame({"id_cnpj": [7], "cnpj": ["12.345.678/0001-95"], "uf": ["sp"]})
    origem = pl.DataFrame({
        "id_cnpj": [7, 7, 7], "ano_base": [2022, 2023, 2024],
        "uf_farmacia": ["sp", "SP", "SP"], "uf_paciente": ["SP", "RJ", "MG"],
        "is_outra_uf": [False, True, True], "qtd_autorizacoes": [10, 20, 30],
        "valor_autorizado": [100.0, 200.0, 300.0],
    })
    monkeypatch.setattr(criticidades, "get_df_dados_farmacia", lambda: farmacias)
    monkeypatch.setattr(criticidades, "scan_geografico_origem_uf", lambda: origem.lazy())

    result = criticidades._build_dispersao_geografica_context(
        cnpj, date(2023, 1, 1), date(2024, 12, 31)
    )
    assert result["periodo_desc"] == "no período de 2023 a 2024"
    assert result["id_cnpj"] == 7 and result["uf_farmacia"] == "SP"
    assert result["total_valor_origem"] == 500
    assert result["total_valor_outra_uf"] == 500
    assert result["total_autorizacoes_outra_uf"] == 50
    assert result["percentual_financeiro_outra_uf"] == 100
    assert [row["uf_paciente"] for row in result["origem_uf_rows"]] == ["MG", "RJ"]
    assert result["origem_uf_rows"][0]["percentual_sobre_outra_uf"] == 60

    same_year = criticidades._build_dispersao_geografica_context(cnpj, date(2024, 1, 1), date(2024, 12, 31))
    assert same_year["periodo_desc"] == "no ano de 2024"
    open_period = criticidades._build_dispersao_geografica_context(cnpj, None, None)
    assert open_period["periodo_desc"] == "no período analisado"
    monkeypatch.setattr(criticidades, "_get_matriz_row_context", lambda *_args: None)
    assert criticidades._build_dispersao_geografica_context(cnpj, None, None) is None
    monkeypatch.setattr(criticidades, "_get_matriz_row_context", lambda *_args: matrix_row)

    monkeypatch.setattr(criticidades, "get_df_dados_farmacia", lambda: pl.DataFrame({"cnpj": [cnpj]}))
    with pytest.raises(RuntimeError, match="Cache de farmacias sem colunas"):
        criticidades._build_dispersao_geografica_context(cnpj, None, None)
    monkeypatch.setattr(criticidades, "get_df_dados_farmacia", lambda: farmacias.head(0))
    with pytest.raises(RuntimeError, match="sem id_cnpj"):
        criticidades._build_dispersao_geografica_context(cnpj, None, None)
    monkeypatch.setattr(criticidades, "get_df_dados_farmacia", lambda: farmacias)
    monkeypatch.setattr(
        criticidades,
        "scan_geografico_origem_uf",
        lambda: pl.DataFrame({"id_cnpj": [7], "bad": [1]}).lazy(),
    )
    with pytest.raises(RuntimeError, match="Cache geografico de origem por UF sem colunas"):
        criticidades._build_dispersao_geografica_context(cnpj, None, None)
    monkeypatch.setattr(criticidades, "scan_geografico_origem_uf", lambda: origem.lazy())
    with pytest.raises(RuntimeError, match="sem registros"):
        criticidades._build_dispersao_geografica_context(cnpj, date(2025, 1, 1), None)
    zero_value = origem.with_columns(pl.lit(0.0).alias("valor_autorizado"))
    monkeypatch.setattr(criticidades, "scan_geografico_origem_uf", lambda: zero_value.lazy())
    with pytest.raises(RuntimeError, match="sem valor autorizado positivo"):
        criticidades._build_dispersao_geografica_context(cnpj, None, None)


def test_geographic_and_clinical_narratives_render_complete_branches(monkeypatch):
    geo_doc = Document()
    map_calls = []
    monkeypatch.setattr(
        criticidades,
        "_add_mapa_geografico_origem_uf",
        lambda doc, name, context: map_calls.append((name, context["id_cnpj"])),
    )
    geographic = {
        "periodo_desc": "no ano de 2024", "percentual_financeiro_outra_uf": 18.5,
        "total_autorizacoes_outra_uf": 35, "total_valor_outra_uf": 2500.0,
        "total_valor_origem": 10000.0, "total_valor_outra_uf": 2500.0,
        "multiplicador_regiao": 2.0, "multiplicador_uf": 1.5, "multiplicador_brasil": 1.2,
        "percentual": 25.0, "id_cnpj": 7,
        "origem_uf_rows": [{
            "uf_paciente": "RJ", "is_outra_uf": True, "qtd_autorizacoes": 35,
            "valor_autorizado": 2500.0, "percentual_sobre_total": 25.0,
            "percentual_sobre_outra_uf": 100.0,
        }],
    }
    criticidades._add_dispersao_geografica_text(geo_doc, "7.3", "Farmácia Teste", geographic, 12, "geo")
    assert map_calls == [("Farmácia Teste", 7)]
    assert len(geo_doc.tables) == 1
    assert "R$ 2.500,00" in " ".join(p.text for p in geo_doc.paragraphs)

    clinical_doc = Document()
    diabetes = {
        "patologia": "DIABETES", "regra_clinica": "IDADE_MENOR_20",
        "objeto": "diabetes", "titulo": "Diabetes", "criterio": "IDADE_MENOR_20",
        "ano_inicio": 2023, "ano_fim": 2024, "qtd_cpfs_distintos": 20,
        "qtd_cpfs_incompativeis": 4, "qtd_autorizacoes_incompativeis": 6,
        "valor_incompativel_pago": 600.0,
        "maior_percentil_regional_qtd_cpfs_incompativeis": 0.9,
        "maior_participacao_cpfs_incompativeis_regiao": 0.2,
        "melhor_rank_regional_qtd_cpfs_incompativeis": 3,
        "evolucao_anual": [{
            "ano_base": 2024, "qtd_cpfs_distintos": 20, "qtd_cpfs_incompativeis": 4,
            "percentual_cpfs_incompativeis": 0.2, "qtd_autorizacoes": 100,
            "qtd_autorizacoes_incompativeis": 6, "valor_incompativel_pago": 600.0,
        }],
        "municipal": {
            "resumo": [
                {"grupo": "Farmácia analisada", "qtd_farmacias": 1, "valor_incompativel_pago": 600.0},
                {"grupo": "Demais farmácias do município", "qtd_farmacias": 4, "valor_incompativel_pago": 400.0},
                {"grupo": "Total do município", "qtd_farmacias": 5, "valor_incompativel_pago": 1000.0},
            ],
            "top20": [{"posicao": 1, "cnpj": "12345678000195", "razao_social": "Farmácia Teste",
                       "valor_incompativel_pago": 600.0, "participacao_municipal": 1.0}],
        },
    }
    parkinson = {
        **diabetes,
        "patologia": "Doença de Parkinson", "regra_clinica": "IDADE_MENOR_50",
        "objeto": "Parkinson", "titulo": "Doença de Parkinson", "criterio": "IDADE_MENOR_50",
        "qtd_cpfs_distintos": 30, "qtd_cpfs_incompativeis": 5,
    }
    demo_calls = []
    monkeypatch.setattr(
        criticidades,
        "_add_parkinson_demografia_text",
        lambda doc, item, number: demo_calls.append((item["titulo"], number)),
    )
    criticidades._add_incompatibilidade_patologica_text(
        clinical_doc,
        "7.2",
        "Farmácia Teste",
        {
            "periodo_desc": "no período de 2023 a 2024", "percentual": 12.5,
            "valor_suspeito": 0, "multiplicador_regiao": 1.0, "multiplicador_uf": 1.5,
            "multiplicador_brasil": 2.0, "ranking_patologias": [diabetes, parkinson],
        },
        20,
        "clinico",
    )
    assert demo_calls == [("Doença de Parkinson", 24)]
    assert len(clinical_doc.tables) == 6
    text = " ".join(p.text for p in clinical_doc.paragraphs)
    assert "R$ 0,00" not in text
    assert "20 CPFs distintos" in text
    assert "Parkinson" in text

    empty_doc = Document()
    criticidades._add_incompatibilidade_patologica_text(
        empty_doc,
        "7.2",
        "Farmácia Teste",
        {"periodo_desc": "no período analisado", "percentual": 10,
         "valor_suspeito": 0, "multiplicador_regiao": 1, "multiplicador_uf": 1,
         "multiplicador_brasil": 1, "ranking_patologias": []},
        1,
    )
    assert len(empty_doc.tables) == 0


def test_clinical_municipal_profile_missing_for_top_ranked_pharmacy_is_rejected():
    item = {"patologia": "Diabetes", "regra_clinica": "IDADE_MENOR_20"}
    clinic = pl.DataFrame(
        {
            "id_cnpj": [1, 2], "patologia": ["DIABETES", "DIABETES"],
            "regra_clinica": ["IDADE_MENOR_20", "IDADE_MENOR_20"],
            "qtd_cpfs_distintos": [10, 20], "qtd_cpfs_incompativeis": [2, 5],
            "qtd_autorizacoes_incompativeis": [3, 7], "valor_incompativel_pago": [30.0, 70.0],
        }
    )
    profiles = pl.DataFrame(
        {"id_cnpj": [1], "cnpj": ["12345678000195"], "razao_social": ["Farmácia alvo"]}
    )

    with pytest.raises(RuntimeError, match="Perfil de estabelecimento ausente para Top 10"):
        criticidades._build_clinica_municipal_context(item, clinic, profiles, 1)


@pytest.mark.parametrize(
    ("evolution", "demography", "message"),
    [
        ([{"qtd_cpfs_distintos": 0, "ano_base": 2022}], None, "Quantidade de CPFs distintos"),
        ([{"qtd_cpfs_distintos": 1, "ano_base": 2022}], pl.DataFrame({
            "id_ibge7": ["1234567"], "ano_censo": [2021], "idade_min": [50], "nu_populacao": [10],
        }), "Demografia IBGE"),
        ([{"qtd_cpfs_distintos": 1, "ano_base": 2022}], pl.DataFrame({
            "id_ibge7": ["1234567"], "ano_censo": [2022], "idade_min": [50], "nu_populacao": [0],
        }), "populacao total"),
        ([{"qtd_cpfs_distintos": 1, "ano_base": 2022}], pl.DataFrame({
            "id_ibge7": ["1234567"], "ano_censo": [2022], "idade_min": [40], "nu_populacao": [10],
        }), r"populacao 50\+"),
        ([{"qtd_cpfs_distintos": 1, "ano_base": 2022}], pl.DataFrame({
            "id_ibge7": ["1234567"] * 2, "ano_censo": [2022] * 2,
            "idade_min": [-1, 50], "nu_populacao": [10, 10],
        }), "idade/populacao invalida"),
        ([{"qtd_cpfs_distintos": 1, "ano_base": 2022}], pl.DataFrame({
            "id_ibge7": ["1234567"] * 3, "ano_censo": [2022] * 3,
            "idade_min": [0, 50, 60], "nu_populacao": [200, -1, 2],
        }), "idade/populacao invalida"),
    ],
)
def test_parkinson_demography_rejects_invalid_observations_and_census_values(
    monkeypatch, evolution, demography, message
):
    if demography is not None:
        monkeypatch.setattr(criticidades, "get_df_dados_ibge_demografia", lambda: demography)
    pharmacy = {"id_ibge7": "1234567", "municipio": "Campinas", "uf": "SP"}

    with pytest.raises(RuntimeError, match=message):
        criticidades._build_parkinson_demografia_context(pharmacy, evolution)


def test_critical_indicator_summary_validates_configuration_and_matrix_values(monkeypatch):
    all_flags = dict(criticidades._INDICATOR_FLAGS)
    all_metadata = dict(criticidades._INDICADOR_QUADRO_META)
    row = _matrix_row_for_indicators(critical_keys={"teto"})
    monkeypatch.setattr(criticidades, "_get_matriz_dinamica_nota", lambda *_args: (pl.DataFrame([row]), row))
    monkeypatch.setattr(criticidades, "_get_criticos_ordenados_por_risco", lambda *_args: [])

    monkeypatch.setattr(criticidades, "_INDICATOR_FLAGS", {"invalid": ("attention", "critical")})
    with pytest.raises(RuntimeError, match="sem mapeamento de colunas"):
        criticidades._build_indicadores_criticos_quadro("12345678000195")

    monkeypatch.setattr(criticidades, "_INDICATOR_FLAGS", {"teto": all_flags["teto"]})
    metadata = dict(all_metadata)
    metadata.pop("teto")
    monkeypatch.setattr(criticidades, "_INDICADOR_QUADRO_META", metadata)
    with pytest.raises(RuntimeError, match="sem metadados de quadro"):
        criticidades._build_indicadores_criticos_quadro("12345678000195")

    monkeypatch.setattr(criticidades, "_INDICATOR_FLAGS", all_flags)
    monkeypatch.setattr(criticidades, "_INDICADOR_QUADRO_META", all_metadata)
    for column, message in (
        ("pct_teto", "sem valor na matriz"),
        ("med_teto_reg", "sem mediana regional"),
        ("risco_teto_reg", "sem valor, mediana regional ou risco regional"),
    ):
        broken = {**row, column: None}
        monkeypatch.setattr(criticidades, "_get_matriz_dinamica_nota", lambda *_args, broken=broken: (pl.DataFrame([row]), broken))
        with pytest.raises(RuntimeError, match=message):
            criticidades._build_indicadores_criticos_quadro("12345678000195")

    volume_col = criticidades._VOLUME_ATIPICO_VALOR_AUMENTO_COL
    volume_row = _matrix_row_for_indicators(critical_keys={"volume_atipico"})
    volume_row[volume_col] = None
    monkeypatch.setattr(criticidades, "_INDICATOR_FLAGS", all_flags)
    monkeypatch.setattr(criticidades, "_INDICADOR_QUADRO_META", all_metadata)
    monkeypatch.setattr(criticidades, "_get_matriz_dinamica_nota", lambda *_args: (pl.DataFrame([volume_row]), volume_row))
    with pytest.raises(RuntimeError, match="volume_atipico sem valor financeiro"):
        criticidades._build_indicadores_criticos_quadro("12345678000195")


def test_critical_indicator_summary_skips_noncritical_rows_and_allows_zero_baseline(monkeypatch):
    row = _matrix_row_for_indicators(critical_keys=set())
    monkeypatch.setattr(criticidades, "_get_matriz_dinamica_nota", lambda *_args: (pl.DataFrame([row]), row))
    monkeypatch.setattr(criticidades, "_get_criticos_ordenados_por_risco", lambda *_args: [])
    assert criticidades._build_indicadores_criticos_quadro("12345678000195") == []

    falecidos = _matrix_row_for_indicators(critical_keys={"falecidos"})
    value_col, median_col, _, _, risk_col, _, _ = criticidades.INDICATOR_MAPPING["falecidos"]
    falecidos[median_col] = None
    falecidos[risk_col] = None
    monkeypatch.setattr(criticidades, "_get_matriz_dinamica_nota", lambda *_args: (pl.DataFrame([falecidos]), falecidos))
    monkeypatch.setattr(criticidades, "_get_criticos_ordenados_por_risco", lambda *_args: ["falecidos"])
    result = criticidades._build_indicadores_criticos_quadro("12345678000195")
    assert result[0]["key"] == "falecidos"
    formato = criticidades._INDICADOR_QUADRO_META["falecidos"][1]
    assert result[0]["valor"] == criticidades._format_indicador_quadro_value(falecidos[value_col], formato)


def test_regional_indicator_context_rejects_bad_scopes_and_missing_comparators(monkeypatch):
    key = "ticket_medio"
    value_col, median_col, _, _, risk_col, _, _ = criticidades.INDICATOR_MAPPING[key]

    def install(frame):
        monkeypatch.setattr(criticidades, "_get_matriz_dinamica_nota", lambda *_args: (frame, {}))

    def frame(cnpjs=("12345678000195",), regions=("R1",), values=(10.0,), medians=(5.0,), risks=(2.0,)):
        return pl.DataFrame({
            "cnpj": list(cnpjs), "id_regiao_saude": list(regions), value_col: list(values),
            median_col: list(medians), risk_col: list(risks),
        })

    metadata = dict(criticidades._INDICADOR_QUADRO_META)
    metadata.pop(key)
    monkeypatch.setattr(criticidades, "_INDICADOR_QUADRO_META", metadata)
    with pytest.raises(RuntimeError, match="sem metadados"):
        criticidades._build_indicador_regional_context("12345678000195", key)
    monkeypatch.setattr(criticidades, "_INDICADOR_QUADRO_META", dict(metadata, **{key: ("Ticket médio", "val")}))

    install(pl.DataFrame({"cnpj": ["12345678000195"]}))
    with pytest.raises(RuntimeError, match="sem colunas obrigatorias"):
        criticidades._build_indicador_regional_context("12345678000195", key)
    install(frame(cnpjs=("99999999000191",)))
    with pytest.raises(RuntimeError, match="nao encontrado"):
        criticidades._build_indicador_regional_context("12345678000195", key)
    install(frame(cnpjs=("12345678000195", "12345678000195"), regions=("R1", "R1"), values=(10.0, 9.0), medians=(5.0, 4.0), risks=(2.0, 2.0)))
    with pytest.raises(RuntimeError, match="mais de uma linha"):
        criticidades._build_indicador_regional_context("12345678000195", key)
    install(frame(regions=("",)))
    with pytest.raises(RuntimeError, match="sem id_regiao_saude"):
        criticidades._build_indicador_regional_context("12345678000195", key)
    install(frame(values=(None,)))
    with pytest.raises(RuntimeError, match="sem linhas para indicador"):
        criticidades._build_indicador_regional_context("12345678000195", key)
    install(frame(
        cnpjs=("12345678000195", "99999999000191"),
        regions=("R1", "R1"),
        values=(None, 20.0),
        medians=(5.0, 10.0),
        risks=(2.0, 4.0),
    ))
    with pytest.raises(RuntimeError, match="sem valor/risco calculado"):
        criticidades._build_indicador_regional_context("12345678000195", key)
    install(frame(medians=(None,)))
    with pytest.raises(RuntimeError, match="sem mediana regional"):
        criticidades._build_indicador_regional_context("12345678000195", key)
    install(frame(risks=(None,)))
    with pytest.raises(RuntimeError, match="sem risco regional"):
        criticidades._build_indicador_regional_context("12345678000195", key)
    install(frame())
    result = criticidades._build_indicador_regional_context("12345678000195", key)
    assert result["total_regional"] == 1 and result["percentil_regional"] == 100.0


def test_criticality_tables_and_parkinson_cross_reference_reject_or_skip_invalid_data():
    with pytest.raises(RuntimeError, match="Contexto regional de indicador ausente"):
        criticidades._add_indicador_regional_table(Document(), {}, 1)
    with pytest.raises(RuntimeError, match="Evolucao anual clinica obrigatoria"):
        criticidades._add_clinica_evolucao_anual_table(Document(), {}, 1)
    with pytest.raises(RuntimeError, match="Comparativo municipal clinico obrigatorio"):
        criticidades._add_clinica_municipio_resumo_table(Document(), {}, 1)
    with pytest.raises(RuntimeError, match="Resumo municipal clinico deve conter"):
        criticidades._add_clinica_municipio_resumo_table(
            Document(), {"municipal": {"resumo": [{"valor_incompativel_pago": 1}]}}, 1
        )
    with pytest.raises(RuntimeError, match="Ranking municipal clinico obrigatorio"):
        criticidades._add_clinica_municipio_top20_table(Document(), {}, 1)
    with pytest.raises(RuntimeError, match="Top 10 municipal clinico sem farmacias"):
        criticidades._add_clinica_municipio_top20_table(Document(), {"municipal": {"top20": []}}, 1)
    with pytest.raises(RuntimeError, match="Comparacao demografica de Parkinson obrigatoria"):
        criticidades._add_parkinson_demografia_text(Document(), {}, 1)
    with pytest.raises(RuntimeError, match="Distribuicao geografica por UF obrigatoria"):
        criticidades._add_dispersao_geografica_origem_uf_table(Document(), "Farmácia", {}, 1)

    doc = Document()
    criticidades._add_parkinson_gtin_sem_comprovacao_text(
        doc, {"ranking_patologias": []}, {"rows": []}, 1
    )
    parkinson = {"ranking_patologias": [{"patologia": "Doença de Parkinson", "regra_clinica": "IDADE_MENOR_50"}]}
    criticidades._add_parkinson_gtin_sem_comprovacao_text(doc, parkinson, {"rows": []}, 2)
    criticidades._add_parkinson_gtin_sem_comprovacao_text(
        doc, parkinson,
        {"rows": [{"patologia": "Doença de Parkinson", "valor_vendas": 0, "valor_sem_comprovacao": 0}]},
        3,
    )
    assert doc.paragraphs == []


def test_clinical_context_conversion_and_narrative_cover_error_and_financial_branches(monkeypatch):
    matrix_row = {key: "inválido" for key in (
        "pct_clinico", "clinico_valor_suspeito", "med_clinico_reg", "med_clinico_uf",
        "med_clinico_br", "risco_clinico_reg", "risco_clinico_uf", "risco_clinico_br",
    )}
    monkeypatch.setattr(criticidades, "_get_matriz_row_context", lambda *_args: matrix_row)
    monkeypatch.setattr(
        criticidades,
        "get_incompatibilidade_patologica_data",
        lambda *_args: SimpleNamespace(model_dump=lambda: {"patologias": []}),
    )
    result = criticidades._build_incompatibilidade_patologica_context("12345678000195", None, None)
    assert result["percentual"] == 0.0 and result["multiplicador_brasil"] == 0.0

    monkeypatch.setattr(
        criticidades,
        "get_incompatibilidade_patologica_data",
        lambda *_args: (_ for _ in ()).throw(ValueError("invalid service payload")),
    )
    with pytest.raises(RuntimeError, match="Detalhamento clinico indisponivel"):
        criticidades._build_incompatibilidade_patologica_context("12345678000195", None, None)

    payload = {
        "patologias": [{
            "municipal_resumo": [], "ranking_municipal": [],
            "demografia_parkinson": {
                "populacao_total": 100, "populacao_50_mais": 20, "casos_esperados": 0,
                "cpfs_observados": 1, "razao_observado_esperado": None, "prevalencia_50_mais": 0.1,
            },
        }]
    }
    monkeypatch.setattr(
        criticidades, "get_incompatibilidade_patologica_data",
        lambda *_args: SimpleNamespace(model_dump=lambda: payload),
    )
    with pytest.raises(RuntimeError, match="Demografia de Parkinson invalida"):
        criticidades._build_incompatibilidade_patologica_context("12345678000195", None, None)

    clinical = Document()
    criticidades._add_incompatibilidade_patologica_text(
        clinical, "7.1", "Farmácia", {
            "periodo_desc": "no período analisado", "percentual": 12.0, "valor_suspeito": 10.0,
            "multiplicador_regiao": 1.0, "multiplicador_uf": 1.0,
            "multiplicador_brasil": 1.0, "ranking_patologias": [],
        }, 1,
    )
    assert "R$ 10,00" in " ".join(paragraph.text for paragraph in clinical.paragraphs)


def test_teto_context_and_text_reject_missing_or_inconsistent_financial_values(monkeypatch):
    valid_row = {
        "pct_teto": "inválido", "teto_valor": 20, "teto_valor_total": 30, "valor_total_vendas": 50,
        "med_teto_reg": "inválido", "med_teto_uf": 6, "med_teto_br": 7,
        "risco_teto_reg": 3, "risco_teto_uf": 2, "risco_teto_br": 1,
    }
    row = dict(valid_row)
    monkeypatch.setattr(criticidades, "_get_matriz_row_context", lambda *_args: row)
    context = criticidades._build_teto_context("12345678000195", None, None)
    assert context["percentual"] == 0.0 and context["mediana_regiao"] == 0.0
    for key, value in (("teto_valor", None), ("teto_valor_total", "bad"), ("valor_total_vendas", -1)):
        broken = {**valid_row, key: value}
        monkeypatch.setattr(criticidades, "_get_matriz_row_context", lambda *_args, broken=broken: broken)
        with pytest.raises(RuntimeError, match=f"{key} (ausente|invalido)"):
            criticidades._build_teto_context("12345678000195", None, None)
    with pytest.raises(RuntimeError, match="Vendas totais inferiores"):
        criticidades._add_teto_text(
            Document(), "7.1", "Farmácia", {
                "periodo_desc": "no período analisado", "percentual": 10, "valor_suspeito": 5,
                "valor_monitorado": 20, "valor_total_vendas": 10,
                "multiplicador_regiao": 1, "multiplicador_uf": 1, "multiplicador_brasil": 1,
            },
        )


def test_geographic_context_coerces_invalid_optional_matrix_values_to_zero(monkeypatch):
    cnpj = "12345678000195"
    monkeypatch.setattr(
        criticidades, "_get_matriz_row_context",
        lambda *_args: {"pct_geografico": "bad", "risco_geografico_reg": "bad"},
    )
    monkeypatch.setattr(
        criticidades, "get_df_dados_farmacia",
        lambda: pl.DataFrame({"id_cnpj": [7], "cnpj": [cnpj], "uf": ["SP"]}),
    )
    monkeypatch.setattr(
        criticidades, "scan_geografico_origem_uf",
        lambda: pl.DataFrame({
            "id_cnpj": [7], "ano_base": [2024], "uf_farmacia": ["SP"], "uf_paciente": ["RJ"],
            "is_outra_uf": [True], "qtd_autorizacoes": [1], "valor_autorizado": [10.0],
        }).lazy(),
    )
    result = criticidades._build_dispersao_geografica_context(cnpj, None, None)
    assert result["percentual"] == 0.0 and result["multiplicador_regiao"] == 0.0
