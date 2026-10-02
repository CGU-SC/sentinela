from datetime import date
from types import SimpleNamespace

from docx import Document
from docx.shared import Inches

from backend.api.services.analytics import nota_tecnica_anexos as anexos


class _Timing:
    def __init__(self):
        self.messages = []

    def mark(self, message):
        self.messages.append(message)


def test_deceased_municipality_formatting_and_cpf_grouping_are_normalized_and_sorted():
    for value in (None, "", "-", "—", "–", "N/I", "NI", "N/A", "NAO INFORMADO", "Não Informado"):
        assert anexos._format_municipio_falecido(value) == "Sem Registro"
    assert anexos._format_municipio_falecido("  SAO PAULO ") == "Sao Paulo"
    assert anexos._format_municipio_uf_falecido_cell({"municipio": "Sem Registro", "uf": "SP"}) == [
        {"text": "Sem Registro", "color": "DC2626"},
        {"text": "/SP", "color": "0F172A"},
    ]
    assert anexos._format_municipio_uf_falecido_cell({"municipio": "Campinas", "uf": "SP"}) == "Campinas/SP"

    transactions = [
        SimpleNamespace(
            cpf="2", nome_falecido="maria SILVA", municipio_falecido="Campinas", uf_falecido="SP",
            dt_obito=date(2020, 1, 1), outros_estabelecimentos=None, valor_total_autorizacao=20,
            dias_apos_obito=90, data_autorizacao=date(2024, 2, 1), num_autorizacao="B",
        ),
        SimpleNamespace(
            cpf="1", nome_falecido="joao santos", municipio_falecido="N/I", uf_falecido=None,
            dt_obito=None, outros_estabelecimentos="outro", valor_total_autorizacao=10,
            dias_apos_obito=30, data_autorizacao=None, num_autorizacao="Z",
        ),
        SimpleNamespace(
            cpf="2", nome_falecido="ignored duplicate", municipio_falecido="ignored", uf_falecido="RJ",
            dt_obito=date(2019, 1, 1), outros_estabelecimentos=None, valor_total_autorizacao=5,
            dias_apos_obito=120, data_autorizacao=date(2024, 1, 1), num_autorizacao="A",
        ),
    ]
    groups = anexos._build_falecidos_grupos(transactions)
    assert [group["cpf"] for group in groups] == ["00000000001", "00000000002"]
    assert groups[0]["nome"] == "Joao Santos" and groups[0]["municipio"] == "Sem Registro"
    assert groups[0]["uf"] == "—" and groups[0]["total_valor"] == 10.0
    assert groups[1]["nome"] == "Maria Silva" and groups[1]["total_valor"] == 25.0
    assert groups[1]["max_dias"] == 120
    assert [item.num_autorizacao for item in groups[1]["transacoes"]] == ["A", "B"]


def test_crm_evidence_annex_resets_its_own_table_number_and_marks_timing(monkeypatch):
    calls = []
    monkeypatch.setattr(
        anexos,
        "_add_crm_evidencias_complementares_body",
        lambda *args: calls.append(args),
    )
    doc = Document()
    timing = _Timing()
    next_number = anexos._add_anexo_crm_evidencias(
        doc, "Farmacia Alvo", {"rows": [1]}, 14, "001/2024", timing=timing, anexo_num="IV"
    )
    assert next_number == 14
    assert calls == [(doc, "Farmacia Alvo", {"rows": [1]}, 0)]
    assert any("ANEXO IV" in paragraph.text for paragraph in doc.paragraphs)
    assert timing.messages == ["anexo IV evidencias CRM"]


def test_deceased_sales_annex_renders_grouped_rows_subtotals_and_totals():
    transactions = [
        SimpleNamespace(
            cpf="12345678901", nome_falecido="maria Silva", municipio_falecido="Campinas", uf_falecido="SP",
            dt_obito=date(2020, 1, 1), outros_estabelecimentos=None, valor_total_autorizacao=25.5,
            dias_apos_obito=30, data_autorizacao=date(2024, 1, 2), num_autorizacao="A1",
        ),
        SimpleNamespace(
            cpf="12345678901", nome_falecido="maria Silva", municipio_falecido="Campinas", uf_falecido="SP",
            dt_obito=date(2020, 1, 1), outros_estabelecimentos=None, valor_total_autorizacao=10.0,
            dias_apos_obito=45, data_autorizacao=date(2024, 1, 3), num_autorizacao="A2",
        ),
        SimpleNamespace(
            cpf="98765432100", nome_falecido="joao Souza", municipio_falecido=None, uf_falecido="RJ",
            dt_obito=date(2019, 2, 1), outros_estabelecimentos=None, valor_total_autorizacao=5.0,
            dias_apos_obito=60, data_autorizacao=date(2024, 2, 1), num_autorizacao="B1",
        ),
    ]
    doc = Document()
    doc.sections[0].page_width = Inches(11)
    doc.sections[0].page_height = Inches(8.5)
    timing = _Timing()
    anexos._add_anexo_falecidos(
        doc,
        "Farmacia Alvo",
        "12.345.678/0001-90",
        {
            "transacoes": transactions,
            "total_autorizacoes": 3,
            "cpfs_distintos": 2,
            "valor_total": 40.5,
            "periodo_desc": "em 2024",
        },
        "001/2024",
        timing=timing,
        anexo_num="IV",
    )

    assert any("ANEXO IV" in paragraph.text for paragraph in doc.paragraphs)
    assert any("em 2024" in paragraph.text for paragraph in doc.paragraphs)
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 7
    assert "TOTAL GERAL - 2 CPF(s) distintos - 3 autorização(ões)" in table.rows[-1].cells[0].text
    assert timing.messages[0].startswith("anexo IV agrupamento (2 CPFs, 3 transacoes)")
    assert timing.messages[-1] == "anexo IV total geral"


def test_deceased_sales_annex_without_transactions_does_not_add_a_page():
    doc = Document()
    anexos._add_anexo_falecidos(doc, "Farmacia", "123", {"transacoes": []}, "001/2024")
    assert len(doc.sections) == 1
    assert doc.paragraphs == []
