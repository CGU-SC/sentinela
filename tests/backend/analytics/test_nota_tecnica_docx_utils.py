from pathlib import Path
import sys
from tempfile import TemporaryDirectory

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, RGBColor

from backend.api.services.analytics import nota_tecnica_docx_utils as utils


def _children(element, tag):
    return element.findall(qn(f"w:{tag}"))


@pytest.fixture
def temp_dir():
    with TemporaryDirectory(prefix="sentinela-nt-tests-", dir=Path.cwd()) as directory:
        yield Path(directory)


def test_asset_resolution_in_development_and_frozen_build(temp_dir, monkeypatch):
    project_root = temp_dir / "sentinela"
    module_path = project_root / "backend/api/services/analytics/nota_tecnica_docx_utils.py"
    module_path.parent.mkdir(parents=True)
    monkeypatch.setattr(utils, "__file__", str(module_path))
    asset = project_root / "frontend/public/img/brasao_republica_mini.jpg"
    asset.parent.mkdir(parents=True)
    asset.write_bytes(b"asset")

    assert utils._resolve_nota_tecnica_asset("img/brasao_republica_mini.jpg") == str(asset)
    assert utils._resolve_brasao_republica_path() == str(asset)

    bundle = temp_dir / "bundle"
    frozen_asset = bundle / "frontend/dist/img/brasao_republica_mini.jpg"
    frozen_asset.parent.mkdir(parents=True)
    frozen_asset.write_bytes(b"frozen")
    monkeypatch.setattr(sys, "frozen", True, raising=False)
    monkeypatch.setattr(sys, "_MEIPASS", str(bundle), raising=False)
    assert utils._resolve_nota_tecnica_asset("img/brasao_republica_mini.jpg") == str(frozen_asset)


@pytest.mark.parametrize("asset_path", ["", ".", "../absolute-looking"])
def test_asset_resolver_rejects_invalid_or_missing_required_asset(asset_path, temp_dir, monkeypatch):
    module_path = temp_dir / "sentinela/backend/api/services/analytics/module.py"
    module_path.parent.mkdir(parents=True)
    monkeypatch.setattr(utils, "__file__", str(module_path))
    with pytest.raises((RuntimeError, FileNotFoundError)):
        utils._resolve_nota_tecnica_asset(asset_path)


def test_asset_resolver_rejects_absolute_paths_and_reports_missing_asset(temp_dir, monkeypatch):
    module_path = temp_dir / "sentinela/backend/api/services/analytics/module.py"
    module_path.parent.mkdir(parents=True)
    monkeypatch.setattr(utils, "__file__", str(module_path))
    with pytest.raises(RuntimeError, match="Caminho relativo"):
        utils._resolve_nota_tecnica_asset(str(temp_dir / "absolute.png"))
    with pytest.raises(FileNotFoundError, match="Asset obrigatorio"):
        utils._resolve_nota_tecnica_asset("missing.png")


def test_bookmark_and_internal_hyperlink_xml_and_validation():
    doc = Document()
    paragraph = doc.add_paragraph("anchor")
    utils._add_bookmark(paragraph, "section_1")
    starts = _children(paragraph._p, "bookmarkStart")
    ends = _children(paragraph._p, "bookmarkEnd")
    assert starts[0].get(qn("w:name")) == "section_1"
    assert starts[0].get(qn("w:id")) == ends[0].get(qn("w:id"))

    second = doc.add_paragraph()
    link = utils._add_internal_hyperlink(second, " Go ", "section_1", color="123456", size=8.5, bold=True, underline=False)
    assert link.get(qn("w:anchor")) == "section_1"
    assert link.find(f".//{qn('w:b')}") is not None
    assert link.find(f".//{qn('w:u')}").get(qn("w:val")) == "none"
    assert link.find(f".//{qn('w:sz')}").get(qn("w:val")) == "17"
    with pytest.raises(RuntimeError, match="bookmark obrigatorio"):
        utils._add_bookmark(doc.add_paragraph(), "")
    with pytest.raises(RuntimeError, match="Destino de hyperlink"):
        utils._add_internal_hyperlink(doc.add_paragraph(), "x", "")


def test_document_fields_and_toc_with_and_without_fallback_entries():
    doc = Document()
    settings = doc.settings.element
    for value in ("false", "false"):
        field = utils.OxmlElement("w:updateFields")
        field.set(qn("w:val"), value)
        settings.append(field)
    utils._mark_document_fields_for_update(doc)
    assert len(settings.findall(qn("w:updateFields"))) == 1
    assert settings.find(qn("w:updateFields")).get(qn("w:val")) == "true"

    paragraph = doc.add_paragraph()
    utils._add_word_toc_field(paragraph, levels="1-3", fallback_entries=[("1", "Intro", "2"), ("2", "Dados", "5")])
    xml = paragraph._p.xml
    assert 'TOC \\o "1-3" \\h \\z \\u' in xml
    assert "Intro\t2" in xml and "Dados\t5" in xml
    assert len(paragraph._p.findall(f".//{qn('w:br')}")) == 1
    default = doc.add_paragraph()
    utils._add_word_toc_field(default)
    assert "Sumário automático" in default._p.xml
    with pytest.raises(RuntimeError, match="Niveis do sumario"):
        utils._add_word_toc_field(doc.add_paragraph(), levels="")


def test_external_hyperlink_adds_relationship_and_preserves_edge_spaces():
    paragraph = Document().add_paragraph()
    link = utils._add_external_hyperlink(paragraph, " site ", "https://example.test", bold=True, underline=False, size=9)
    assert paragraph.part.rels[link.get(qn("r:id"))].target_ref == "https://example.test"
    assert link.find(f".//{qn('w:t')}").get("{http://www.w3.org/XML/1998/namespace}space") == "preserve"
    assert link.find(f".//{qn('w:b')}") is not None
    assert link.find(f".//{qn('w:u')}").get(qn("w:val")) == "none"
    normal = utils._add_external_hyperlink(paragraph, "site", "https://other.test")
    assert normal.find(f".//{qn('w:t')}").get("{http://www.w3.org/XML/1998/namespace}space") is None
    with pytest.raises(RuntimeError, match="URL de hyperlink"):
        utils._add_external_hyperlink(paragraph, "x", "")


def test_cell_shading_borders_and_table_border_modes_replace_previous_xml():
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    run = cell.paragraphs[0].add_run("text")
    utils._cell_bg(cell, "ABCDEF")
    utils._cell_bg(cell, "123456")
    utils._cell_bg_run(run, "FEDCBA")
    utils._cell_borders(cell, left={"sz": "8", "color": "010203"}, top={"sz": "2", "color": "040506"})
    utils._cell_borders(cell)
    tc_pr = cell._tc.get_or_add_tcPr()
    assert len(tc_pr.findall(qn("w:shd"))) == 1
    assert tc_pr.find(qn("w:shd")).get(qn("w:fill")) == "123456"
    assert len(tc_pr.findall(qn("w:tcBorders"))) == 1
    assert tc_pr.find(qn("w:tcBorders")).find(qn("w:left")).get(qn("w:val")) == "none"
    assert run._r.get_or_add_rPr().find(qn("w:shd")).get(qn("w:fill")) == "FEDCBA"

    utils._tbl_no_borders(table)
    utils._tbl_no_borders(table)
    utils._set_table_open_borders(table, color="AA0000", sz="7")
    utils._set_table_grid_borders(table)
    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    assert len(table._tbl.tblPr.findall(qn("w:tblBorders"))) == 1
    assert borders.find(qn("w:insideV")).get(qn("w:val")) == "single"
    assert borders.find(qn("w:insideV")).get(qn("w:color")) == "CBD5E1"
    # Exercise table-property creation when it is absent.
    raw_table = utils.OxmlElement("w:tbl")
    utils._set_table_open_borders(type("TableProxy", (), {"_tbl": raw_table})())
    assert raw_table.find(qn("w:tblPr")) is not None
    raw_table = utils.OxmlElement("w:tbl")
    proxy = type("TableProxy", (), {"_tbl": raw_table})()
    utils._tbl_no_borders(proxy)
    utils._set_table_grid_borders(type("TableProxy", (), {"_tbl": utils.OxmlElement("w:tbl")})())


def test_rows_table_together_and_block_paragraph_formatters():
    doc = Document()
    title = doc.add_paragraph("title")
    table = doc.add_table(rows=1, cols=2)
    trailing = [doc.add_paragraph("note"), doc.add_paragraph("last")]
    utils._keep_small_table_together(title, table, trailing)
    assert title.paragraph_format.keep_with_next
    assert table.rows[0]._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
    assert trailing[-1].paragraph_format.keep_with_next is False
    assert table.cell(0, 1).paragraphs[0].paragraph_format.keep_with_next
    utils._row_cant_split(table.rows[0])
    utils._repeat_table_header(table.rows[0])
    utils._repeat_table_header(table.rows[0])
    assert len(table.rows[0]._tr.get_or_add_trPr().findall(qn("w:tblHeader"))) == 1

    empty_table = doc.add_table(rows=0, cols=1)
    no_trailing_title = doc.add_paragraph()
    utils._keep_small_table_together(no_trailing_title, empty_table)
    single = doc.add_table(rows=1, cols=1)
    only_title = doc.add_paragraph()
    utils._keep_small_table_together(only_title, single)
    assert single.cell(0, 0).paragraphs[-1].paragraph_format.keep_with_next is False

    centered = doc.add_paragraph()
    utils._format_block_title(centered, space_before=3, space_after=4, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    assert centered.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert centered.paragraph_format.keep_with_next and centered.paragraph_format.keep_together
    footnote = doc.add_paragraph()
    utils._format_block_footnote(footnote, alignment=WD_ALIGN_PARAGRAPH.LEFT)
    assert footnote.paragraph_format.keep_with_next is False
    picture = doc.add_paragraph()
    utils._format_picture_paragraph(picture, keep_with_next=False)
    assert picture.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert picture.paragraph_format.keep_with_next is False


def test_rgb_and_run_formatting_including_bold_color_adjustment():
    assert utils._rgb("12ABEF") == RGBColor(18, 171, 239)
    doc = Document()
    normal = utils._run(doc.add_paragraph(), "bold", bold=True, color="0F172A", italic=True, underline=True)
    assert normal.bold and normal.italic and normal.underline
    assert normal.font.color.rgb == RGBColor(51, 65, 85)
    custom = utils._run(doc.add_paragraph(), "custom", color="FF0000", size=12, bold=True)
    assert custom.font.color.rgb == RGBColor(255, 0, 0)
    assert custom.font.size.pt == 12


def test_footnote_creation_reuse_links_and_validation():
    doc = Document()
    paragraph = doc.add_paragraph()
    ref = utils._footnote_ref(doc, paragraph, 1, "A source")
    assert ref._r.find(qn("w:footnoteReference")).get(qn("w:id")) == "1"
    part = doc.part.part_related_by(utils.RT.FOOTNOTES)
    assert len([f for f in part.element.findall(qn("w:footnote")) if f.get(qn("w:id")) == "1"]) == 1
    utils._ensure_footnote(doc, 1, "ignored on duplicate")
    assert len([f for f in part.element.findall(qn("w:footnote")) if f.get(qn("w:id")) == "1"]) == 1

    utils._footnote_ref(doc, paragraph, 2, "See source URL", hyperlink_text="source", hyperlink_url="https://example.test")
    assert any(relationship.target_ref == "https://example.test" for relationship in part.rels.values())
    with pytest.raises(RuntimeError, match="URL obrigatoria"):
        utils._ensure_footnote(doc, 3, "linked text", hyperlink_text="text")
    with pytest.raises(RuntimeError, match="Texto do hyperlink nao encontrado"):
        utils._ensure_footnote(doc, 3, "different text", hyperlink_text="absent", hyperlink_url="https://example.test")
    with pytest.raises(RuntimeError, match="Texto e URL obrigatorios"):
        utils._append_footnote_hyperlink(part, utils.OxmlElement("w:p"), "", "https://example.test")
    with pytest.raises(RuntimeError, match="Texto e URL obrigatorios"):
        utils._append_footnote_hyperlink(part, utils.OxmlElement("w:p"), "link", "")
    linked_paragraph = utils.OxmlElement("w:p")
    utils._append_footnote_hyperlink(part, linked_paragraph, " trailing ", "https://space.test")
    assert linked_paragraph.find(f".//{qn('w:t')}").get("{http://www.w3.org/XML/1998/namespace}space") == "preserve"

    p = utils.OxmlElement("w:p")
    utils._append_footnote_text_run(p, " edge ")
    assert p.find(f".//{qn('w:t')}").get("{http://www.w3.org/XML/1998/namespace}space") == "preserve"
    before = len(p)
    utils._append_footnote_text_run(p, "")
    assert len(p) == before


def test_table_width_writing_and_fast_row_xml():
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    widths = [Inches(1.25), Inches(2.5)]
    utils._set_table_fixed_widths(table, widths)
    assert table.autofit is False
    assert table._tbl.tblPr.find(qn("w:tblLayout")).get(qn("w:type")) == "fixed"
    assert table._tbl.tblPr.find(qn("w:tblW")).get(qn("w:w")) == str(sum(int(w.twips) for w in widths))
    assert len(table._tbl.tblGrid.findall(qn("w:gridCol"))) == 2
    cell = table.cell(0, 0)
    utils._set_cell_width(cell, widths[0])
    tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
    assert tc_w.get(qn("w:type")) == "dxa"
    cell._tc.get_or_add_tcPr().remove(tc_w)

    class CellWithoutWidth:
        _tc = utils.OxmlElement("w:tc")

        @property
        def width(self):
            return None

        @width.setter
        def width(self, value):
            pass

    raw_cell = CellWithoutWidth()
    utils._set_cell_width(raw_cell, widths[0])
    assert raw_cell._tc.get_or_add_tcPr().find(qn("w:tcW")) is not None
    utils._write_cell(cell, "texto", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    assert cell.text == "texto"
    assert cell.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    utils._write_cell(cell, "normal")

    utils._append_table_row_fast(
        table,
        ["plain", ["first", {"text": "second", "color": "FF0000", "size": 8, "break": False}, {"break": True}], [{"text": " spaced "}, 42]],
        widths,
        alignments=["center"],
        sizes=[7],
        color="112233",
        fill="ABCDEF",
    )
    row = table._tbl.tr_lst[-1]
    cells = row.tc_lst
    assert len(cells) == 3
    assert cells[0].find(f".//{qn('w:t')}").text == "plain"
    assert cells[1].findall(f".//{qn('w:br')}")
    assert cells[2].find(f".//{qn('w:t')}").get("{http://www.w3.org/XML/1998/namespace}space") == "preserve"
    assert cells[0].find(qn("w:tcPr")).find(qn("w:shd")).get(qn("w:fill")) == "ABCDEF"

    utils._append_table_row_fast(table, ["no width", "no alignment", "no sizes"], [])
    assert table._tbl.tr_lst[-1].tc_lst[1].find(f".//{qn('w:jc')}") is None

    table_properties = table._tbl.tblPr
    layout = table_properties.find(qn("w:tblLayout"))
    table_width = table_properties.find(qn("w:tblW"))
    if layout is not None:
        table_properties.remove(layout)
    if table_width is not None:
        table_properties.remove(table_width)

    class TableWithoutSizing:
        _tbl = table._tbl
        columns = table.columns
        autofit = None

    utils._set_table_fixed_widths(TableWithoutSizing(), widths)


def test_fast_row_preserves_default_segments_and_empty_dict_values():
    doc = Document()
    table = doc.add_table(rows=0, cols=1)
    utils._append_table_row_fast(table, [["left", {"text": "right", "break": True}, {"text": "", "color": "AA0000"}]], [], sizes=[5.5])
    row_xml = table._tbl.tr_lst[0].xml
    assert "left" in row_xml and "right" not in row_xml
    assert row_xml.count("<w:br") == 1
