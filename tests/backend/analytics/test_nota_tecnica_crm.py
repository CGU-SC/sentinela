from datetime import date, datetime, timezone
from types import SimpleNamespace

import polars as pl
import pytest
from docx import Document

from backend.api.services.analytics import nota_tecnica_crm as crm_nt


def test_scalar_formatting_and_required_values():
    assert crm_nt._as_float("2.5") == 2.5
    assert crm_nt._as_float("bad") == 0.0
    assert crm_nt._as_optional_float(None) is None
    assert crm_nt._as_optional_float("bad") is None
    assert crm_nt._as_optional_float(float("nan")) is None
    assert crm_nt._as_optional_float("2.5") == 2.5
    assert crm_nt._as_int("2.9") == 2
    assert crm_nt._as_int("bad") == 0
    assert crm_nt._required_positive_float("2", "value", "test") == 2
    assert crm_nt._required_nonnegative_float(0, "value", "test") == 0
    assert crm_nt._required_nonnegative_int("2.9", "value", "test") == 2
    assert crm_nt._format_optional_decimal_pt(None, empty="—") == "—"
    assert crm_nt._format_optional_decimal_pt(12.5, 2) == "12,50"
    assert crm_nt._volume_horario_multiplicador_value({"mediana_hora": 0, "nu_prescricoes": 3}) == 3
    assert crm_nt._format_volume_horario_multiplicador({"mediana_hora": 2, "nu_prescricoes": 5}) == "2,5x"
    assert crm_nt._date_to_competencia(date(2024, 6, 2)) == 202406
    assert crm_nt._date_to_competencia(None) is None
    assert crm_nt._select_crm_concentracao_table_rows([{"taxa_hora": 30}, {"taxa_hora": 31}]) == [{"taxa_hora": 31}]
    assert crm_nt._select_crm_unico_concentracao_table_rows([{"taxa_hora": 31}]) == [{"taxa_hora": 31}]
    assert crm_nt._crm_event_key({"dt": "2024-02-03T12:00", "hr": "8"}) == ("2024-02-03", 8)
    assert crm_nt._vez_ou_vezes(1) == "vez"
    assert crm_nt._vez_ou_vezes(1.1) == "vezes"
    assert crm_nt._crm_num_uf("123/ SP ") == ("123", "SP")
    assert crm_nt._crm_num_uf("123") == ("123", "")
    assert crm_nt._crm_num_uf(None) == ("Não informado", "")
    assert crm_nt._format_competencia_crm(202401) == "01/2024"
    assert crm_nt._format_competencia_crm("invalid") == "invalid"
    assert crm_nt._format_competencia_crm(None) == "—"
    assert crm_nt._required_competencia_crm("202401", "test") == "01/2024"
    assert crm_nt._format_time_hour("8") == "08:00"
    assert crm_nt._format_time_hour("bad") == "—"
    assert crm_nt._format_hour_range(23) == "23:00 até 00:00"
    assert crm_nt._format_hour_range("bad") == "—"
    assert crm_nt._format_janela_minutos(0) == "Mesmo instante"
    assert crm_nt._format_janela_minutos(15) == "15 min"
    assert crm_nt._format_janela_minutos(120) == "2h"
    assert crm_nt._format_janela_minutos(125) == "2h 5min"
    assert crm_nt._plural(1, "singular", "plural") == "singular"
    assert crm_nt._plural(2, "singular", "plural") == "plural"

    with pytest.raises(ValueError, match="ausente ou invalido"):
        crm_nt._required_positive_float(0, "value", "test")
    with pytest.raises(ValueError, match="ausente ou invalido"):
        crm_nt._required_nonnegative_float(-1, "value", "test")
    with pytest.raises(ValueError, match="ausente ou invalido"):
        crm_nt._required_nonnegative_int("bad", "value", "test")
    with pytest.raises(ValueError, match="ausente ou invalido"):
        crm_nt._required_nonnegative_int(-1, "value", "test")
    with pytest.raises(ValueError, match="ausente ou invalido"):
        crm_nt._volume_horario_multiplicador_value({"mediana_hora": -1, "nu_prescricoes": 1})
    with pytest.raises(ValueError, match="ausente ou invalido"):
        crm_nt._required_competencia_crm("2024", "test")
    with pytest.raises(ValueError, match="ausente ou invalido"):
        crm_nt._select_crm_unico_concentracao_table_rows([{"taxa_hora": 0}])


def test_date_datetime_and_alert_window_helpers():
    assert crm_nt._format_date_br(None) == "Não localizada"
    assert crm_nt._format_date_br(datetime(2024, 2, 3, 4)) == "03.02.2024"
    assert crm_nt._format_date_br(date(2024, 2, 3)) == "03.02.2024"
    assert crm_nt._format_date_br("   ") == crm_nt._format_date_br(None)
    assert crm_nt._format_date_br("2024-02-03 10:20:00") == "03.02.2024"
    assert crm_nt._format_date_br("03/02/2024") == "03.02.2024"
    assert crm_nt._format_date_br("unparsed") == "unparsed"
    assert crm_nt._as_date_for_weekend_marker(datetime(2024, 2, 3, 1)) == date(2024, 2, 3)
    assert crm_nt._as_date_for_weekend_marker(date(2024, 2, 3)) == date(2024, 2, 3)
    assert crm_nt._as_date_for_weekend_marker("   ") is None
    assert crm_nt._as_date_for_weekend_marker("2024-02-03") == date(2024, 2, 3)
    assert crm_nt._as_date_for_weekend_marker("bad") is None
    assert crm_nt._parse_datetime_crm(None) is None
    assert crm_nt._parse_datetime_crm(" ") is None
    class TimestampLike:
        @staticmethod
        def to_pydatetime():
            return datetime(2024, 2, 3, 12, 30, tzinfo=timezone.utc)

    assert crm_nt._parse_datetime_crm(TimestampLike()) == datetime(2024, 2, 3, 12, 30)
    assert crm_nt._parse_datetime_crm(date(2024, 2, 3)) == datetime(2024, 2, 3)
    assert crm_nt._parse_datetime_crm(datetime(2024, 2, 3, tzinfo=None)) == datetime(2024, 2, 3)
    assert crm_nt._parse_datetime_crm("2024-02-03T12:30:00Z") == datetime(2024, 2, 3, 12, 30)
    assert crm_nt._parse_datetime_crm("2024-02-03 12:30:00.123") == datetime(2024, 2, 3, 12, 30, 0, 123000)
    assert crm_nt._parse_datetime_crm("invalid") is None
    assert crm_nt._format_datetime_br_minute("bad") == "N/d"
    assert crm_nt._format_datetime_br_minute("2024-02-03 12:30:00") == "03/02/24 12:30"
    assert crm_nt._datetime_in_alert_window(datetime(2024, 2, 3, 10, 1), datetime(2024, 2, 3, 10), datetime(2024, 2, 3, 10, 2))
    assert crm_nt._datetime_in_alert_window(datetime(2024, 2, 3, 10, 1), datetime(2024, 2, 3, 10, 0, 30), datetime(2024, 2, 3, 10, 1, 30))
    assert not crm_nt._datetime_in_alert_window(datetime(2024, 2, 3, 10, 1, 1), datetime(2024, 2, 3, 10), datetime(2024, 2, 3, 10, 1))
    assert crm_nt._datetime_windows_overlap("2024-02-03 10:00", "2024-02-03 10:10", "2024-02-03 10:10", "2024-02-03 10:20")
    assert crm_nt._datetime_windows_overlap("2024-02-03 10:10", "2024-02-03 10:00", "2024-02-03 10:05", "2024-02-03 10:15")
    assert crm_nt._datetime_windows_overlap("2024-02-03 10:00", "2024-02-03 10:10", "2024-02-03 10:20", "2024-02-03 10:10")
    assert not crm_nt._datetime_windows_overlap("bad", "bad", "2024-02-03", "2024-02-04")


def test_crm_alert_context_labels_and_document_cells():
    labels = crm_nt._crm_alertas_contexto_labels({
        "flag_crm_invalido": 1,
        "flag_prescricao_antes_registro": 1,
        "flag_crm_exclusivo": 1,
        "qtd_alertas_crm_unico": 1,
        "alertas_crm_multiplos": [{}],
        "alerta5_geografico": True,
        "flag_robo": 1,
        "flag_robo_oculto": 1,
    })
    assert labels == [
        "CRM não localizado", "CRM irregular", "CRM exclusivo",
        "Concentração em único CRM", "Concentração em múltiplos CRMs",
        "Registro geográfico", "Volume diário atípico local/Brasil",
    ]
    assert crm_nt._crm_alertas_contexto_labels({"nu_prescricoes_dia": crm_nt.CRM_DAILY_RATE_ALERT_THRESHOLD + 1}) == ["Volume diário atípico local"]
    assert crm_nt._crm_alertas_contexto_labels({"prescricoes_dia_total_brasil": crm_nt.CRM_DAILY_RATE_ALERT_THRESHOLD + 1}) == ["Volume diário atípico Brasil"]

    doc = Document()
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    crm_nt._write_crm_alertas_cell(cell, [])
    assert "Sem alerta" in cell.text
    crm_nt._write_crm_alertas_cell(cell, ["Concentração em único CRM", "Custom"])
    assert "Muitas autorizações em sequência pelo mesmo CRM" in cell.text
    assert "Custom" in cell.text
    crm_nt._add_crm_alert_legend(doc, [])
    crm_nt._add_crm_alert_legend(doc, [{"alertas_contexto": ["CRM irregular", "Concentração em único CRM"]}])
    assert len(doc.tables) == 3
    crm_nt._write_date_cell_with_weekend_marker(cell, "2024-02-03")
    assert "(fim de semana)" in cell.text
    crm_nt._write_date_cell_with_weekend_marker(cell, "2024-02-05")
    assert "(fim de semana)" not in cell.text


def test_enrich_crm_alert_values_from_real_parquet(tmp_path, monkeypatch):
    monkeypatch.setattr(crm_nt, "_get_cnpj_cache_dir", lambda _cnpj: str(tmp_path))
    (tmp_path / crm_nt.CRM_RAIOX_TX_PARQUET).unlink(missing_ok=True)
    unico = [{
        "id_medico": "123/SP", "dt_ini_hora": "2024-02-03 10:00:30", "dt_fim_hora": "2024-02-03 10:01:30"
    }, {"id_medico": "", "dt_ini_hora": None, "dt_fim_hora": None}]
    crm_nt._enrich_crm_unico_valores("123", unico)
    assert unico[0]["valor_alerta"] == 0
    assert unico[0]["valor_alerta_disponivel"] is False
    assert unico[1]["valor_alerta_disponivel"] is False

    pl.DataFrame({
        "dt_janela": ["2024-02-03", "2024-02-03", "2024-02-03", "2024-02-04", "2024-02-03"],
        "data_hora": ["2024-02-03 10:01:00", "2024-02-03 10:02:00", "2024-02-03 10:01:00", "2024-02-04 10:01:00", "invalid"],
        "id_medico": ["123/SP", "123/SP", "999/RJ", "123/SP", "123/SP"],
        "valor_pago": [12.5, 9.0, 100.0, 88.0, 1000.0],
    }).write_parquet(tmp_path / crm_nt.CRM_RAIOX_TX_PARQUET)
    crm_nt._enrich_crm_unico_valores("123", unico[:1])
    assert unico[0]["valor_alerta"] == 21.5
    assert unico[0]["valor_alerta_disponivel"] is True

    multiplo = [{"dt_ini_hora": "2024-02-03 10:00:00", "dt_fim_hora": "2024-02-03 10:03:00"}]
    crm_nt._enrich_crm_multiplo_valores("123", multiplo)
    assert multiplo[0]["valor_alerta"] == 121.5
    assert multiplo[0]["valor_alerta_disponivel"] is True
    monkeypatch.setattr(crm_nt.pl, "scan_parquet", lambda *_args, **_kwargs: (_ for _ in ()).throw(OSError("broken cache")))
    crm_nt._enrich_crm_unico_valores("123", unico[:1])
    crm_nt._enrich_crm_multiplo_valores("123", multiplo)
    assert unico[0]["valor_alerta_disponivel"] is False
    assert multiplo[0]["valor_alerta_disponivel"] is False


def test_principais_crms_context_rows_include_out_of_top_ten_alerts():
    rows = [
        {
            "id_medico": f"{i}/SP", "no_medico": f"medico {i}", "nu_prescricoes": 20 - i,
            "vl_total_prescricoes": 200 - i, "nu_estabelecimentos": 2, "competencia_nu_estabelecimentos": "202401",
        }
        for i in range(12)
    ]
    rows[11]["flag_crm_invalido"] = 1
    result = crm_nt._build_principais_crms_contexto_rows(rows)
    assert len(result) == 11
    assert result[-1]["id_medico"] == "11/SP"
    assert result[-1]["alertas_contexto"] == ["CRM não localizado"]
    assert result[0]["competencia_nu_estabelecimentos"] == "01/2024"
    assert crm_nt._build_principais_crms_contexto_rows([]) == []
    with pytest.raises(ValueError, match="ausente ou invalido"):
        crm_nt._build_principais_crms_contexto_rows([{"nu_estabelecimentos": -1, "competencia_nu_estabelecimentos": "bad"}])


def test_load_crm_evidencias_detalhadas_success_and_failures(monkeypatch):
    hourly = pl.DataFrame({
        "dt_janela": ["2024-01-10", "2024-01-11"], "hr_janela": [10, 11], "nu_prescricoes": [40, 2],
        "nu_crms_diferentes": [3, 1], "mediana_hora": [2.0, 1.0], "is_volume_horario_anomalo": [1, 0],
    })
    monkeypatch.setattr(crm_nt, "load_or_sync_crm_timeline_hora", lambda _cnpj: SimpleNamespace(error=None, df=hourly))
    monkeypatch.setattr(crm_nt, "load_or_sync_crm_unico_alertas", lambda _cnpj: SimpleNamespace(error=None, df=pl.DataFrame()))
    monkeypatch.setattr(crm_nt, "load_or_sync_crm_multi_alertas", lambda _cnpj: SimpleNamespace(error=None, df=pl.DataFrame()))
    monkeypatch.setattr(crm_nt, "sync_crm_raiox_tx", lambda _cnpj: SimpleNamespace(error=None))
    monkeypatch.setattr(crm_nt, "_build_crm_unico_alertas_por_medico", lambda _df: {"1/SP": [{"nu_prescricoes": 3}]})
    monkeypatch.setattr(crm_nt, "_build_alertas_crm_multiplos_por_medico", lambda *_args: {"2/SP": [{"nu_presc_total": 4}]})
    monkeypatch.setattr(crm_nt, "_get_cnpj_cache_dir", lambda _cnpj: "cache")
    alerts, unico, multi = crm_nt._load_crm_evidencias_detalhadas("123", date(2024, 1, 1), date(2024, 1, 31))
    assert alerts == [{"tipo": "VOLUME", "dt": "2024-01-10", "hr": 10, "nu_prescricoes": 40, "nu_crms": 3, "mediana_hora": 2.0}]
    assert unico["1/SP"][0]["nu_prescricoes"] == 3
    assert multi["2/SP"][0]["nu_presc_total"] == 4

    cases = [
        ("load_or_sync_crm_timeline_hora", SimpleNamespace(error="hourly", df=None), "horario"),
        ("load_or_sync_crm_unico_alertas", SimpleNamespace(error="unique", df=None), "unico"),
        ("load_or_sync_crm_multi_alertas", SimpleNamespace(error="multi", df=None), "multiplos"),
        ("sync_crm_raiox_tx", SimpleNamespace(error="raiox"), "Raio-X"),
    ]
    for name, result, match in cases:
        monkeypatch.setattr(crm_nt, "load_or_sync_crm_timeline_hora", lambda _cnpj: SimpleNamespace(error=None, df=hourly))
        monkeypatch.setattr(crm_nt, "load_or_sync_crm_unico_alertas", lambda _cnpj: SimpleNamespace(error=None, df=pl.DataFrame()))
        monkeypatch.setattr(crm_nt, "load_or_sync_crm_multi_alertas", lambda _cnpj: SimpleNamespace(error=None, df=pl.DataFrame()))
        monkeypatch.setattr(crm_nt, "sync_crm_raiox_tx", lambda _cnpj: SimpleNamespace(error=None))
        monkeypatch.setattr(crm_nt, name, lambda *_args, result=result: result)
        with pytest.raises(RuntimeError, match=match):
            crm_nt._load_crm_evidencias_detalhadas("123", None, None)

    monkeypatch.setattr(crm_nt, "load_or_sync_crm_timeline_hora", lambda _cnpj: SimpleNamespace(error=None, df=pl.DataFrame({"bad": [1]})))
    with pytest.raises(RuntimeError, match="colunas obrigatorias"):
        crm_nt._load_crm_evidencias_detalhadas("123", None, None)


def test_hhi_context_validation_and_document_rendering(monkeypatch):
    cnpj = "12345678000195"
    matrix = pl.DataFrame({
        "cnpj": [cnpj], "val_hhi_crm": [250.0], "risco_crm_reg": [2.0], "risco_crm_uf": [3.0], "risco_crm_br": [4.0],
    })
    monkeypatch.setattr(crm_nt, "_build_dynamic_matriz_risco", lambda **_kwargs: matrix)
    data = SimpleNamespace(crms_interesse=[
        {"id_medico": "100/SP", "no_medico": "ana maria", "nu_prescricoes": 80, "vl_total_prescricoes": 800,
         "dt_inscricao_crm": "2001-01-01"},
        {"id_medico": "200/RJ", "nu_prescricoes": 20, "vl_total_prescricoes": 200},
    ])
    result = crm_nt._build_hhi_crm_context(cnpj, date(2024, 1, 1), date(2024, 12, 31), 1000, data)
    assert result["periodo_intervalo"] == "de 01.01.2024 a 31.12.2024"
    assert result["total_autorizacoes"] == 100
    assert result["principal"]["id_medico"] == "100/SP"
    assert result["indice_hhi"] == 250
    doc = Document()
    rendered = []
    crm_nt._add_hhi_crm_text(doc, "7.1", "Farmácia Teste", "12.345.678/0001-95", result, 7, lambda: rendered.append(True), "hhi-crm")
    assert rendered == [True]
    assert len(doc.tables) == 1
    assert "CRM/SP" not in doc.paragraphs[0].text
    assert doc.tables[0].rows[1].cells[0].text == "100/SP"

    with pytest.raises(RuntimeError, match="sem registro obrigatorio"):
        monkeypatch.setattr(crm_nt, "_build_dynamic_matriz_risco", lambda **_kwargs: matrix.clear())
        crm_nt._build_hhi_crm_context(cnpj, None, None, 1000, data)


def test_hhi_context_handles_absence_validation_and_top_crm_limits(monkeypatch):
    cnpj = "12345678000195"
    matrix = pl.DataFrame({
        "cnpj": [cnpj], "val_hhi_crm": [250.0], "risco_crm_reg": [2.0],
        "risco_crm_uf": [3.0], "risco_crm_br": [4.0],
    })
    monkeypatch.setattr(crm_nt, "_build_dynamic_matriz_risco", lambda **_kwargs: matrix)
    monkeypatch.setattr(
        crm_nt,
        "get_crm_data",
        lambda *_args: (_ for _ in ()).throw(RuntimeError("crm offline")),
    )
    assert crm_nt._build_hhi_crm_context(cnpj, None, None, 100) is None
    assert crm_nt._build_hhi_crm_context(cnpj, None, None, 100, SimpleNamespace(crms_interesse=[])) is None

    data = SimpleNamespace(crms_interesse=[{"nu_prescricoes": 1}])
    monkeypatch.setattr(crm_nt, "_build_dynamic_matriz_risco", lambda **_kwargs: pl.DataFrame({"cnpj": [cnpj]}))
    with pytest.raises(RuntimeError, match="colunas obrigatorias"):
        crm_nt._build_hhi_crm_context(cnpj, None, None, 100, data)
    monkeypatch.setattr(crm_nt, "_build_dynamic_matriz_risco", lambda **_kwargs: matrix.with_columns(pl.lit("99999999000191").alias("cnpj")))
    with pytest.raises(RuntimeError, match="sem registro obrigatorio"):
        crm_nt._build_hhi_crm_context(cnpj, None, None, 100, data)

    monkeypatch.setattr(crm_nt, "_build_dynamic_matriz_risco", lambda **_kwargs: matrix)
    for key in ("val_hhi_crm", "risco_crm_reg", "risco_crm_uf", "risco_crm_br"):
        invalid = matrix.with_columns(pl.lit(0.0).alias(key))
        monkeypatch.setattr(crm_nt, "_build_dynamic_matriz_risco", lambda **_kwargs: invalid)
        with pytest.raises(ValueError, match="ausente ou invalido"):
            crm_nt._build_hhi_crm_context(cnpj, None, None, 100, data)

    monkeypatch.setattr(crm_nt, "_build_dynamic_matriz_risco", lambda **_kwargs: matrix)
    with pytest.raises(RuntimeError, match="Total financeiro da farmacia obrigatorio"):
        crm_nt._build_hhi_crm_context(cnpj, None, None, 0, data)
    assert crm_nt._build_hhi_crm_context(
        cnpj, None, None, 100, SimpleNamespace(crms_interesse=[{"nu_prescricoes": 0}])
    ) is None

    five_of_six = SimpleNamespace(crms_interesse=[
        {"id_medico": f"{i}/SP", "nu_prescricoes": 20, "vl_total_prescricoes": 100}
        for i in range(6)
    ])
    result = crm_nt._build_hhi_crm_context(cnpj, date(2024, 1, 1), None, 600, five_of_six)
    assert len(result["top_crms"]) == 5
    assert result["periodo_intervalo"] == "a partir de 01.01.2024"

    ten_of_twelve = SimpleNamespace(crms_interesse=[
        {"id_medico": f"{i}/SP", "nu_prescricoes": 1, "vl_total_prescricoes": 1}
        for i in range(12)
    ])
    result = crm_nt._build_hhi_crm_context(cnpj, None, date(2024, 12, 31), 100, ten_of_twelve)
    assert len(result["top_crms"]) == 10
    assert result["periodo_intervalo"] == "até 31.12.2024"

    many_crms = SimpleNamespace(crms_interesse=[
        {"id_medico": f"{i}/SP", "nu_prescricoes": 1, "vl_total_prescricoes": 1}
        for i in range(101)
    ])
    result = crm_nt._build_hhi_crm_context(cnpj, None, None, 101, many_crms)
    assert len(result["top_crms"]) == 10
    assert result["periodo_intervalo"] == "no período analisado"


def test_irregular_crm_context_handles_matrix_and_crm_failures(monkeypatch):
    cnpj = "12345678000195"
    valid_matrix = pl.DataFrame({
        "cnpj": [cnpj], "pct_crms_irregulares": [10.0], "crms_irregulares_valor_total": [1000.0],
        "crms_irregulares_valor": [100.0], "risco_crms_irregulares_reg": ["bad"],
        "risco_crms_irregulares_uf": [2.0], "risco_crms_irregulares_br": [3.0],
    })
    monkeypatch.setattr(crm_nt, "_build_dynamic_matriz_risco", lambda **_kwargs: valid_matrix)
    monkeypatch.setattr(
        crm_nt,
        "get_crm_data",
        lambda *_args: (_ for _ in ()).throw(RuntimeError("crm offline")),
    )
    result = crm_nt._build_crms_irregulares_context(cnpj, date(2024, 1, 1), None)
    assert result["periodo_intervalo"] == "a partir de 01.01.2024"
    assert result["total_autorizacoes"] == 0
    assert result["top_irregulares"] == []
    assert result["multiplicador_regiao"] == 0

    data = SimpleNamespace(crms_interesse=[])
    monkeypatch.setattr(crm_nt, "_build_dynamic_matriz_risco", lambda **_kwargs: pl.DataFrame({"cnpj": [cnpj]}))
    with pytest.raises(RuntimeError, match="colunas obrigatorias"):
        crm_nt._build_crms_irregulares_context(cnpj, None, None, data)
    monkeypatch.setattr(crm_nt, "_build_dynamic_matriz_risco", lambda **_kwargs: valid_matrix.with_columns(pl.lit("99999999000191").alias("cnpj")))
    with pytest.raises(RuntimeError, match="sem registro obrigatorio"):
        crm_nt._build_crms_irregulares_context(cnpj, None, None, data)

    for key in ("crms_irregulares_valor_total", "crms_irregulares_valor", "pct_crms_irregulares"):
        invalid = valid_matrix.with_columns(pl.lit(-1.0).alias(key))
        monkeypatch.setattr(crm_nt, "_build_dynamic_matriz_risco", lambda **_kwargs: invalid)
        with pytest.raises(ValueError, match="ausente ou invalido"):
            crm_nt._build_crms_irregulares_context(cnpj, None, None, data)

    valid_crms = SimpleNamespace(crms_interesse=[{
        "id_medico": "123/SP", "nu_prescricoes": 5, "vl_total_prescricoes": 50,
        "flag_crm_invalido": 1, "flag_prescricao_antes_registro": 1,
    }])
    monkeypatch.setattr(crm_nt, "_build_dynamic_matriz_risco", lambda **_kwargs: valid_matrix)
    result = crm_nt._build_crms_irregulares_context(cnpj, date(2024, 1, 1), date(2024, 12, 31), valid_crms)
    assert result["periodo_intervalo"] == "de 01.01.2024 a 31.12.2024"
    assert result["qtd_invalidos"] == result["qtd_antes_registro"] == 1

    doc = Document()
    no_details = {**result, "top_irregulares": []}
    crm_nt._add_crms_irregulares_text(doc, "7.1", "Farmácia", cnpj, no_details, 1)
    assert len(doc.tables) == 0


def test_crms_irregulares_context_and_document_rendering(monkeypatch):
    cnpj = "12345678000195"
    matrix = pl.DataFrame({
        "cnpj": [cnpj], "pct_crms_irregulares": [25.0], "crms_irregulares_valor_total": [1000.0],
        "crms_irregulares_valor": [250.0], "risco_crms_irregulares_reg": [1.0],
        "risco_crms_irregulares_uf": [2.0], "risco_crms_irregulares_br": [3.0],
    })
    monkeypatch.setattr(crm_nt, "_build_dynamic_matriz_risco", lambda **_kwargs: matrix)
    data = SimpleNamespace(crms_interesse=[
        {"id_medico": "100/SP", "no_medico": "ana maria", "nu_prescricoes": 5, "vl_total_prescricoes": 200, "flag_crm_invalido": 1},
        {"id_medico": "200/RJ", "nu_prescricoes": 4, "vl_total_prescricoes": 100, "flag_prescricao_antes_registro": 1},
        {"id_medico": "300/SP", "nu_prescricoes": 10, "vl_total_prescricoes": 500},
    ])
    result = crm_nt._build_crms_irregulares_context(cnpj, None, date(2024, 12, 31), data)
    assert result["periodo_intervalo"] == "até 31.12.2024"
    assert result["total_autorizacoes"] == 19
    assert result["qtd_invalidos"] == 1
    assert result["qtd_antes_registro"] == 1
    no_period = crm_nt._build_crms_irregulares_context(cnpj, None, None, data)
    assert no_period["periodo_intervalo"] == "no período analisado"
    doc = Document()
    crm_nt._add_crms_irregulares_text(doc, "7.2", "Farmácia Teste", cnpj, result, 8, "crm-irregular")
    assert len(doc.tables) == 1
    assert "R$ 250,00" in " ".join(p.text for p in doc.paragraphs)
    assert "CRM não localizado" in doc.tables[0].rows[1].cells[3].text

    with pytest.raises(RuntimeError, match="colunas obrigatorias"):
        monkeypatch.setattr(crm_nt, "_build_dynamic_matriz_risco", lambda **_kwargs: pl.DataFrame({"cnpj": [cnpj]}))
        crm_nt._build_crms_irregulares_context(cnpj, None, None, data)


def test_complementary_evidence_context_and_all_docx_renderers(monkeypatch):
    monkeypatch.setattr(crm_nt, "_load_crm_evidencias_detalhadas", lambda *_args: ([], {}, {}))
    monkeypatch.setattr(crm_nt, "_enrich_crm_unico_valores", lambda _cnpj, _rows: None)
    monkeypatch.setattr(crm_nt, "_enrich_crm_multiplo_valores", lambda _cnpj, _rows: None)
    medico = {
        "id_medico": "123/SP", "no_medico": "ana maria", "nu_prescricoes": 50, "vl_total_prescricoes": 500,
        "pct_participacao": 50, "nu_prescricoes_dia": crm_nt.CRM_DAILY_RATE_ALERT_THRESHOLD + 1,
        "prescricoes_dia_total_brasil": crm_nt.CRM_DAILY_RATE_ALERT_THRESHOLD + 2,
        "nu_estabelecimentos": 4, "competencia_nu_estabelecimentos": "202401", "flag_robo": 1,
        "alertas_crm_unico": [{"dt": "2024-02-03", "nu_prescricoes": 40, "nu_minutos": 15,
            "nu_minutos_intervalo": 10, "taxa_hora": 240, "id_severidade": 2,
            "dt_ini_hora": "2024-02-03 10:00", "dt_fim_hora": "2024-02-03 10:10"}],
        "alertas_crm_multiplos": [{"dt": "2024-02-03", "hr": 11, "nu_presc_crm": 30,
            "nu_presc_total": 50, "nu_crms_total": 2, "nu_minutos": 5, "taxa_hora": 600,
            "dt_ini_hora": "2024-02-03 11:00", "dt_fim_hora": "2024-02-03 11:05", "id_severidade": 1}],
    }
    data = SimpleNamespace(crms_interesse=[medico], cnpj_alerts=[
        {"tipo": "VOLUME", "dt": "2024-02-03", "hr": 12, "nu_prescricoes": 60, "nu_crms": 2, "mediana_hora": 0},
        {"tipo": "MULTIPLO", "dt": "2024-02-03", "hr": 11, "nu_prescricoes": 50, "nu_crms": 2,
         "nu_minutos": 5, "taxa_hora": 600, "dt_ini_hora": "2024-02-03 11:00", "dt_fim_hora": "2024-02-03 11:05"},
    ])
    context = crm_nt._build_crm_evidencias_complementares_context("123", None, None, data)
    assert context["intensiva"]["qtd_local"] == 1
    assert context["volume_horario"]["maior_multiplicador"] == 60
    assert context["crm_unico"]["qtd_medicos"] == 1
    assert context["crms_multiplos"]["qtd_surtos"] == 1
    assert context["principais_crms_contexto"]["rows"][0]["competencia_nu_estabelecimentos"] == "01/2024"

    doc = Document()
    next_table = crm_nt._add_crm_evidencias_complementares_body(doc, "Farmácia Teste", context, 10)
    assert next_table > 10
    assert len(doc.tables) >= 4

    # Each renderer also handles an empty detail set after writing its narrative.
    empty_doc = Document()
    crm_nt._add_crm_intensiva_complementar_text(empty_doc, "a", "Farmácia", {"qtd_medicos": 0}, 1)
    crm_nt._add_crm_unico_complementar_text(empty_doc, "b", "Farmácia", {"qtd_alertas": 0}, 1)
    crm_nt._add_crms_multiplos_complementar_text(empty_doc, "c", "Farmácia", {"qtd_surtos": 0}, 1)
    crm_nt._add_crm_volume_horario_complementar_text(empty_doc, "d", "Farmácia", {"qtd_alertas": 0}, 1)
    crm_nt._add_principais_crms_contexto_text(empty_doc, "Farmácia", {"rows": []}, 1)
    assert len(empty_doc.tables) == 0

    crm_nt._add_crm_unico_complementar_text(empty_doc, "e", "Farmácia", {
        "qtd_medicos": 1,
        "qtd_alertas": 1,
        "maior_qtd": 1,
        "rows": [{
            "id_medico": "123/SP", "nu_prescricoes": 1, "nu_minutos_intervalo": None,
            "nu_minutos": 5, "taxa_hora": 12, "dt": "2024-02-03",
            "valor_alerta_disponivel": False,
        }],
    }, 2)
    crm_nt._add_crms_multiplos_complementar_text(empty_doc, "f", "Farmácia", {
        "qtd_medicos": 2, "qtd_surtos": 1, "maior_qtd": 1, "eventos": [],
    }, 3)
    assert len(empty_doc.tables) == 1


def test_crm_enrichment_normalizes_reversed_windows_and_skips_unusable_rows(tmp_path, monkeypatch):
    missing_cache = tmp_path / "missing"
    monkeypatch.setattr(crm_nt, "_get_cnpj_cache_dir", lambda _cnpj: str(missing_cache))

    unico = [
        {"id_medico": "1/SP", "dt_ini_hora": "2024-02-03 10:10", "dt_fim_hora": "2024-02-03 10:00"},
        {"id_medico": "", "dt_ini_hora": None, "dt_fim_hora": None},
    ]
    crm_nt._enrich_crm_unico_valores("123", unico)
    assert all(row["valor_alerta"] == 0 for row in unico)
    assert all(row["valor_alerta_disponivel"] is False for row in unico)
    crm_nt._enrich_crm_unico_valores("123", [{"id_medico": "", "dt_ini_hora": None, "dt_fim_hora": None}])

    multiplo = [
        {"dt_ini_hora": "2024-02-03 10:10", "dt_fim_hora": "2024-02-03 10:00"},
        {"dt_ini_hora": None, "dt_fim_hora": "2024-02-03 10:00"},
    ]
    crm_nt._enrich_crm_multiplo_valores("123", multiplo)
    assert all(row["valor_alerta"] == 0 for row in multiplo)
    assert all(row["valor_alerta_disponivel"] is False for row in multiplo)
    crm_nt._enrich_crm_multiplo_valores("123", [{"dt_ini_hora": None, "dt_fim_hora": None}])


def test_complementary_crm_evidence_derives_best_multiple_event(monkeypatch):
    monkeypatch.setattr(crm_nt, "_load_crm_evidencias_detalhadas", lambda *_args: ([{
        "tipo": "VOLUME", "dt": "2024-02-03", "hr": 10,
        "nu_prescricoes": 4, "nu_crms": 2, "mediana_hora": 2,
    }], {}, {}))
    monkeypatch.setattr(crm_nt, "_enrich_crm_unico_valores", lambda _cnpj, _rows: None)
    monkeypatch.setattr(crm_nt, "_enrich_crm_multiplo_valores", lambda _cnpj, _rows: None)

    alerts = [
        {
            "dt": "2024-02-03", "hr": 11, "nu_presc_crm": 12, "nu_presc_total": 15,
            "nu_crms_total": 3, "nu_minutos": 5, "taxa_hora": 100,
            "dt_ini_hora": "2024-02-03 11:00", "dt_fim_hora": "2024-02-03 11:05",
        },
        {
            "dt": "2024-02-03", "hr": 11, "nu_presc_crm": 8, "nu_presc_total": 10,
            "nu_crms_total": 2, "nu_minutos": 4, "taxa_hora": 200,
            "dt_ini_hora": "2024-02-03 11:00", "dt_fim_hora": "2024-02-03 11:04",
        },
        {
            "dt": "2024-02-03", "hr": 11, "nu_presc_crm": 2, "nu_presc_total": 8,
            "nu_crms_total": 2, "nu_minutos": 3, "taxa_hora": 50,
            "dt_ini_hora": "2024-02-03 11:00", "dt_fim_hora": "2024-02-03 11:03",
        },
        {"nu_presc_total": 0},
    ]
    data = SimpleNamespace(
        crms_interesse=[
            {
                "id_medico": "123/SP", "no_medico": "ana maria", "nu_prescricoes": 20,
                "vl_total_prescricoes": 200, "nu_estabelecimentos": 1,
                "competencia_nu_estabelecimentos": "202401", "alertas_crm_multiplos": alerts,
                "alertas_crm_unico": [{"nu_prescricoes": 0}],
            },
            {
                "id_medico": "456/RJ", "nu_prescricoes": 2, "vl_total_prescricoes": 20,
                "nu_prescricoes_dia": crm_nt.CRM_DAILY_RATE_ALERT_THRESHOLD + 1,
                "prescricoes_dia_total_brasil": 0, "nu_estabelecimentos": 1,
                "competencia_nu_estabelecimentos": "202401",
            },
            {
                "id_medico": "789/BA", "nu_prescricoes": 2, "vl_total_prescricoes": 20,
                "nu_prescricoes_dia": 0,
                "prescricoes_dia_total_brasil": crm_nt.CRM_DAILY_RATE_ALERT_THRESHOLD + 1,
                "nu_estabelecimentos": 1, "competencia_nu_estabelecimentos": "202401",
            },
        ],
        cnpj_alerts=[],
    )

    context = crm_nt._build_crm_evidencias_complementares_context("123", None, None, data)

    assert context["crms_multiplos"]["qtd_surtos"] == 1
    assert context["crms_multiplos"]["eventos"][0]["taxa_hora"] == 200
    assert context["crms_multiplos"]["eventos"][0]["nu_prescricoes"] == 10
    assert context["volume_horario"]["qtd_alertas"] == 1
    assert context["intensiva"]["qtd_local"] == 1
    assert context["intensiva"]["qtd_brasil"] == 1


def test_complementary_crm_evidence_returns_none_when_unavailable_or_without_alerts(monkeypatch):
    monkeypatch.setattr(
        crm_nt,
        "get_crm_data",
        lambda *_args: (_ for _ in ()).throw(RuntimeError("crm unavailable")),
    )
    assert crm_nt._build_crm_evidencias_complementares_context("123", None, None) is None

    monkeypatch.setattr(crm_nt, "_load_crm_evidencias_detalhadas", lambda *_args: ([], {}, {}))
    empty = SimpleNamespace(crms_interesse=[], cnpj_alerts=[])
    assert crm_nt._build_crm_evidencias_complementares_context("123", None, None, empty) is None

    no_alerts = SimpleNamespace(
        crms_interesse=[{
            "id_medico": "123/SP", "nu_prescricoes": 0, "vl_total_prescricoes": 0,
            "nu_estabelecimentos": 1, "competencia_nu_estabelecimentos": "202401",
        }],
        cnpj_alerts=[{"tipo": "VOLUME", "nu_prescricoes": 0}],
    )
    assert crm_nt._build_crm_evidencias_complementares_context("123", None, None, no_alerts) is None


def test_complementary_crm_event_without_doctor_detail_uses_event_crm_count(monkeypatch):
    monkeypatch.setattr(crm_nt, "_load_crm_evidencias_detalhadas", lambda *_args: ([], {}, {}))
    monkeypatch.setattr(crm_nt, "_enrich_crm_unico_valores", lambda _cnpj, _rows: None)
    monkeypatch.setattr(crm_nt, "_enrich_crm_multiplo_valores", lambda _cnpj, _rows: None)
    data = SimpleNamespace(
        crms_interesse=[{
            "id_medico": "123/SP", "nu_prescricoes": 1, "vl_total_prescricoes": 10,
            "nu_estabelecimentos": 1, "competencia_nu_estabelecimentos": "202401",
        }],
        cnpj_alerts=[{
            "tipo": "MULTIPLO", "dt": "2024-02-03", "hr": 11,
            "nu_prescricoes": 6, "nu_crms": 3, "nu_minutos": 5, "taxa_hora": 72,
            "dt_ini_hora": "2024-02-03 11:00", "dt_fim_hora": "2024-02-03 11:05",
        }],
    )

    context = crm_nt._build_crm_evidencias_complementares_context("123", None, None, data)

    assert context["crms_multiplos"]["qtd_medicos"] == 3
    assert context["crms_multiplos"]["qtd_surtos"] == 1


def test_crm_volume_renderer_explains_nonzero_median():
    doc = Document()
    crm_nt._add_crm_volume_horario_complementar_text(
        doc,
        "a",
        "Farmácia Teste",
        {
            "qtd_alertas": 1,
            "rows": [{
                "dt": "2024-02-03", "hr": 10, "nu_prescricoes": 4,
                "nu_crms": 2, "mediana_hora": 2,
            }],
        },
        1,
    )
    assert len(doc.tables) == 1
    assert "2,0 vezes" in doc.paragraphs[1].text
    assert "mediana histórica de 2,0 autorizações" in doc.paragraphs[1].text
