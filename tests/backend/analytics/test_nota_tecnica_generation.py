import base64
from datetime import date, datetime

import polars as pl
import pytest
from docx import Document

from backend.api.services.analytics import nota_tecnica as nota


def test_timing_settings_and_log_output(tmp_path, monkeypatch):
    monkeypatch.setattr(nota, "_project_root", lambda: tmp_path)
    monkeypatch.delenv("NOTA_TECNICA_TIMING_LOG", raising=False)
    assert nota._nota_tecnica_timing_enabled() is False

    (tmp_path / ".env").write_text("# comentario\nOUTRA=1\nNOTA_TECNICA_TIMING_LOG='yes'\n", encoding="utf-8")
    assert nota._nota_tecnica_timing_enabled() is True
    monkeypatch.setenv("NOTA_TECNICA_TIMING_LOG", "  OFF ")
    assert nota._nota_tecnica_timing_enabled() is False
    monkeypatch.setenv("NOTA_TECNICA_TIMING_LOG", " On ")
    assert nota._nota_tecnica_timing_enabled() is True

    timing = nota._NotaTecnicaTiming("123", date(2024, 1, 1), None)
    timing.mark("carregou dados")
    timing.write(status="ERRO", error="fonte indisponivel")
    log = (tmp_path / "logs" / "nota_tecnica_timing.log").read_text(encoding="utf-8")
    assert "CNPJ 123" in log and "2024-01-01 a fim-aberto" in log
    assert "carregou dados" in log and "Erro: fonte indisponivel" in log


def test_nota_tecnica_value_and_identifier_formatters():
    assert nota._risk_color("CRÍTICO", 0) == ("F87171", "CRÍTICO")
    assert nota._risk_color("médio", 0) == ("F97316", "ATENÇÃO")
    assert nota._risk_color("normal", 99) == ("334155", "NORMAL")
    assert nota._risk_color(None, 21) == ("F87171", "CRÍTICO")
    assert nota._risk_color(None, 11) == ("F97316", "ATENÇÃO")
    assert nota._risk_color(None, 10) == ("334155", "NORMAL")
    assert nota._vez_ou_vezes(-1) == "vez"
    assert nota._vez_ou_vezes(2) == "vezes"

    generated = nota._build_codigo_verificacao("12.345/0001-99", datetime(2024, 3, 4, 5, 6))
    assert generated.startswith("NT-12345000199-20240304-") and len(generated.rsplit("-", 1)[1]) == 8
    assert nota._build_codigo_verificacao("", datetime(2024, 3, 4)).startswith("NT-SEM-CNPJ-20240304-")
    assert nota._resolve_numero_nota_input(None) == ("XXX", True)
    assert nota._resolve_numero_nota_input(" 0042 ") == ("0042", False)
    with pytest.raises(ValueError, match="apenas dígitos"):
        nota._resolve_numero_nota_input("42-A")
    assert nota._resolve_numero_processo_input(None, 2024) == ("00XXX.XXXXXX/2024-XX", True)
    assert nota._resolve_numero_processo_input("12345678901234567", 2024) == ("12345.678901/2345-67", False)
    assert nota._resolve_numero_processo_input("12345.678901/2024-67", 2024) == ("12345.678901/2024-67", False)
    with pytest.raises(ValueError, match="17 dígitos"):
        nota._resolve_numero_processo_input("123", 2024)


def test_nota_tecnica_signatories_are_normalized_and_validated():
    assert nota._resolve_assinantes_tecnicos(None) == [
        {"nome": "Fulano de Tal", "cargo": "Cargo"}, {"nome": "Cicrano de Tal", "cargo": "Cargo"}
    ]
    assert nota._resolve_assinantes_tecnicos([{"nome": " Ana ", "cargo": " Auditora "}]) == [
        {"nome": "Ana", "cargo": "Auditora"}
    ]
    assert nota._resolve_assinantes_tecnicos([])[0]["nome"] == "Fulano de Tal"
    assert nota._resolve_assinantes_tecnicos([{}])[0]["nome"] == "Fulano de Tal"
    for invalid in ("Ana", ["Ana"], [{"nome": "Ana"}], [{"cargo": "Auditora"}], [{}, {}, {}, {}]):
        with pytest.raises(ValueError):
            nota._resolve_assinantes_tecnicos(invalid)


def test_project_root_resolves_repository_directory():
    assert nota._project_root().name == "sentinela"


def test_nota_tecnica_docx_section_footer_watermark_and_headers(tmp_path):
    doc = Document()
    (tmp_path / "brasao.png").write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+j1ioAAAAASUVORK5CYII="
    ))
    nota._add_sumario_official_header(
        doc, str(tmp_path / "brasao.png"),
        {"nome_unidade": "CGU", "linha_endereco": "Brasília", "linha_contato": "contato", "codigo": "SP"},
        2024, "XXX", True, "00XXX.XXXXXX/2024-XX", True,
    )
    assert "NOTA TÉCNICA Nº" in " ".join(p.text for p in doc.paragraphs)
    with pytest.raises(FileNotFoundError):
        nota._add_sumario_official_header(
            doc, str(tmp_path / "ausente.png"), {"nome_unidade": "CGU"}, 2024, "1", False, "123", False
        )

    section = nota._start_section(doc, footer_lines=["linha um", "linha dois"])
    assert len(doc.sections) == 2
    assert section.footer.paragraphs[0].text == "linha um\nlinha dois"
    nota._apply_codigo_verificacao_footer(doc, "NT-123")
    assert "NT-123" in doc.sections[-1].footer.paragraphs[0].text
    heading = nota._format_main_heading(doc.add_heading("Seção", level=1))
    assert heading.paragraph_format.keep_with_next is True
    nota._add_confidential_watermark(section)
    assert b"ACESSO RESTRITO" in section.header._element.xml.encode()


def test_nota_tecnica_localidade_resolution_and_summary_helpers(monkeypatch):
    assert nota._valor_estimado_por_percentual(1000, 25) == 250
    assert nota._valor_estimado_por_percentual(0, 25) == 0
    assert nota._valor_estimado_por_percentual(1000, 0) == 0
    assert nota._normalize_id_ibge7(123.0) == "0000123"
    assert nota._normalize_id_ibge7("1234567.0") == "1234567"
    assert nota._normalize_id_ibge7("X") == "X"
    assert nota._resolve_unidade_pf({}) == "Delegacia de Polícia Federal competente"

    monkeypatch.setattr(nota, "get_localidades_df", lambda: pl.DataFrame(
        {"id_ibge7": [1234567], "unidade_pf": [" Delegacia PF. "]}
    ))
    assert nota._resolve_unidade_pf({"id_ibge7": "1234567.0"}) == "Delegacia PF"
    monkeypatch.setattr(nota, "get_localidades_df", lambda: pl.DataFrame({"municipio": ["X"]}))
    assert nota._resolve_unidade_pf({"id_ibge7": 1234567}) == "Delegacia de Polícia Federal competente"
    monkeypatch.setattr(nota, "get_localidades_df", lambda: pl.DataFrame(
        {"id_ibge7": [7654321], "unidade_pf": ["Delegacia de Outra Cidade"]}
    ))
    assert nota._resolve_unidade_pf({"id_ibge7": 1234567}) == "Delegacia de Polícia Federal competente"
    monkeypatch.setattr(nota, "get_localidades_df", lambda: (_ for _ in ()).throw(RuntimeError("offline")))
    assert nota._resolve_unidade_pf({"id_ibge7": 1234567}) == "Delegacia de Polícia Federal competente"

    assert nota._build_resumo_falecidos("7.1", {
        "periodo_desc": "2020 a 2024.", "total_autorizacoes": 1200, "cpfs_distintos": 35, "valor_total": 1234.5,
    }).startswith("[Subitem 7.1]: Registros, 2020 a 2024, de 1.200 vendas")
    doc = Document()
    nota._add_resumo_criticidades_conclusao(doc, [])
    assert doc.paragraphs == []
    nota._add_resumo_criticidades_conclusao(doc, ["Resumo de criticidade"])
    assert "Resumo de criticidade" in " ".join(p.text for p in doc.paragraphs)


@pytest.mark.parametrize(
    ("key", "comp", "expected"),
    [
        ("teto", {"valor_suspeito": 20.0, "multiplicador_regiao": 2.0}, "R$ 20,00"),
        ("polimedicamento", {"percentual": 10.0, "multiplicador_regiao": 1.5}, "R$ 100,00"),
        ("alto_custo", {"percentual": 10.0, "multiplicador_regiao": 1.5}, "R$ 100,00"),
        ("vendas_rapidas", {"percentual": 10.0, "multiplicador_regiao": 1.5}, "R$ 100,00"),
        ("recorrencia_sistemica", {"percentual": 10.0, "multiplicador_regiao": 1.5}, "R$ 100,00"),
        ("dias_pico", {"percentual": 10.0, "multiplicador_regiao": 1.5}, "R$ 100,00"),
        ("ticket_medio", {"valor": 5.5, "multiplicador_regiao": 1.5}, "R$ 5,50"),
        ("receita_paciente", {"valor": 5.5, "multiplicador_regiao": 1.5}, "R$ 5,50"),
        ("per_capita", {"valor": 5.5, "multiplicador_regiao": 1.5}, "R$ 5,50"),
        ("hhi_crm", {"principal": {"id_medico": "CRM/SP"}, "principal_autorizacoes": 1234,
                      "principal_valor": 20.0, "pct_valor": 15.0}, "1.234 autorizações"),
        ("crms_irregulares", {"pct_irregular": 15.0, "valor_irregular": 10.0}, "15,00%"),
        ("dispersao_geografica", {"percentual_financeiro_outra_uf": 3.0, "total_valor_outra_uf": 50.0}, "R$ 50,00"),
    ],
)
def test_nota_tecnica_criticidade_summaries(key, comp, expected):
    result = nota._build_resumo_criticidade("7.1", key, comp, total_mov=1000)
    assert expected in result


def test_nota_tecnica_criticidade_summary_requires_valid_clinical_value():
    assert "12,50%" in nota._build_resumo_criticidade(
        "7.1", "incompatibilidade_patologica", {"percentual": 12.5, "multiplicador_regiao": 2.0}, 100
    )
    with pytest.raises(RuntimeError, match="Valor clinico total obrigatorio"):
        nota._build_resumo_criticidade(
            "7.1", "incompatibilidade_patologica", {"percentual": 12.5, "ranking_patologias": ["A"]}, 100
        )
    with pytest.raises(RuntimeError, match="positivo e finito"):
        nota._build_resumo_criticidade(
            "7.1", "incompatibilidade_patologica",
            {"percentual": 12.5, "ranking_patologias": ["A"], "valor_suspeito": float("inf")}, 100,
        )
    assert "valor total identificado de R$ 25,00" in nota._build_resumo_criticidade(
        "7.1", "incompatibilidade_patologica",
        {"percentual": 12.5, "multiplicador_regiao": 2.0, "ranking_patologias": ["A"], "valor_suspeito": 25},
        100,
    )
    assert nota._build_resumo_criticidade("7.1", "indicador_desconhecido", {}, 100) is None


def test_nota_tecnica_critical_items_numbering_and_word_contents():
    first = nota._iter_criticidade_items({"teto", "ticket_medio"}, "Farmácia Alfa", ordered_keys=["ticket_medio", "teto"])
    assert [item[:2] for item in first] == [("ticket_medio", "7.1"), ("teto", "7.2")]
    remaining = nota._iter_criticidade_items(
        {"teto", "ticket_medio"}, "Farmácia Alfa", start_index=2,
        exclude_keys={"ticket_medio"}, ordered_keys=["ticket_medio", "teto", "desconhecido"],
    )
    assert len(remaining) == 1 and remaining[0][:2] == ("teto", "7.2")
    assert "Farmácia Alfa" in remaining[0][2]
    assert nota._iter_criticidade_items(
        {"desconhecido"}, "Farmácia Alfa", ordered_keys=["desconhecido"]
    ) == []

    doc = Document()
    nota._add_toc_entry(doc, "  7.1", "Análise", "8")
    assert "7.1 Análise" in doc.paragraphs[0].text and doc.paragraphs[0].paragraph_format.left_indent.inches == 0.4
    nota._build_sumario(doc, {"teto"}, "Farmácia Alfa", "00.000.000/0001-91", ["teto"])
    assert any("SUMÁRIO" in paragraph.text for paragraph in doc.paragraphs)
    assert len(doc.tables) == 0


def test_generate_nota_tecnica_runs_all_sections_and_criticality_branches(tmp_path, monkeypatch):
    class Payload:
        def __init__(self, values):
            self._values = values
            self.porte_empresa = values.get("porte_empresa", "DEMAIS")
            self.situacao_rf = values.get("situacao_rf", "ATIVA")

        def model_dump(self):
            return dict(self._values)

    cnpj = "12345678000195"
    cadastro = Payload(
        {
            "cnpj": cnpj, "razao_social": "Farmacia Teste", "nome_fantasia": "Teste", "municipio": "Brasilia",
            "uf": "DF", "logradouro": "Rua A", "tipo_logradouro": "RUA", "numero": "10", "bairro": "Centro",
            "cep": "70000000", "capital_social": 100000, "is_matriz": False,
            "id_cnae_principal": 4771701, "cnae_principal": "Farmacia", "cnaes_secundarios": [],
        }
    )
    cnpj_data = Payload(
        {
            "cnpj": cnpj, "razao_social": "Farmacia Teste", "municipio": "Brasilia", "uf": "DF",
            "percValSemComp": 25.0, "score_risco_final": 30, "classificacao_risco": "CRITICO",
            "is_conexao_ativa": True, "porte_empresa": "DEMAIS", "situacao_rf": "ATIVA",
            "totalMov": 1000.0, "valSemComp": 250.0,
        }
    )
    all_keys = [key for key, _, _ in nota._SECAO5_MAP]
    monkeypatch.setattr(nota, "_nota_tecnica_timing_enabled", lambda: False)
    monkeypatch.setattr(nota, "get_dados_farmacia", lambda _: cadastro)
    monkeypatch.setattr(nota, "get_dashboard_data", lambda *args, **kwargs: type(
        "Resumo", (), {"resultado_cnpjs": [cnpj_data]}
    )())
    monkeypatch.setattr(nota, "get_socios_farmacia", lambda _: type("Socios", (), {"socios": []})())
    monkeypatch.setattr(nota, "get_crm_data", lambda *args, **kwargs: {"crm": "loaded"})
    monkeypatch.setattr(nota, "resolve_nota_tecnica_regional", lambda _: {
        "codigo": "DF", "estado": "Distrito Federal", "cidade_uf": "Brasília/DF",
        "nome_unidade": "CGU Regional DF", "linha_endereco": "Brasília", "linha_contato": "Contato",
        "superintendente": "Responsável", "cargo_superintendente": "Superintendente",
    })
    monkeypatch.setattr(nota, "_resolve_brasao_republica_path", lambda: str(tmp_path / "brasao.png"))
    (tmp_path / "brasao.png").write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+j1ioAAAAASUVORK5CYII="
    ))
    monkeypatch.setattr(nota, "_get_criticos", lambda *args: set(all_keys))
    monkeypatch.setattr(nota, "_get_criticos_ordenados_por_risco", lambda _cnpj, _keys, *_: list(all_keys))
    monkeypatch.setattr(nota, "_build_falecidos_context", lambda *args: {
        "periodo_desc": "2020 a 2024", "total_autorizacoes": 3, "cpfs_distintos": 2, "valor_total": 50.0,
    })
    monkeypatch.setattr(nota, "_build_anexo_ii_context", lambda *args: {"rows": []})
    monkeypatch.setattr(nota, "_build_crm_evidencias_complementares_context", lambda *args, **kwargs: {"evidencias": []})
    monkeypatch.setattr(nota, "_build_ultimo_mes_sav_context", lambda *args: {"total": 100, "mes_formatado": "jan/2024"})
    monkeypatch.setattr(nota, "_build_esocial_context", lambda *args: {})
    monkeypatch.setattr(nota, "_add_esocial_context_text", lambda _doc, *_args: 0)
    monkeypatch.setattr(nota, "_build_regional_comparison_context", lambda *args: {
        "multiplicador": 2.0, "multiplicador_uf": 1.5, "multiplicador_brasil": 1.2,
        "qtd_farmacias": 2, "municipios": ["Brasilia", "Gama"], "uf": "DF",
    })
    monkeypatch.setattr(nota, "_build_posicionamento_regional_context", lambda *args: {})
    monkeypatch.setattr(nota, "_build_gtin_sem_comprovacao_context", lambda *args: {
        "total_gtins": 2, "representativos_count": 1, "total_valor": 250.0,
        "representativos_valor": 200.0, "representativos_pct": 80.0,
    })
    monkeypatch.setattr(nota, "_build_evolucao_financeira_context", lambda *args: {
        "rows": [], "semestres_atipicos": [{
            "semestre_fmt": "2024-S1", "chave_semestre": 202401, "aumento_valor_semestre": 100.0,
            "taxa_crescimento_pct": 10.0,
        }], "total": 1000.0, "periodo_meses": "2024", "top_irregulares": [
            {"chave_semestre": 202302, "semestre_fmt": "2023-S2", "irregular": 200.0}
        ], "medicamentos_aumento_atipico": [{"gtin": "1"}],
    })
    monkeypatch.setattr(nota, "_build_socios_volume_atipico_context", lambda *args: [{"socio": "A"}])
    monkeypatch.setattr(nota, "_build_repasses_anuais_context", lambda *args: {
        "sem_repasses": True, "total": 0.0, "anos": [],
    })

    monkeypatch.setattr(nota, "_build_indicadores_criticos_quadro", lambda *args: [{"indicador": "teto"}])
    monkeypatch.setattr(nota, "_build_indicador_regional_context", lambda *args: {"rows": []})
    monkeypatch.setattr(nota, "_build_incompatibilidade_patologica_context", lambda *args: {"unavailable": True})
    for name in (
        "_build_teto_context", "_build_polimedicamento_context", "_build_ticket_medio_context",
        "_build_receita_paciente_context", "_build_per_capita_context", "_build_alto_custo_context",
        "_build_vendas_rapidas_context", "_build_recorrencia_sistemica_context", "_build_dias_pico_context",
        "_build_dispersao_geografica_context",
    ):
        monkeypatch.setattr(nota, name, lambda *args: {
            "percentual": 10.0, "multiplicador_regiao": 2.0, "valor_suspeito": 100.0, "valor": 10.0,
            "percentual_financeiro_outra_uf": 3.0, "total_valor_outra_uf": 30.0,
        })
    monkeypatch.setattr(nota, "_build_hhi_crm_context", lambda *args, **kwargs: {
        "principal": {"id_medico": "CRM/DF"}, "principal_autorizacoes": 10,
        "principal_valor": 100.0, "pct_valor": 10.0,
    })
    monkeypatch.setattr(nota, "_build_crms_irregulares_context", lambda *args, **kwargs: {
        "top_irregulares": [], "pct_irregular": 10.0, "valor_irregular": 100.0,
    })
    monkeypatch.setattr(nota, "_count_incompatibilidade_patologica_tables", lambda _: 1)
    monkeypatch.setattr(nota, "_build_resumo_criticidade", lambda *_: "Resumo de criticidade")

    noop_names = (
        "_add_sumario_official_header", "_build_sumario", "_add_quadro_identificacao", "_add_quadro_53",
        "_add_quadro_comparativo_regional", "_add_figura_posicionamento_regional", "_add_tabela_gtins_sem_comprovacao",
        "_add_quadro_evolucao_financeira", "_add_tabela_medicamentos_aumento_atipico", "_add_quadro_socios_volume_atipico",
        "_add_figura_evolucao_financeira", "_add_tabela_repasses_anuais", "_add_indicadores_criticos_quadro",
        "_add_indicador_regional_table", "_add_falecidos_criticidade_text", "_add_incompatibilidade_patologica_text",
        "_add_parkinson_gtin_sem_comprovacao_text", "_add_teto_text", "_add_polimedicamento_text",
        "_add_ticket_medio_text", "_add_receita_paciente_text", "_add_per_capita_text", "_add_alto_custo_text",
        "_add_vendas_rapidas_text", "_add_recorrencia_sistemica_text", "_add_dias_pico_text",
        "_add_dispersao_geografica_text", "_add_dispersao_geografica_regional_text", "_add_hhi_crm_text",
        "_add_crms_irregulares_text", "_add_anexo_crm_evidencias", "_add_anexo_ii_memoria_calculo", "_add_anexo_falecidos",
    )
    for name in noop_names:
        monkeypatch.setattr(nota, name, lambda *args, **kwargs: None)

    document_type = type(Document())
    styles_property = document_type.styles
    style_access = {"failed_once": False}

    class FlakyStyles:
        def __init__(self, styles):
            self._styles = styles

        def __getitem__(self, name):
            if name.startswith("Heading") and not style_access["failed_once"]:
                style_access["failed_once"] = True
                raise RuntimeError("heading styles unavailable")
            return self._styles[name]

    def styles_with_one_heading_failure(document):
        return FlakyStyles(styles_property.fget(document))

    with monkeypatch.context() as style_failure:
        style_failure.setattr(document_type, "styles", property(styles_with_one_heading_failure))
        output = nota.generate_nota_tecnica(
            db=None, cnpj=cnpj, numero_nota="42", numero_processo="12345.678901/2024-67",
            assinantes_tecnicos=[{"nome": "Auditora Teste", "cargo": "Auditora Federal"}],
        )
    assert style_access["failed_once"] is True

    assert output.read(2) == b"PK"
    assert output.tell() == 2
    output.seek(0)
    generated_doc = Document(output)
    text = "\n".join(paragraph.text for paragraph in generated_doc.paragraphs)
    assert "1. ASSUNTO" in text and "5. SOBRE A FARMÁCIA" in text
    assert "8. CONCLUSÃO E ENCAMINHAMENTO" in text
    assert "Auditora Teste" in text and "Responsável" in text
    assert len(generated_doc.sections) >= 6

    def render_variant(*, inicio=None, fim=None):
        variant = nota.generate_nota_tecnica(
            db=None,
            cnpj=cnpj,
            data_inicio=inicio,
            data_fim=fim,
            numero_nota="42",
            numero_processo="12345.678901/2024-67",
            assinantes_tecnicos=[{"nome": "Auditora Teste", "cargo": "Auditora Federal"}],
        )
        assert variant.read(2) == b"PK"

    cadastro._values["nome_fantasia"] = "null"
    cadastro._values["capital_social"] = 0
    cnpj_data.porte_empresa = "Microempresa"
    monkeypatch.setattr(nota, "_build_incompatibilidade_patologica_context", lambda *args: {"rows": []})
    monkeypatch.setattr(nota, "_build_teto_context", lambda *args: None)
    monkeypatch.setattr(nota, "_build_crms_irregulares_context", lambda *args, **kwargs: {
        "top_irregulares": [{"id_medico": "1/SP"}], "pct_irregular": 10.0, "valor_irregular": 100.0,
    })
    monkeypatch.setattr(nota, "_build_repasses_anuais_context", lambda *args: {
        "sem_repasses": False, "total": 500.0, "anos": [2024],
    })
    render_variant(inicio=date(2024, 1, 1), fim=date(2024, 12, 31))

    cadastro._values["capital_social"] = 100000
    cnpj_data.porte_empresa = "Empresa de pequeno porte"
    render_variant(inicio=date(2024, 1, 1))
    cnpj_data.porte_empresa = "Empresa de médio porte"
    render_variant(fim=date(2024, 12, 31))
    cnpj_data.porte_empresa = "Empresa de grande porte"
    render_variant()

    monkeypatch.setattr(nota, "_get_criticos", lambda *_args: {"falecidos"})
    monkeypatch.setattr(nota, "_get_criticos_ordenados_por_risco", lambda *_args: ["falecidos"])
    monkeypatch.setattr(nota, "_build_falecidos_context", lambda *_args: None)
    with pytest.raises(RuntimeError, match="detalhamento de falecidos nao foi encontrado"):
        render_variant()

    monkeypatch.setattr(nota, "_get_criticos", lambda *_args: set())
    monkeypatch.setattr(nota, "_get_criticos_ordenados_por_risco", lambda *_args: [])
    monkeypatch.setattr(
        nota,
        "get_crm_data",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(RuntimeError("CRM offline")),
    )
    monkeypatch.setattr(nota, "_nota_tecnica_timing_enabled", lambda: True)
    monkeypatch.setattr(nota._NotaTecnicaTiming, "write", lambda _self, **_kwargs: None)
    render_variant()
