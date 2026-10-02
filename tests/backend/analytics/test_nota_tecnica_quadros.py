from datetime import date
from decimal import Decimal
from types import SimpleNamespace

from docx import Document

from backend.api.services.analytics import nota_tecnica_quadros as quadros


def test_secondary_cnae_selection_and_required_contract():
    assert quadros._format_cnae_secundario_destacado([]) == "—"
    assert quadros._format_cnae_secundario_destacado(
        [
            {"id_cnae": "6201-5/00", "descricao": "DESENVOLVIMENTO"},
            {"id_cnae": "4771-7/02", "descricao": "FARMÁCIA"},
            {"id_cnae": "4771-7/01", "descricao": "DROGARIA"},
        ]
    ) == "4771701 - DROGARIA"
    assert quadros._format_cnae_secundario_destacado(
        [{"id_cnae": "6201/5", "descricao": "  Software  "}]
    ) == "62015 - Software"

    for invalid in (None, "4771701", ["4771701"], [{}], [{"descricao": "sem código"}], [{"id_cnae": "x"}],
                    [{"id_cnae": "1", "descricao": 2}]):
        try:
            quadros._format_cnae_secundario_destacado(invalid)
        except RuntimeError:
            pass
        else:
            raise AssertionError(f"O contrato deveria rejeitar {invalid!r}")


def test_partnership_regional_and_gtin_tables_render_rows():
    doc = Document()
    quadros._add_quadro_socios_volume_atipico(
        doc,
        [
            {"nome_socio": "Pessoa A", "entrada_txt": "2024-S1", "semestre_fmt": "2024 S1", "taxa_crescimento_pct": 25, "distancia_semestres": 0},
            {"nome_socio": "Pessoa B", "entrada_txt": "2024-S1", "semestre_fmt": "2024 S2", "taxa_crescimento_pct": 10, "distancia_semestres": 1},
            {"nome_socio": "Pessoa B", "entrada_txt": "2024-S1", "semestre_fmt": "2024 S2", "taxa_crescimento_pct": None, "distancia_semestres": 2},
        ],
        2,
    )
    quadros._add_quadro_socios_volume_atipico(doc, [], 3)
    quadros._add_quadro_comparativo_regional(
        doc,
        {"mediana_regional": 12.5, "multiplicador": 2, "qtd_farmacias": 80},
        {"percValSemComp": 25},
        "2020 a 2024",
        4,
    )
    quadros._add_tabela_gtins_sem_comprovacao(
        doc,
        "Farmácia X",
        "12.345.678/0001-95",
        {
            "rows": [
                {"gtin": "123", "descricao": "MEDICAMENTO A", "patologia": "diabetes", "qtd_sem_comprovacao": 1200, "valor_sem_comprovacao": 450.25, "pct_sem_comprovacao": 12.5}
            ],
            "total_qtd": 1200,
            "total_valor": 450.25,
            "total_pct_sem_comprovacao": 12.5,
        },
        "2020 a 2024",
        5,
    )
    assert len(doc.tables) == 3
    assert "1 semestre após a entrada" in " ".join(
        cell.text for row in doc.tables[0].rows for cell in row.cells
    )
    assert "1.200" in " ".join(cell.text for row in doc.tables[-1].rows for cell in row.cells)


def test_financial_evolution_and_atypical_medicine_tables_cover_highlights():
    doc = Document()
    quadros._add_quadro_evolucao_financeira(
        doc,
        "Farmácia X",
        "12.345.678/0001-95",
        {
            "primeiro_semestre": "2020-S1",
            "ultimo_semestre": "2024-S2",
            "primeiro_semestre_fmt": "1º semestre de 2020",
            "ultimo_semestre_fmt": "2º semestre de 2024",
            "rows": [
                {"semestre_fmt": "1º semestre de 2024", "semestre": "2024-S1", "total": 1000, "regular": 200, "irregular": 800, "pct_irregular": 80, "volume_atipico": True, "taxa_crescimento_pct": 65},
                {"semestre_fmt": "2º semestre de 2024", "semestre": "2024-S2", "total": 100, "regular": 100, "irregular": 0, "pct_irregular": 0, "volume_atipico": False, "taxa_crescimento_pct": None},
            ],
            "total": 1100,
            "regular": 300,
            "irregular": 800,
            "pct_irregular": 72.72,
        },
        6,
    )
    quadros._add_tabela_medicamentos_aumento_atipico(
        doc,
        [
            {"semestre_fmt": "2024 S1", "gtin": "1", "descricao": "Produto A", "valor_anterior": 10, "valor_atual": 20, "aumento_valor": 10, "aumento_relativo_pct": 100},
            {"semestre_fmt": "2024 S1", "gtin": "2", "descricao": "Produto B", "valor_anterior": 5, "valor_atual": 15, "aumento_valor": 10, "aumento_relativo_pct": None},
            {"semestre_fmt": "2024 S2", "gtin": "3", "descricao": "Produto C", "valor_anterior": 2, "valor_atual": 8, "aumento_valor": 6, "aumento_relativo_pct": 300},
        ],
        7,
    )
    quadros._add_tabela_medicamentos_aumento_atipico(doc, [], 8)
    assert len(doc.tables) == 2
    assert "tabela_evolucao_financeira" in doc.element.xml
    assert "+100,0%" in " ".join(cell.text for row in doc.tables[1].rows for cell in row.cells)


def test_identification_table_renders_cadastral_and_partner_scenarios():
    doc = Document()
    base = {
        "cnpj_fmt": "12.345.678/0001-95",
        "razao_social": "Farmácia X",
        "nome_fantasia": "Drogaria X",
        "natureza_juridica": "Sociedade Empresária",
        "id_cnae_principal": "4771701",
        "cnae_principal": "Comércio varejista de produtos farmacêuticos",
        "cnaes_secundarios": [{"id_cnae": "4771702", "descricao": "Farmácia"}],
        "data_abertura": date(2000, 1, 2),
        "situacao_rf": "ATIVA",
        "porte_empresa": "DEMAIS",
        "endereco_completo": "Rua A, 10",
        "bairro": "Centro",
        "municipio": "São Paulo",
        "uf": "SP",
        "cep": "01000-000",
        "telefone_1": "1111",
        "telefone_2": "2222",
        "email": "x@example.test",
        "data_processamento": date(2025, 3, 4),
        "socios_ativos": [
            SimpleNamespace(cpf_cnpj_socio="12345678901", data_entrada_sociedade=date(2020, 5, 6), nome_socio="Pessoa X")
        ],
    }
    quadros._add_quadro_identificacao(doc, base, Decimal("100000.50"), "2020 a 2024")
    no_partners = {
        **base,
        "data_abertura": None,
        "data_processamento": None,
        "socios_ativos": [],
        "endereco_completo": None,
        "telefone_1": None,
        "telefone_2": None,
    }
    quadros._add_quadro_identificacao(doc, no_partners, Decimal("0"), "2020 a 2024")
    text = " ".join(p.text for p in doc.paragraphs) + " " + " ".join(
        cell.text for table in doc.tables for row in table.rows for cell in row.cells
    )
    assert "CPF: 123.456.789-01" in text
    assert "Informação de sócios não disponível" in text
    assert "Capital Social" in text


def test_esocial_and_program_summary_tables_render_optional_data():
    doc = Document()
    quadros._add_quadro_esocial(doc, "Farmácia X", "12.345.678/0001-95", {"rows": []}, 9)
    quadros._add_quadro_esocial_trabalhadores(doc, "Farmácia X", "12.345.678/0001-95", {"trabalhador_detalhe_rows": []})
    quadros._add_quadro_esocial(
        doc,
        "Farmácia X",
        "12.345.678/0001-95",
        {"rows": [{"ano_base": 2024, "qtd_trabalhadores_vinculo_ano": 4, "qtd_farmaceuticos_vinculo_ano": 2}], "dt_carga_fonte_txt": "01/2025"},
        9,
    )
    quadros._add_quadro_esocial_trabalhadores(
        doc,
        "Farmácia X",
        "12.345.678/0001-95",
        {
            "trabalhador_detalhe_rows": [
                {"ano_base": 2024, "cpf_trabalhador": "123", "cbo": 223405, "titulo_cbo": "Farmacêutico", "dt_admissao_txt": "01/01/2024", "dt_rescisao_txt": None},
                {"ano_base": 2023, "cpf_trabalhador": None, "cbo": None, "titulo_cbo": None, "dt_admissao_txt": None, "dt_rescisao_txt": "01/02/2023"},
            ],
            "trabalhador_detalhe_modo": "anos_criticos",
        },
    )
    quadros._add_quadro_esocial_trabalhadores(
        doc,
        "Farmácia X",
        "12.345.678/0001-95",
        {"trabalhador_detalhe_rows": [{"ano_base": 2022, "cbo": 1234}]},
    )
    quadros._add_quadro_53(
        doc,
        "Farmácia X",
        "12.345.678/0001-95",
        {"totalMov": 1000.5, "totalQtde": 100, "valSemComp": 200.25, "qtdeSemComp": 20, "percValSemComp": 20, "percQtdeSemComp": 20},
        "2020 a 2024",
        10,
    )
    assert len(doc.tables) == 4
    all_text = " ".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    assert "Farmacêuticos com vínculo" in all_text
    assert "223405" in all_text
    assert "20,00%" in all_text


def test_annual_payment_table_covers_rows_and_no_payment_notice():
    doc = Document()
    common = {"periodo_fmt": "no período analisado"}
    quadros._add_tabela_repasses_anuais(
        doc,
        "Farmácia X",
        "12.345.678/0001-95",
        {**common, "rows": [{"ano": 2023, "valor": 1000}, {"ano": 2024, "valor": 1500}], "total": 2500},
        11,
    )
    quadros._add_tabela_repasses_anuais(
        doc,
        "Farmácia X",
        "12.345.678/0001-95",
        {**common, "rows": [], "total": 0, "sem_repasses": True},
        12,
    )
    assert len(doc.tables) == 2
    text = " ".join(p.text for p in doc.paragraphs)
    assert "possível repasse dos recursos" in text
    assert "tentativa" in text
