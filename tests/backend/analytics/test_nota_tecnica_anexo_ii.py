from datetime import date, datetime
from types import SimpleNamespace

import pytest
from docx import Document

from api.services.analytics import nota_tecnica_anexo_ii as anexo


class Engine:
    def connect(self):
        return "engine"


class Session:
    bind = "session-bind"

    def get_bind(self):
        return "active-bind"


class ModelWithDump:
    def model_dump(self):
        return {"value": 1}


class ModelWithDict:
    def dict(self):
        return {"value": 2}


class Timing:
    def __init__(self):
        self.marks = []

    def mark(self, label):
        self.marks.append(label)


def _row(tipo, **values):
    return {"tipo_linha": tipo, **values}


def test_model_conversion_and_engine_resolution_support_expected_inputs():
    assert anexo._model_to_dict({"value": 0}) == {"value": 0}
    assert anexo._model_to_dict(ModelWithDump()) == {"value": 1}
    assert anexo._model_to_dict(ModelWithDict()) == {"value": 2}
    assert anexo._model_to_dict(object()) == {}
    engine = Engine()
    assert anexo._resolve_engine(engine) is engine
    assert anexo._resolve_engine(Session()) == "active-bind"
    assert anexo._resolve_engine(SimpleNamespace(bind="session-bind")) == "session-bind"
    raw = object()
    assert anexo._resolve_engine(raw) is raw


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (None, None),
        ("", None),
        ("-", None),
        ("—", None),
        (date(2024, 2, 29), (datetime(2024, 2, 29), "29.02.2024")),
        (" 29.02.2024 ", (datetime(2024, 2, 29), "29.02.2024")),
        ("29/02/2024", (datetime(2024, 2, 29), "29.02.2024")),
        ("2024-02-29T00:00:00", (datetime(2024, 2, 29), "29.02.2024")),
        ("not-a-date", None),
    ],
)
def test_parse_period_date_accepts_supported_formats_and_rejects_invalid_values(value, expected):
    assert anexo._parse_period_date(value) == expected


def test_pick_period_boundary_keeps_minimum_or_maximum_and_ignores_invalid_dates():
    jan = anexo._parse_period_date("2024-01-01")
    feb = anexo._parse_period_date("2024-02-01")
    assert anexo._pick_period_boundary(None, "2024-02-01", "min") == feb
    assert anexo._pick_period_boundary(feb, "2024-01-01", "min") == jan
    assert anexo._pick_period_boundary(jan, "2024-02-01", "max") == feb
    assert anexo._pick_period_boundary(jan, "2024-01-01", "max") == jan
    assert anexo._pick_period_boundary(jan, "sem data", "min") == jan


def test_build_sections_groups_gtins_and_ignores_rows_before_first_header():
    rows = [
        _row("resumo_parcial", medicamento="ignorado"),
        _row("header_medicamento", gtin="111"),
        _row("resumo_parcial", medicamento="Medicamento A", valor_irregular=10),
        _row("header_colunas"),
        _row("venda_normal", medicamento="", valor=20),
        _row("venda_irregular", medicamento="", valor=10),
        _row("outra_linha", value=1),
        _row("header_medicamento", gtin="222"),
        _row("venda_normal", medicamento="Medicamento B"),
    ]
    sections = anexo._build_sections(rows)
    assert len(sections) == 2
    assert sections[0]["gtin"] == "111"
    assert sections[0]["medicamento"] == "Medicamento A"
    assert sections[0]["subtotal"]["valor_irregular"] == 10
    assert [row["tipo_linha"] for row in sections[0]["rows"]] == ["venda_normal", "venda_irregular", "outra_linha"]
    assert sections[1]["gtin"] == "222"
    assert sections[1]["medicamento"] == "Medicamento B"
    assert anexo._build_sections([]) == []


def test_build_context_calculates_periods_subtotals_and_percentages(monkeypatch):
    rows = [
        _row("header_medicamento", gtin=111),
        _row("venda_normal", medicamento="ANTIBIOTICO", vendas=3, valor=30, estoque_final=50),
        _row(
            "venda_irregular", medicamento="ANTIBIOTICO", vendas=2, vendas_irregular=1,
            valor=20, valor_irregular=12.5, periodo_inicio_irregular="2020-02-01",
            periodo_final="2020-11-30", estoque_final=40,
        ),
        _row("resumo_parcial", gtin=111, medicamento="ANTIBIOTICO", valor=50, valor_irregular=12.5,
             vendas=5, vendas_irregular=1, estoque_final=None),
        _row("header_medicamento", gtin=222),
        _row(
            "venda_irregular", medicamento="ANALGESICO", vendas=4, vendas_irregular=2,
            valor=80, valor_irregular=7.25, periodo_inicio_irregular="01/03/2021",
            periodo_final="31.12.2021", estoque_final=9,
        ),
        _row("resumo_parcial", gtin=222, valor=80, valor_irregular=0, vendas=4, vendas_irregular=0, estoque_final=8),
        _row("header_medicamento", gtin=333),
        _row("venda_irregular", medicamento="SEM SUBTOTAL", valor=10, valor_irregular=3,
             vendas=2, vendas_irregular=1, estoque_final=7),
        _row("resumo_parcial", valor_irregular=0, vendas=0, vendas_irregular=0, estoque_final=None),
        _row("header_medicamento", gtin=444),
        _row("venda_irregular", medicamento="SEM DATAS", valor=5, valor_irregular=2,
             vendas=1, vendas_irregular=1, periodo_inicio_irregular="x", periodo_final="y", estoque_final=6),
        _row("resumo_parcial", valor_irregular=2, vendas=1, vendas_irregular=1, estoque_final=None),
        _row("header_medicamento", gtin=555),
        _row("venda_irregular", medicamento="ZERO", valor_irregular=0),
    ]
    movimentacao = SimpleNamespace(summary={"valor_irregular": 25.0}, rows=rows)
    captured = {}

    def load(cnpj, engine):
        captured.update(cnpj=cnpj, engine=engine)
        return movimentacao

    monkeypatch.setattr(anexo, "get_movimentacao_data", load)
    engine = Engine()
    context = anexo._build_anexo_ii_context("123", engine)
    assert captured == {"cnpj": "123", "engine": engine}
    assert context["summary"] == {"valor_irregular": 25.0}
    assert context["total_gtins_irregulares"] == 4
    assert [item["gtin"] for item in context["consolidado"]] == ["111", "222", "333", "444"]
    first = context["consolidado"][0]
    assert first["periodo_sem_comprovacao"] == "01.02.2020 a 30.11.2020"
    assert first["estoque_final"] == 40
    assert first["vendas"] == 5
    assert first["valor_irregular"] == 12.5
    assert first["pct_prejuizo_total"] == 50.0
    assert context["consolidado"][1]["valor_irregular"] == 7.25
    assert context["detalhes"][1]["rows"][0]["tipo_linha"] == "venda_irregular"
    assert context["consolidado"][2]["periodo_sem_comprovacao"] == "—"
    assert context["consolidado"][3]["pct_prejuizo_total"] == 8.0


def test_build_context_supports_non_database_movement_contract(monkeypatch):
    rows = [
        _row("header_medicamento", gtin=None),
        _row("venda_irregular", medicamento=None, valor=1, valor_irregular=5,
             vendas=1, vendas_irregular=1, periodo_inicio_irregular="01/02/2020", periodo_final="2020-03-01"),
        _row("resumo_parcial", valor_irregular=5, estoque_final=None),
    ]
    monkeypatch.setattr(
        anexo,
        "get_movimentacao_data",
        lambda *args: SimpleNamespace(summary={"valor_irregular": 0}, rows=rows),
    )
    context = anexo._build_anexo_ii_context("123", object())
    item = context["consolidado"][0]
    assert item["gtin"] == ""
    assert item["medicamento"] == "NÃO IDENTIFICADO"
    assert item["estoque_final"] == 0
    assert item["vendas"] == 1
    assert item["valor"] == 1.0
    assert item["pct_prejuizo_total"] == 0.0


def test_detail_renderer_skips_empty_inputs_and_draws_rows_and_subtotal():
    doc = Document()
    assert anexo._add_anexo_ii_detalhamento(doc, []) is None
    timing = Timing()
    details = [
        {"gtin": "skip", "rows": []},
        {
            "gtin": "123", "medicamento": "MEDICAMENTO A", "vendas": 10,
            "vendas_irregular": 2, "valor": 100.5, "valor_irregular": 20.25,
            "rows": [
                {
                    "periodo_inicial": "01.01.2020", "periodo_inicio_irregular": "01.06.2020",
                    "periodo_final": "30.06.2020", "estoque_inicial": 1000, "estoque_final": 20,
                    "vendas": 10, "vendas_irregular": 2, "valor": 100.5,
                    "valor_irregular": 20.25, "notas": "NF 1", "tipo_linha": "venda_irregular",
                },
                {
                    "periodo_inicial": None, "periodo_inicio_irregular": None, "periodo_final": None,
                    "estoque_inicial": 0, "estoque_final": 0, "vendas": 1, "vendas_irregular": 0,
                    "valor": 5, "valor_irregular": 0, "notas": None, "tipo_linha": "venda_normal",
                },
            ],
        },
    ]
    anexo._add_anexo_ii_detalhamento(doc, details, timing=timing, anexo_num="IV")
    assert len(doc.tables) == 1
    assert len(doc.tables[0].rows) == 4
    assert timing.marks == ["anexo IV detalhe GTIN 2 (2 linhas)"]


@pytest.mark.parametrize("with_consolidated", [True, False])
def test_main_annex_renderer_creates_summary_and_consolidated_table(with_consolidated):
    doc = Document()
    timing = Timing()
    item = {
        "gtin": "123456", "medicamento": "MEDICAMENTO A", "periodo_sem_comprovacao": "01.01.2020 a 31.12.2020",
        "estoque_final": 20, "vendas": 10, "vendas_irregular": 2,
        "valor": 100.5, "valor_irregular": 20.25, "pct_prejuizo_total": 80.0,
    }
    detail = {
        **item,
        "rows": [{
            "periodo_inicial": "01.01.2020", "periodo_inicio_irregular": "01.06.2020",
            "periodo_final": "31.12.2020", "estoque_inicial": 30, "estoque_final": 20,
            "vendas": 10, "vendas_irregular": 2, "valor": 100.5,
            "valor_irregular": 20.25, "notas": "NF 1", "tipo_linha": "venda_irregular",
        }],
    }
    context = {
        "summary": {
            "total_vendas": 100, "total_vendas_irregular": 2, "valor_total": 1000,
            "valor_irregular": 20.25, "pct_irregular": 2.025,
        },
        "consolidado": [item] if with_consolidated else [],
        "detalhes": [detail] if with_consolidated else [],
    }
    anexo._add_anexo_ii_memoria_calculo(
        doc, "Farmacia Teste", "12.345.678/0001-90", "2020 a 2024", context,
        "NT-10", timing=timing, anexo_num="IV",
    )
    assert len(doc.sections) == 2
    assert doc.sections[-1].orientation
    assert len(doc.tables) == (3 if with_consolidated else 2)
    assert any("ANEXO IV" in paragraph.text for paragraph in doc.paragraphs)
    assert any("detalhamento total" in mark for mark in timing.marks)
