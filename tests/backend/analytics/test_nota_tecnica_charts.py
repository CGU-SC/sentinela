import io
import json
from pathlib import Path

import pytest
from docx import Document

from backend.api.services.analytics import nota_tecnica_charts as charts


def test_currency_axis_rounding_and_svg_escaping():
    assert charts._axis_currency_label(999) == "R$ 999"
    assert charts._axis_currency_label(1000) == "R$ 1 mil"
    assert charts._axis_currency_label(1_000_000) == "R$ 1,0 mi"
    assert charts._nice_axis_max(-1) == 1
    assert charts._nice_axis_max(1) == 1
    assert charts._nice_axis_max(15) == 20
    assert charts._nice_axis_max(33) == 50
    assert charts._nice_axis_max(75) == 100
    assert charts._svg_escape('<A & "B">') == "&lt;A &amp; &quot;B&quot;&gt;"
    assert charts._svg_escape("ação") == "a&#231;&#227;o"
    assert charts._svg_currency_axis_label(1000) == "R$ 1 mil"


def test_geojson_polygon_helpers_support_polygons_holes_and_degenerate_rings():
    polygon = {"type": "Polygon", "coordinates": [[[0, 0], [2, 0], [2, 2], [0, 2], [0, 0]], [[1, 1]]]}
    multipolygon = {"type": "MultiPolygon", "coordinates": [[[[0, 0], [1, 0], [1, 1]]], [[[4, 4], [5, 4], [5, 5]]]]}
    assert len(list(charts._iter_geojson_polygon_rings(polygon))) == 1
    assert len(list(charts._iter_geojson_polygon_rings(multipolygon))) == 2
    assert list(charts._iter_geojson_polygon_rings({"type": "Point", "coordinates": [0, 0]})) == []
    area, cx, cy = charts._polygon_area_and_centroid([(0, 0), (2, 0), (2, 2), (0, 2)])
    assert (area, cx, cy) == (4, 1, 1)
    assert charts._polygon_area_and_centroid([(1, 2), (3, 4)]) == (0, 0, 0)
    area, cx, cy = charts._polygon_area_and_centroid([(0, 0), (1, 1), (2, 2)])
    assert area == 0 and (cx, cy) == (1, 1)
    assert charts._weighted_centroid_from_rings([]) is None
    assert charts._weighted_centroid_from_rings([[(0, 0), (2, 0), (2, 2), (0, 2)]]) == (1, 1)
    assert charts._weighted_centroid_from_rings([[(1, 2), (3, 4)]]) == (2, 3)


def test_svg_curve_helpers_handle_empty_single_duplicate_and_plateau_points():
    assert charts._svg_smooth_path([]) == ""
    assert charts._svg_smooth_path([(3, 4)]) == "M 3.00 4.00"
    path = charts._svg_smooth_path([(0, 0), (4, 2)])
    assert "C 2.00 0.00, 2.00 2.00, 4.00 2.00" in path
    assert charts._percentile_curve_anchors([], []) == [(1.0, 0.0), (100.0, 0.0)]
    anchors = charts._percentile_curve_anchors([40, 10, 20, 20, 30], [3, 1, 1, 9, 2])
    assert anchors[0] == (10.0, 1)
    assert anchors[-1] == (40.0, 3)
    assert all(left[0] <= right[0] for left, right in zip(anchors, anchors[1:]))
    assert charts._svg_int_label(12345.6) == "12.346"

    import numpy as np

    assert charts._monotone_smooth_curve([2, 1], [5, 3], np=np) == ([2.0, 1.0], [5, 3])
    x_smooth, y_smooth = charts._monotone_smooth_curve(
        [30, 10, 20, 40, 50], [6, 1, 1, 8, 8], np=np
    )
    assert len(x_smooth) >= 1200
    assert all(left <= right for left, right in zip(y_smooth, y_smooth[1:]))
    assert charts._monotone_smooth_curve([4, 4, 4], [2, 1, 3], np=np) == (
        [4.0], [2.0]
    )


def test_financial_and_risk_svg_builders_render_axes_labels_and_empty_data():
    financial = charts._build_evolucao_financeira_chart_svg(
        {
            "rows": [
                {"semestre_fmt": "1º semestre", "regular": 800.0, "irregular": 200.0, "total": 1000.0},
                {"semestre": "2024-S2", "regular": 0.0, "irregular": 0.0, "total": 0.0},
            ]
        }
    )
    assert "Evolu&#231;&#227;o semestral" in financial
    assert "1&#176; semestre" in financial or "1º semestre" in financial
    assert "R$ 1 mil" in financial
    empty_financial = charts._build_evolucao_financeira_chart_svg({"rows": []})
    assert "<svg" in empty_financial and "</svg>" in empty_financial

    risk = charts._build_percentil_risco_chart_svg(
        {
            "metric_label": "Risco <atual>", "current_value": 25, "percentile_rank": 80,
            "percentiles": [{"percentile": 20, "score": 3}, {"percentile": 60, "score": 12},
                            {"percentile": 90, "score": 28}],
        }
    )
    assert "Percentil 80" in risk and "&lt;atual&gt;" in risk
    flat_risk = charts._build_percentil_risco_chart_svg({"percentiles": [], "current_value": 0})
    assert "Estabelecimento" in flat_risk and "Percentil 100" in flat_risk


def test_regional_position_svg_uses_current_and_comparison_pharmacies():
    svg = charts._build_posicionamento_regional_chart_svg(
        {
            "metric_label": "% sob análise",
            "rows": [
                {"total_mov": 5000, "pct_sem_comprovacao": 4, "is_current": False},
                {"total_mov": 10000, "pct_sem_comprovacao": 10, "is_current": True},
                {"total_mov": 20000, "pct_sem_comprovacao": 8, "is_current": False},
            ],
            "current": {"total_mov": 10000, "pct_sem_comprovacao": 10},
        }
    )
    assert "Posicionamento regional" in svg
    assert "Mediana regional" in svg
    assert "% sob an&#225;lise" in svg
    assert "Outras farm&#225;cias" in svg
    empty = charts._build_posicionamento_regional_chart_svg({})
    assert "<svg" in empty and "Estabelecimento" in empty


def test_parkinson_svg_charts_render_values_and_reject_invalid_age_bands():
    demographics = {
        "casos_esperados": 120, "qtd_cpfs_distintos_observado": 15,
        "ano_observado": 2024, "municipio": "São Paulo", "uf": "SP", "razao_observado_esperado": 0.125,
        "populacao_50_mais": 10000, "percentual_50_mais": 0.35,
        "faixas_etarias": [
            {"faixa": "0-49", "populacao": 15000, "destacar_50_mais": False},
            {"faixa": "50+", "populacao": 10000, "destacar_50_mais": True},
        ],
    }
    comparison = charts._build_parkinson_demografia_chart_svg(demographics)
    assert "S&#227;o Paulo/SP" in comparison and "Casos esperados" in comparison and "120" in comparison
    age_bands = charts._build_parkinson_faixas_etarias_chart_svg(demographics)
    assert "50+: 10.000 pessoas (35,00%)" in age_bands
    assert "url(#age50)" in age_bands and "0-49" in age_bands
    with pytest.raises(RuntimeError, match="Faixas etarias obrigatorias"):
        charts._build_parkinson_faixas_etarias_chart_svg({"faixas_etarias": []})
    with pytest.raises(RuntimeError, match="Populacao por faixa etaria invalida"):
        charts._build_parkinson_faixas_etarias_chart_svg({"faixas_etarias": [{"faixa": "50+", "populacao": 0}]})


def test_map_builder_uses_geojson_values_and_validates_missing_inputs(tmp_path, monkeypatch):
    geojson = {
        "type": "FeatureCollection",
        "features": [
            {
                "type": "Feature", "properties": {"UF": "SP"},
                "geometry": {"type": "Polygon", "coordinates": [[[0, 0], [2, 0], [2, 2], [0, 2], [0, 0]]]},
            },
            {
                "type": "Feature", "properties": {"UF": "RJ"},
                "geometry": {"type": "Polygon", "coordinates": [[[3, 0], [5, 0], [5, 2], [3, 2], [3, 0]]]},
            },
            {
                "type": "Feature", "properties": {"UF": "AC"},
                "geometry": {"type": "Polygon", "coordinates": [[]]},
            },
        ],
    }
    path = tmp_path / "brasil-uf.json"
    path.write_text(json.dumps(geojson), encoding="utf-8")
    monkeypatch.setattr(charts, "_brasil_uf_geojson_path", lambda: path)
    output = charts._build_mapa_geografico_origem_uf(
        {
            "uf_farmacia": "SP",
            "origem_uf_rows": [
                {"uf_paciente": "SP", "percentual_sobre_total": 80, "valor_autorizado": 1000},
                {"uf_paciente": "RJ", "percentual_sobre_total": 20, "valor_autorizado": 250},
            ],
        }
    )
    assert output.read(8) == b"\x89PNG\r\n\x1a\n"
    with pytest.raises(RuntimeError, match="Dados de origem geografica por UF ausentes"):
        charts._build_mapa_geografico_origem_uf({"origem_uf_rows": []})
    path.write_text(json.dumps({"features": [{"properties": {"UF": "SP"}, "geometry": {"type": "Point", "coordinates": [1, 2]}}]}), encoding="utf-8")
    with pytest.raises(RuntimeError, match="sem coordenadas validas"):
        charts._build_mapa_geografico_origem_uf({"uf_farmacia": "SP", "origem_uf_rows": [{"uf_paciente": "SP"}]})


def test_png_wrappers_and_svg_preference_paths(monkeypatch):
    png = lambda _: io.BytesIO(b"png-data")
    monkeypatch.setattr(charts, "_svg_to_png_stream", lambda _svg: io.BytesIO(b"svg-png"))
    assert charts._build_evolucao_financeira_chart_svg_png({"rows": []}).read() == b"svg-png"
    assert charts._build_percentil_risco_chart_svg_png({}).read() == b"svg-png"
    assert charts._build_posicionamento_regional_chart_svg_png({}).read() == b"svg-png"
    assert charts._build_parkinson_demografia_chart_svg_png({}).read() == b"svg-png"
    valid_demography = {"faixas_etarias": [{"populacao": 1}]}
    assert charts._build_parkinson_faixas_etarias_chart_svg_png(valid_demography).read() == b"svg-png"

    monkeypatch.setattr(charts, "_build_evolucao_financeira_chart_svg_png", png)
    assert charts._build_evolucao_financeira_chart_prefer_svg({}).read() == b"png-data"
    monkeypatch.setattr(charts, "_build_percentil_risco_chart_svg_png", png)
    assert charts._build_percentil_risco_chart_prefer_svg({}).read() == b"png-data"
    monkeypatch.setattr(charts, "_build_posicionamento_regional_chart_svg_png", png)
    assert charts._build_posicionamento_regional_chart_prefer_svg({}).read() == b"png-data"
    monkeypatch.setattr(charts, "_build_parkinson_demografia_chart_svg_png", png)
    assert charts._build_parkinson_demografia_chart_prefer_svg({}).read() == b"png-data"
    monkeypatch.setattr(charts, "_build_parkinson_faixas_etarias_chart_svg_png", png)
    assert charts._build_parkinson_faixas_etarias_chart_prefer_svg({}).read() == b"png-data"

    monkeypatch.setattr(charts, "_build_evolucao_financeira_chart_svg_png", lambda _: (_ for _ in ()).throw(ValueError("svg")))
    monkeypatch.setattr(charts, "_build_evolucao_financeira_chart", lambda _: io.BytesIO(b"matplotlib"))
    assert charts._build_evolucao_financeira_chart_prefer_svg({}).read() == b"matplotlib"
    monkeypatch.setattr(charts, "_build_percentil_risco_chart_svg_png", lambda _: (_ for _ in ()).throw(ValueError("svg")))
    monkeypatch.setattr(charts, "_build_percentil_risco_chart", lambda _: io.BytesIO(b"matplotlib"))
    assert charts._build_percentil_risco_chart_prefer_svg({}).read() == b"matplotlib"
    monkeypatch.setattr(charts, "_build_posicionamento_regional_chart_svg_png", lambda _: (_ for _ in ()).throw(ValueError("svg")))
    monkeypatch.setattr(charts, "_build_posicionamento_regional_chart", lambda _: io.BytesIO(b"matplotlib"))
    assert charts._build_posicionamento_regional_chart_prefer_svg({}).read() == b"matplotlib"


def test_chart_docx_figures_use_centered_title_and_footnote(monkeypatch):
    doc = Document()
    image_bytes = __import__("base64").b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+j1ioAAAAASUVORK5CYII="
    )
    monkeypatch.setattr(charts, "_build_evolucao_financeira_chart_prefer_svg", lambda _: io.BytesIO(
        image_bytes
    ))
    monkeypatch.setattr(
        charts,
        "_build_parkinson_faixas_etarias_chart_prefer_svg",
        lambda _: io.BytesIO(image_bytes),
    )
    charts._add_figura_evolucao_financeira(doc, "Farmácia X", "00.000.000/0001-91", {"rows": []})
    charts._add_figura_posicionamento_regional(
        doc, "Farmácia X", "00.000.000/0001-91", {}, figure_number=2
    )
    charts._add_figura_percentil_risco(
        doc, "Farmácia X", "00.000.000/0001-91", {}, figure_number=3
    )
    charts._add_figura_percentil_risco(
        doc, "Farmácia X", "00.000.000/0001-91", {}, figure_number=4, show_title=False
    )
    charts._add_figura_parkinson_comparacao(doc, {}, figure_number=5)
    charts._add_figura_parkinson_faixas_etarias(doc, {}, figure_number=6)
    assert len(doc.paragraphs) >= 3
    assert doc.paragraphs[0].alignment is not None


def test_matplotlib_renderers_produce_png_for_fallback_and_edge_cases():
    financial = charts._build_evolucao_financeira_chart(
        {
            "rows": [
                {"semestre_fmt": "2024 S1", "regular": 1200, "irregular": 450},
                {"semestre": "2024 S2", "regular": 0, "irregular": 800, "total": 800},
                {"semestre": "2025 S1", "regular": 0, "irregular": 0, "total": 0},
            ]
        }
    )
    assert financial.read(8) == b"\x89PNG\r\n\x1a\n"

    percentile = charts._build_percentil_risco_chart(
        {
            "percentiles": [
                {"percentile": 20, "score": 2},
                {"percentile": 40, "score": 2},
                {"percentile": 60, "score": 5},
                {"percentile": 80, "score": 9},
            ],
            "current_value": 6,
            "percentile_rank": 70,
            "metric_label": "Indicador auditável",
        }
    )
    assert percentile.read(8) == b"\x89PNG\r\n\x1a\n"

    empty_percentile = charts._build_percentil_risco_chart({"percentiles": []})
    assert empty_percentile.read(8) == b"\x89PNG\r\n\x1a\n"

    regional = charts._build_posicionamento_regional_chart(
        {
            "metric_label": "Percentual sob análise",
            "rows": [
                {"total_mov": 100, "pct_sem_comprovacao": 3},
                {"total_mov": 700, "pct_sem_comprovacao": 8, "is_current": True},
                {"total_mov": 300, "pct_sem_comprovacao": 5},
            ],
            "current": {"total_mov": 700, "pct_sem_comprovacao": 8},
        }
    )
    assert regional.read(8) == b"\x89PNG\r\n\x1a\n"

    empty_regional = charts._build_posicionamento_regional_chart({})
    assert empty_regional.read(8) == b"\x89PNG\r\n\x1a\n"


def test_svg_converter_returns_nonempty_png_bytes():
    png = charts._svg_to_png_stream(
        '<svg xmlns="http://www.w3.org/2000/svg" width="2" height="2">'
        '<rect width="2" height="2" fill="#fff"/></svg>'
    )
    assert png.tell() == 0
    assert png.read(8) == b"\x89PNG\r\n\x1a\n"


def test_svg_converter_rejects_unsupported_and_empty_renderer_results(monkeypatch):
    import resvg_py

    monkeypatch.setattr(resvg_py, "svg_to_bytes", lambda **_kwargs: bytearray(b"not bytes"))
    with pytest.raises(RuntimeError, match="unsupported type"):
        charts._svg_to_png_stream("<svg/>")

    monkeypatch.setattr(resvg_py, "svg_to_bytes", lambda **_kwargs: b"")
    with pytest.raises(RuntimeError, match="empty PNG bytes"):
        charts._svg_to_png_stream("<svg/>")


def test_map_asset_path_handles_frozen_runtime_and_missing_geojson(tmp_path, monkeypatch):
    import sys

    monkeypatch.setattr(sys, "frozen", True, raising=False)
    monkeypatch.setattr(sys, "_MEIPASS", str(tmp_path), raising=False)
    assert charts._brasil_uf_geojson_path() == tmp_path / "frontend" / "dist" / "geo" / "brasil-uf.json"

    missing = tmp_path / "does-not-exist.json"
    monkeypatch.setattr(charts, "_brasil_uf_geojson_path", lambda: missing)
    with pytest.raises(RuntimeError, match="GeoJSON de UF do Brasil nao encontrado"):
        charts._build_mapa_geografico_origem_uf({
            "origem_uf_rows": [{"uf_paciente": "SP", "percentual_sobre_total": 100, "valor_autorizado": 1}],
        })


def test_map_asset_path_uses_frontend_public_assets_outside_frozen_runtime(monkeypatch):
    import sys

    monkeypatch.setattr(sys, "frozen", False, raising=False)
    monkeypatch.delattr(sys, "_MEIPASS", raising=False)

    expected = Path(charts.__file__).resolve().parents[4] / "frontend" / "public" / "geo" / "brasil-uf.json"
    assert charts._brasil_uf_geojson_path() == expected


def test_map_figure_wrapper_adds_title_image_and_source(monkeypatch):
    image_bytes = __import__("base64").b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+j1ioAAAAASUVORK5CYII="
    )
    monkeypatch.setattr(
        charts,
        "_build_mapa_geografico_origem_uf",
        lambda _context: io.BytesIO(image_bytes),
    )
    doc = Document()

    charts._add_mapa_geografico_origem_uf(
        doc,
        "Farmácia Teste",
        {"origem_uf_rows": [{"uf_paciente": "SP"}]},
    )

    assert len(doc.inline_shapes) == 1
    assert "Mapa 01" in doc.paragraphs[0].text
    assert "Fonte: Sentinela" in doc.paragraphs[-1].text
