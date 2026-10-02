from datetime import date

from docx import Document
import pytest

from backend.api.services.analytics import nota_tecnica_esocial as esocial


def test_esocial_plural_month_distance_and_summary_helpers_validate_periods():
    assert esocial._plural(1, "vínculo", "vínculos") == "vínculo"
    assert esocial._plural(0, "vínculo", "vínculos") == "vínculos"
    assert esocial._capitalize_first("") == ""
    assert esocial._capitalize_first("setembro") == "Setembro"
    assert esocial._month_distance(date(2023, 12, 1), date(2024, 2, 1)) == 2
    with pytest.raises(RuntimeError, match="Periodo inicial sem ano/mes"):
        esocial._month_distance("2023-01", date(2023, 2, 1))
    with pytest.raises(RuntimeError, match="Periodo final sem ano/mes"):
        esocial._month_distance(date(2023, 1, 1), "2023-02")
    with pytest.raises(RuntimeError, match="Periodo final anterior"):
        esocial._month_distance(date(2024, 1, 1), date(2023, 12, 1))

    annual = esocial._format_annual_summary(
        [
            {"ano_base": 2023, "qtd_trabalhadores_vinculo_ano": 1, "qtd_farmaceuticos_vinculo_ano": 1},
            {"ano_base": 2024, "qtd_trabalhadores_vinculo_ano": 2, "qtd_farmaceuticos_vinculo_ano": 0},
        ]
    )
    assert "2023" in annual and "2024" in annual
    assert esocial._format_single_worker_summary(
        [
            {"ano_base": 2023, "trabalhador_unico_cbo_txt": "223405", "trabalhador_unico_dt_admissao_txt": "01/01/2023"},
            {"ano_base": 2024, "trabalhador_unico_cbo_txt": None, "trabalhador_unico_dt_admissao_txt": None},
        ]
    ) == "2023 (223405, admissão em 01/01/2023) e 2024 (CBO não identificado)"


def test_esocial_context_without_data_returns_same_table_number():
    doc = Document()
    result = esocial._add_esocial_context_text(doc, "Farmacia", "123", {"has_data": False}, 8)
    assert result == 8
    assert any("Não foram identificados registros" in paragraph.text for paragraph in doc.paragraphs)
    assert doc.tables == []


def test_esocial_context_one_year_with_pharmacist_and_active_composition(monkeypatch):
    calls = []
    monkeypatch.setattr(esocial, "_add_quadro_esocial", lambda *args: calls.append(("annual", args[-1])))
    monkeypatch.setattr(esocial, "_add_quadro_esocial_trabalhadores", lambda *args: calls.append(("workers", None)))
    doc = Document()
    context = {
        "has_data": True,
        "dt_carga_fonte_txt": "31/12/2024",
        "rows": [
            {
                "ano_base": 2024, "qtd_trabalhadores_vinculo_ano": 2, "qtd_farmaceuticos_vinculo_ano": 1,
                "qtd_trabalhadores_ativos_competencia": 1, "qtd_farmaceuticos_ativos_competencia": 0,
                "competencia_txt": "dezembro de 2024",
            }
        ],
        "anos_um_trabalhador": [
            {"ano_base": 2024, "trabalhador_unico_cbo_txt": "223405", "trabalhador_unico_dt_admissao_txt": "02/01/2024"}
        ],
    }
    result = esocial._add_esocial_context_text(doc, "Farmacia", "CNPJ", context, 5)
    joined = " ".join(paragraph.text for paragraph in doc.paragraphs)
    assert result == 6
    assert "somente durante o ano de 2024" in joined
    assert "CBO de farmacêutico" in joined
    assert "permaneciam ativos 1 trabalhador" in joined
    assert "apenas um trabalhador" in joined
    assert calls == [("annual", 6), ("workers", None)]


def test_esocial_context_one_year_without_pharmacist_adds_legal_context(monkeypatch):
    monkeypatch.setattr(esocial, "_add_quadro_esocial", lambda *_args: None)
    monkeypatch.setattr(esocial, "_add_quadro_esocial_trabalhadores", lambda *_args: None)
    doc = Document()
    result = esocial._add_esocial_context_text(
        doc,
        "Farmacia",
        "CNPJ",
        {
            "has_data": True,
            "rows": [{"ano_base": 2024, "qtd_trabalhadores_vinculo_ano": 1, "qtd_farmaceuticos_vinculo_ano": 0}],
        },
        0,
    )
    joined = " ".join(paragraph.text for paragraph in doc.paragraphs)
    assert result == 1
    assert "nenhum deles com CBO de farmacêutico" in joined
    assert "ausência de vínculo" in joined


@pytest.mark.parametrize(
    ("missing_years", "expected"),
    [
        ([2023, 2024], "Em nenhum dos anos considerados"),
        ([2024], "Nos anos de 2024"),
        ([], "Em todos os anos considerados"),
    ],
)
def test_esocial_context_multiple_year_summaries_and_missing_professionals(monkeypatch, missing_years, expected):
    monkeypatch.setattr(esocial, "_add_quadro_esocial", lambda *_args: None)
    monkeypatch.setattr(esocial, "_add_quadro_esocial_trabalhadores", lambda *_args: None)
    doc = Document()
    rows = [
        {"ano_base": 2023, "qtd_trabalhadores_vinculo_ano": 2, "qtd_farmaceuticos_vinculo_ano": 0},
        {"ano_base": 2024, "qtd_trabalhadores_vinculo_ano": 1, "qtd_farmaceuticos_vinculo_ano": 1},
    ]
    context = {
        "has_data": True,
        "periodo_anos_txt": "2023 a 2024",
        "rows": rows,
        "anos_sem_farmaceutico": [{"ano_base": year} for year in missing_years],
    }
    result = esocial._add_esocial_context_text(doc, "Farmacia", "CNPJ", context, 10)
    joined = " ".join(paragraph.text for paragraph in doc.paragraphs)
    assert result == 11
    assert expected in joined
    assert "composição anual" in joined


def test_esocial_unstaffed_movement_alert_checks_contract_and_renders_summary(monkeypatch):
    assert esocial._add_movimentacao_sem_funcionario_alerta(Document(), {}, 3) == 3
    with pytest.raises(RuntimeError, match="campos obrigatorios"):
        esocial._add_movimentacao_sem_funcionario_alerta(
            Document(), {"movimentacao_sem_funcionario_alerta": {"ultimo_periodo_movimentacao": None}}, 3
        )

    calls = []
    monkeypatch.setattr(esocial, "_add_movimentacao_sem_funcionario_table", lambda doc, **kwargs: calls.append(kwargs))
    doc = Document()
    context = {
        "movimentacao_sem_funcionario_alerta": {
            "ultimo_periodo_movimentacao": date(2024, 3, 1),
            "ultimo_periodo_movimentacao_txt": "março de 2024",
            "ultimo_mes_trabalhador_ativo": date(2024, 1, 1),
            "ultimo_mes_trabalhador_ativo_txt": "janeiro de 2024",
            "valor_pfpb_periodo_sem_funcionario": 1234.5,
            "qtd_autorizacoes_periodo_sem_funcionario": 2,
            "ano_ultima_movimentacao": 2024,
            "ano_esocial_referencia_ultima_movimentacao": 2023,
            "is_sem_esocial_no_ano_ultima_movimentacao": True,
        }
    }
    assert esocial._add_movimentacao_sem_funcionario_alerta(doc, context, 6) == 7
    joined = " ".join(paragraph.text for paragraph in doc.paragraphs)
    assert "Não foram identificados vínculos no eSocial para 2024" in joined
    assert calls[0]["qtd_meses_sem_funcionario"] == 2
    assert calls[0]["qtd_autorizacoes_periodo_txt"] == "2"


def test_esocial_movement_summary_table_writes_required_rows():
    doc = Document()
    esocial._add_movimentacao_sem_funcionario_table(
        doc,
        tabela_num=2,
        ultimo_mes_ativo_txt="janeiro de 2024",
        periodo_txt="março de 2024",
        qtd_meses_txt="2",
        qtd_meses_sem_funcionario=2,
        valor_periodo_txt="1.234,50",
        qtd_autorizacoes_periodo_txt="3",
    )
    assert len(doc.tables) == 1
    assert len(doc.tables[0].rows) == 6
    assert "Tabela 2" in " ".join(paragraph.text for paragraph in doc.paragraphs)
